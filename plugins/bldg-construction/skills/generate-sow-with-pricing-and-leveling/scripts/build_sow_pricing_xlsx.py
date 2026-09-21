#!/usr/bin/env python3
"""Build the SOW + pricing + leveling workbook.

usage: build_sow_pricing_xlsx.py <data.json> <out.xlsx>

Schema of <data.json> -- see references/workbook-spec.md for the full contract.
{
  "project","address","date","prepared_by","job","client",
  "trades":[{"trade":str,"csi":[str],"lines":[
      {"scope":str,"qty":float|null,"unit":str,"bullets":[str]}]}],
  "notes":[{"type":"EXCLUSION"|"CLARIFICATION","item":str,"text":str}],
  "basis":[{"trade","csi","scope","qty","unit","basis"}]      # optional, hidden sheet
}
A null qty renders "--", shades the row red, and is excluded from every total.
"""
import json, math, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

W="FFFFFFFF"; BLACK="FF000000"; CHAR="FF3A3A3A"; MED="FF7A7A7A"
LT="FFE8E8E8"; VLT="FFF5F5F5"; RED="FFFFD6D6"
FONT="Calibri"; CUR='$#,##0.00'
HAIR=Side("thin",color="FFB0B0B0"); MEDS=Side("medium",color=CHAR); BLKS=Side("medium",color=BLACK)
COLW=[56,10,11,13,14,2.5,15,15,15,15]          # A..J
LAST=10; GUT=6                                  # gutter column
HEADS=["Scope Item","Qty","Unit Type","Unit Rate","Total","","Sub #1","Sub #2","Sub #3","Sub #4"]

def ch(spec,mn=22,lp=13):
    """Row height that fits wrapped text. NEVER write a flat height -- it kills auto-fit."""
    n=max([1]+[math.ceil(len(str(t))/c) for t,c in spec if t])
    return max(mn,n*lp)

def paint(ws,r1,c1,r2,c2,edge,inner=HAIR):
    """Border EVERY cell in the range. Excel draws a merged range's edges from all of its
    constituent cells, so bordering only the anchor leaves the box three-quarters open."""
    for r in range(r1,r2+1):
        for c in range(c1,c2+1):
            cur=ws.cell(r,c).border
            ws.cell(r,c).border=Border(
                left = edge if c==c1 else (inner or cur.left),
                right= edge if c==c2 else (inner or cur.right),
                top  = edge if r==r1 else (inner or cur.top),
                bottom=edge if r==r2 else (inner or cur.bottom))

def erow(ws,r,c1,c2,top=None,bottom=None):
    for c in range(c1,c2+1):
        b=ws.cell(r,c).border
        ws.cell(r,c).border=Border(left=b.left,right=b.right,top=top or b.top,bottom=bottom or b.bottom)

def ecol(ws,r1,r2,c,l=None,rt=None):
    for r in range(r1,r2+1):
        b=ws.cell(r,c).border
        ws.cell(r,c).border=Border(left=l or b.left,right=rt or b.right,top=b.top,bottom=b.bottom)

def build(d,out):
    wb=Workbook(); ws=wb.active; ws.title="Scope of Work"
    for col,w in zip("ABCDEFGHIJ",COLW): ws.column_dimensions[col].width=w
    for col in "FGHIJ":
        ws.column_dimensions[col].outline_level=1; ws.column_dimensions[col].hidden=False
    M=[]   # merges applied AFTER borders, never before

    ws.cell(1,1,d["project"]).font=Font(FONT,18,bold=True); ws.row_dimensions[1].height=28; M.append((1,1,1,LAST))
    ws.cell(2,1,d.get("address","")).font=Font(FONT,11); M.append((2,1,2,LAST))
    ws.cell(3,1,"Date: %s"%d.get("date","")).font=Font(FONT,10); M.append((3,1,3,1))
    ws.cell(3,2,"Prepared By: %s   |   Job %s   |   Client: %s"%(
        d.get("prepared_by",""),d.get("job",""),d.get("client",""))).font=Font(FONT,10)
    M.append((3,2,3,LAST)); ws.row_dimensions[4].height=5

    x=ws.cell(5,1,"F/I = Furnish and Install  |  F/O = Furnish Only  |  I/O = Install Only  |  "
                  "Remove = Demolition  |  Provide = Services  |  LS = Lump Sum")
    x.font=Font(FONT,9,italic=True,color=CHAR)
    for c in range(1,LAST+1): ws.cell(5,c).fill=PatternFill("solid",fgColor=VLT)
    ws.row_dimensions[5].height=16; M.append((5,1,5,LAST))
    x=ws.cell(6,1,"Lines shown without a quantity require a plan count and are not included in "
                  "any total. See Notes & Clarifications.")
    x.font=Font(FONT,9,bold=True,color=CHAR)
    for c in range(1,LAST+1): ws.cell(6,c).fill=PatternFill("solid",fgColor=RED)
    ws.row_dimensions[6].height=18; M.append((6,1,6,LAST))
    erow(ws,5,1,LAST,top=MEDS); erow(ws,6,1,LAST,bottom=MEDS)   # rules only, no verticals
    ws.row_dimensions[7].height=7

    HDR=8
    for col,t in enumerate(HEADS,1):
        c=ws.cell(HDR,col,t); c.font=Font(FONT,11,bold=True,color=W)
        c.fill=PatternFill("solid",fgColor=CHAR)
        c.alignment=Alignment("left" if col==1 else "center","center")
    ws.row_dimensions[HDR].height=22
    paint(ws,HDR,1,HDR,LAST,BLKS,inner=Side("thin",color=CHAR))
    ws.freeze_panes="A%d"%(HDR+1)

    r=HDR+1; trade_rows=[]
    for tr in d["trades"]:
        ws.row_dimensions[r].height=7; r+=1
        h1,h2=r,r+1
        c=ws.cell(h1,1,"%s  --  CSI %s"%(tr["trade"].upper(),", ".join(tr.get("csi",[]))))
        c.font=Font(FONT,11,bold=True,color=W); c.alignment=Alignment("left","center")
        for rr in (h1,h2):
            for cc in range(1,GUT): ws.cell(rr,cc).fill=PatternFill("solid",fgColor=MED)
        M.append((h1,1,h2,5))
        for col in range(GUT+1,LAST+1):
            a=ws.cell(h1,col,"Sub #%d"%(col-GUT)); a.font=Font(FONT,10,bold=True)
            a.fill=PatternFill("solid",fgColor=LT); a.alignment=Alignment("center","center")
            b=ws.cell(h2,col); b.number_format=CUR
            b.fill=PatternFill("solid",fgColor=VLT); b.alignment=Alignment("center","center")
        ws.row_dimensions[h1].height=19; ws.row_dimensions[h2].height=19
        r=h2+1; first=r; band=0
        for ln in tr["lines"]:
            q=ln.get("qty"); bullets=ln.get("bullets") or []
            fg = RED if q is None else (VLT if band%2 else None); band+=1
            s=ws.cell(r,1,ln["scope"]); s.font=Font(FONT,10,bold=bool(bullets))
            s.alignment=Alignment("left","center",wrap_text=True)
            qa=ws.cell(r,2, q if q is not None else "--")
            if q is not None: qa.number_format='#,##0.00'
            qa.alignment=Alignment("right","center"); qa.font=Font(FONT,10,bold=True)
            u=ws.cell(r,3,ln.get("unit","")); u.alignment=Alignment("center","center"); u.font=Font(FONT,10)
            ur=ws.cell(r,4); ur.number_format=CUR; ur.alignment=Alignment("right","center")
            tot=ws.cell(r,5,'=IF(AND(ISNUMBER($B%d),ISNUMBER($D%d)),$B%d*$D%d,"")'%(r,r,r,r))
            tot.number_format=CUR; tot.font=Font(FONT,10,bold=True); tot.alignment=Alignment("right","center")
            for col in range(GUT+1,LAST+1):
                p=ws.cell(r,col); p.number_format=CUR; p.alignment=Alignment("right","center")
            if fg:
                for cc in list(range(1,GUT))+list(range(GUT+1,LAST+1)):
                    ws.cell(r,cc).fill=PatternFill("solid",fgColor=fg)
            ws.row_dimensions[r].height=ch([(ln["scope"],COLW[0])]); r+=1
            for b in bullets:
                bc=ws.cell(r,1,"- "+b); bc.font=Font(FONT,9,italic=True,color=CHAR)
                bc.alignment=Alignment("left","center",wrap_text=True,indent=2)
                ws.row_dimensions[r].height=ch([(b,COLW[0]-6)],mn=14,lp=11); r+=1
        last=r-1; tot_r=r
        c=ws.cell(tot_r,1,"TRADE TOTAL"); c.font=Font(FONT,10,bold=True)
        c.alignment=Alignment("left","center"); M.append((tot_r,1,tot_r,3))
        te=ws.cell(tot_r,5,'=IF(SUM(E%d:E%d)=0,"-",SUM(E%d:E%d))'%(first,last,first,last))
        te.number_format=CUR; te.font=Font(FONT,10,bold=True); te.alignment=Alignment("right","center")
        for col in range(GUT+1,LAST+1):
            L=get_column_letter(col)
            # sums from the LUMP-SUM row down, so lump-sum and itemised bidders both roll up
            p=ws.cell(tot_r,col,'=IF(SUM(%s%d:%s%d)=0,"-",SUM(%s%d:%s%d))'%(L,h2,L,last,L,h2,L,last))
            p.number_format=CUR; p.font=Font(FONT,10,bold=True); p.alignment=Alignment("right","center")
        for cc in list(range(1,GUT))+list(range(GUT+1,LAST+1)):
            ws.cell(tot_r,cc).fill=PatternFill("solid",fgColor=LT)
        ws.row_dimensions[tot_r].height=19
        paint(ws,h1,1,tot_r,LAST,MEDS,inner=HAIR)
        erow(ws,h2,1,LAST,bottom=MEDS); erow(ws,tot_r,1,LAST,top=MEDS)
        ecol(ws,h1,tot_r,GUT,l=MEDS,rt=MEDS); ecol(ws,h1,tot_r,GUT+1,l=MEDS)
        trade_rows.append((tr["trade"],", ".join(tr.get("csi",[])),tot_r))
        r=tot_r+1

    # ---- SUBTOTAL / markup ladder / TOTAL ----
    MONEY=[5]+list(range(GUT+1,LAST+1))
    ws.row_dimensions[r].height=7; r+=1
    r_sub=r
    c=ws.cell(r_sub,1,"SUBTOTAL"); c.font=Font(FONT,11,bold=True); c.alignment=Alignment("left","center")
    M.append((r_sub,1,r_sub,3))
    for col in MONEY:
        L=get_column_letter(col)
        ref="+".join("%s%d"%(L,tr) for _,_,tr in trade_rows) or "0"
        cc=ws.cell(r_sub,col,"=%s"%ref); cc.number_format=CUR
        cc.font=Font(FONT,11,bold=True); cc.alignment=Alignment("right","center")
    for cc in list(range(1,GUT))+list(range(GUT+1,LAST+1)):
        ws.cell(r_sub,cc).fill=PatternFill("solid",fgColor=LT)
    ws.row_dimensions[r_sub].height=20; r+=1

    r_m1=r; n_mk=int(d.get("markup_rows",6))
    labels=d.get("markup_labels") or []
    for i in range(n_mk):
        rr=r_m1+i
        lb=ws.cell(rr,1, labels[i] if i<len(labels) else None)
        lb.font=Font(FONT,10); lb.alignment=Alignment("left","center")
        lb.fill=PatternFill("solid",fgColor=VLT)
        M.append((rr,1,rr,3))
        pc=ws.cell(rr,4); pc.number_format='0.00%'
        pc.alignment=Alignment("right","center"); pc.font=Font(FONT,10,bold=True)
        pc.fill=PatternFill("solid",fgColor=VLT)
        for col in MONEY:
            L=get_column_letter(col)
            base="%s%d"%(L,r_sub) if i==0 else "%s%d+SUM(%s%d:%s%d)"%(L,r_sub,L,r_m1,L,rr-1)
            f='=IF($D%d="","",ROUND((%s)*$D%d,2))'%(rr,base,rr)
            cc=ws.cell(rr,col,f); cc.number_format=CUR; cc.alignment=Alignment("right","center")
        ws.row_dimensions[rr].height=18
    r=r_m1+n_mk
    r_tot=r
    c=ws.cell(r_tot,1,"TOTAL"); c.font=Font(FONT,12,bold=True); c.alignment=Alignment("left","center")
    M.append((r_tot,1,r_tot,3))
    for col in MONEY:
        L=get_column_letter(col)
        cc=ws.cell(r_tot,col,'=ROUND(%s%d+SUM(%s%d:%s%d),2)'%(L,r_sub,L,r_m1,L,r_m1+n_mk-1))
        cc.number_format=CUR; cc.font=Font(FONT,12,bold=True); cc.alignment=Alignment("right","center")
    for cc in list(range(1,GUT))+list(range(GUT+1,LAST+1)):
        ws.cell(r_tot,cc).fill=PatternFill("solid",fgColor=LT)
    ws.row_dimensions[r_tot].height=24
    paint(ws,r_sub,1,r_tot,LAST,MEDS,inner=HAIR)
    erow(ws,r_sub,1,LAST,bottom=MEDS); erow(ws,r_tot,1,LAST,top=MEDS)
    ecol(ws,r_sub,r_tot,GUT,l=MEDS,rt=MEDS); ecol(ws,r_sub,r_tot,GUT+1,l=MEDS)
    r=r_tot+1

    for m in M: ws.merge_cells(start_row=m[0],start_column=m[1],end_row=m[2],end_column=m[3])
    ws.print_area="A1:J%d"%(r-1)
    ws.page_setup.orientation="landscape"; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_margins.left=ws.page_margins.right=0.4
    ws.page_margins.top=ws.page_margins.bottom=0.5
    ws.print_title_rows="1:%d"%HDR
    ws.oddHeader.left.text="%s -- Scope of Work"%d["project"]; ws.oddHeader.right.text="Page &P of &N"
    ws.oddFooter.center.text=d.get("prepared_by",""); ws.oddFooter.right.text=d.get("date","")
    ws.sheet_properties.outlinePr.summaryRight=False; ws.sheet_view.showGridLines=False

    # ---- SOV: trade roll-up, our total vs sub bids, low bid highlighted ----
    from openpyxl.formatting.rule import FormulaRule
    sv=wb.create_sheet("SOV")
    SW=[30,24,15,15,15,15,15,15,16,16]
    for col,w in zip("ABCDEFGHIJ",SW): sv.column_dimensions[col].width=w
    sv.cell(1,1,"%s -- Schedule of Values"%d["project"]).font=Font(FONT,16,bold=True)
    sv.row_dimensions[1].height=24; sv.merge_cells("A1:J1")
    x=sv.cell(2,1,"Job %s  |  %s  |  %s  |  %s"%(d.get("job",""),d.get("address",""),
                                                 d.get("prepared_by",""),d.get("date","")))
    x.font=Font(FONT,10,italic=True)
    for c in range(1,11): sv.cell(2,c).fill=PatternFill("solid",fgColor=VLT)
    sv.row_dimensions[2].height=18; sv.merge_cells("A2:J2")
    x=sv.cell(3,1,"Our Total and the sub columns are linked to the Scope of Work sheet. "
                  "Markup percentages are entered once on Scope of Work and mirrored here. "
                  "The low bid in each row is highlighted.")
    x.font=Font(FONT,9,italic=True,color=CHAR); sv.merge_cells("A3:J3")
    for c in range(1,11): sv.cell(3,c).fill=PatternFill("solid",fgColor=VLT)
    sv.row_dimensions[3].height=16
    erow(sv,2,1,10,top=MEDS); erow(sv,3,1,10,bottom=MEDS)
    shr=5
    for col,t in enumerate(["Trade","CSI","Our Total","Sub #1","Sub #2","Sub #3","Sub #4",
                            "Low Bid","Low Bidder","Low Bid vs Ours"],1):
        c=sv.cell(shr,col,t); c.font=Font(FONT,11,bold=True,color=W)
        c.fill=PatternFill("solid",fgColor=CHAR)
        c.alignment=Alignment("left" if col<=2 else "center","center")
    sv.row_dimensions[shr].height=22
    paint(sv,shr,1,shr,10,BLKS,inner=Side("thin",color=CHAR)); sv.freeze_panes="A%d"%(shr+1)
    rr=shr+1; first_sov=rr
    for name,csi,tr_row in trade_rows:
        sv.cell(rr,1,name).font=Font(FONT,10,bold=True)
        sv.cell(rr,1).alignment=Alignment("left","center",wrap_text=True)
        sv.cell(rr,2,csi).font=Font(FONT,9,color=CHAR)
        sv.cell(rr,2).alignment=Alignment("left","center",wrap_text=True)
        c=sv.cell(rr,3,"='Scope of Work'!E%d"%tr_row); c.number_format=CUR; c.font=Font(FONT,10,bold=True)
        for i,col in enumerate(range(GUT+1,LAST+1)):
            L=get_column_letter(col)
            c=sv.cell(rr,4+i,"=IF('Scope of Work'!%s%d=\"-\",\"\",'Scope of Work'!%s%d)"%(L,tr_row,L,tr_row))
            c.number_format=CUR
        sv.cell(rr,8,'=IF(COUNT(D%d:G%d)=0,"",MIN(D%d:G%d))'%(rr,rr,rr,rr)).number_format=CUR
        sv.cell(rr,8).font=Font(FONT,10,bold=True)
        sv.cell(rr,9,'=IF(H%d="","",INDEX($D$%d:$G$%d,MATCH(H%d,D%d:G%d,0)))'%(rr,shr,shr,rr,rr,rr))
        sv.cell(rr,10,'=IF(OR(H%d="",N(C%d)=0),"",H%d-C%d)'%(rr,rr,rr,rr)).number_format=CUR
        for col in range(3,11):
            sv.cell(rr,col).alignment=Alignment("center" if col==9 else "right","center")
        sv.row_dimensions[rr].height=ch([(name,30),(csi,24)],mn=20,lp=13); rr+=1
    last_sov=rr-1
    sv.conditional_formatting.add("D%d:G%d"%(first_sov,last_sov),
        FormulaRule(formula=['AND(ISNUMBER(D%d),COUNT($D%d:$G%d)>0,D%d=MIN($D%d:$G%d))'
                             %(first_sov,first_sov,first_sov,first_sov,first_sov,first_sov)],
                    fill=PatternFill("solid",fgColor="FFD6F0D6"), font=Font(bold=True)))
    s_sub=rr
    c=sv.cell(s_sub,1,"SUBTOTAL"); c.font=Font(FONT,11,bold=True); sv.merge_cells(start_row=s_sub,start_column=1,end_row=s_sub,end_column=2)
    for col in range(3,9):
        L=get_column_letter(col)
        cc=sv.cell(s_sub,col,'=IF(SUM(%s%d:%s%d)=0,"",SUM(%s%d:%s%d))'%(L,first_sov,L,last_sov,L,first_sov,L,last_sov))
        cc.number_format=CUR; cc.font=Font(FONT,11,bold=True); cc.alignment=Alignment("right","center")
    for col in range(1,11): sv.cell(s_sub,col).fill=PatternFill("solid",fgColor=LT)
    sv.row_dimensions[s_sub].height=20
    s_m1=s_sub+1
    for i in range(n_mk):
        rw=s_m1+i
        sv.cell(rw,1,"=IF('Scope of Work'!A%d=\"\",\"\",'Scope of Work'!A%d)"%(r_m1+i,r_m1+i)).font=Font(FONT,10)
        c=sv.cell(rw,2,"=IF('Scope of Work'!D%d=\"\",\"\",'Scope of Work'!D%d)"%(r_m1+i,r_m1+i))
        c.number_format='0.00%'; c.alignment=Alignment("right","center"); c.font=Font(FONT,10,bold=True)
        for col in range(3,9):
            L=get_column_letter(col)
            base="%s%d"%(L,s_sub) if i==0 else "%s%d+SUM(%s%d:%s%d)"%(L,s_sub,L,s_m1,L,rw-1)
            cc=sv.cell(rw,col,'=IF(OR($B%d="",N(%s%d)=0),"",ROUND((%s)*$B%d,2))'%(rw,L,s_sub,base,rw))
            cc.number_format=CUR; cc.alignment=Alignment("right","center")
        for col in range(1,11): sv.cell(rw,col).fill=PatternFill("solid",fgColor=VLT)
        sv.row_dimensions[rw].height=18
    s_tot=s_m1+n_mk
    c=sv.cell(s_tot,1,"TOTAL"); c.font=Font(FONT,12,bold=True); sv.merge_cells(start_row=s_tot,start_column=1,end_row=s_tot,end_column=2)
    for col in range(3,9):
        L=get_column_letter(col)
        cc=sv.cell(s_tot,col,'=IF(N(%s%d)=0,"",ROUND(%s%d+SUM(%s%d:%s%d),2))'%(L,s_sub,L,s_sub,L,s_m1,L,s_m1+n_mk-1))
        cc.number_format=CUR; cc.font=Font(FONT,12,bold=True); cc.alignment=Alignment("right","center")
    for col in range(1,11): sv.cell(s_tot,col).fill=PatternFill("solid",fgColor=LT)
    sv.row_dimensions[s_tot].height=24
    paint(sv,first_sov,1,s_tot,10,MEDS,inner=HAIR)
    erow(sv,last_sov,1,10,bottom=MEDS); erow(sv,s_tot,1,10,top=MEDS)
    sv.page_setup.orientation="landscape"; sv.page_setup.fitToWidth=1; sv.page_setup.fitToHeight=0
    sv.sheet_properties.pageSetUpPr.fitToPage=True
    sv.print_title_rows="1:%d"%shr; sv.sheet_view.showGridLines=False

    # ---- Notes & Clarifications (client-facing) ----
    n=wb.create_sheet("Notes & Clarifications")
    for col,w in zip("ABC",[16,34,104]): n.column_dimensions[col].width=w
    n.cell(1,1,"%s -- Notes & Clarifications"%d["project"]).font=Font(FONT,16,bold=True)
    n.row_dimensions[1].height=24; n.merge_cells("A1:C1")
    x=n.cell(2,1,"Job %s  |  %s  |  %s  |  %s"%(d.get("job",""),d.get("address",""),
                                                d.get("prepared_by",""),d.get("date","")))
    x.font=Font(FONT,10,italic=True)
    for c in range(1,4): n.cell(2,c).fill=PatternFill("solid",fgColor=VLT)
    n.row_dimensions[2].height=18; n.merge_cells("A2:C2")
    erow(n,2,1,3,top=MEDS,bottom=MEDS)
    hr=4
    for col,t in enumerate(["Type","Item","Clarification"],1):
        c=n.cell(hr,col,t); c.font=Font(FONT,11,bold=True,color=W)
        c.fill=PatternFill("solid",fgColor=CHAR); c.alignment=Alignment("left","center")
    n.row_dimensions[hr].height=20
    paint(n,hr,1,hr,3,BLKS,inner=Side("thin",color=CHAR)); n.freeze_panes="A5"
    rr=hr+1
    for nt in d.get("notes",[]):
        for col,v in enumerate([nt["type"],nt["item"],nt["text"]],1):
            c=n.cell(rr,col,v); c.font=Font(FONT,10,bold=(col==1))
            c.alignment=Alignment("left","top",wrap_text=True)
            c.fill=PatternFill("solid",fgColor=(RED if nt["type"]=="EXCLUSION" else VLT))
        n.row_dimensions[rr].height=ch([(nt["item"],34),(nt["text"],104)],mn=26,lp=13); rr+=1
    if rr>hr+1: paint(n,hr+1,1,rr-1,3,MEDS,inner=HAIR)
    n.page_setup.orientation="landscape"; n.page_setup.fitToWidth=1; n.page_setup.fitToHeight=0
    n.sheet_properties.pageSetUpPr.fitToPage=True
    n.print_title_rows="1:%d"%hr; n.sheet_view.showGridLines=False

    # ---- Quantity Basis (internal, hidden -- see SKILL.md, hidden is not removed) ----
    basis=d.get("basis") or []
    if basis:
        h=wb.create_sheet("Quantity Basis")
        for col,w in zip("ABCDEF",[22,10,54,12,8,70]): h.column_dimensions[col].width=w
        h.cell(1,1,"Internal -- quantity basis by line").font=Font(FONT,14,bold=True)
        h.merge_cells("A1:F1"); hr2=3
        for col,t in enumerate(["Trade","CSI","Scope Item","Qty","Unit","Basis"],1):
            c=h.cell(hr2,col,t); c.font=Font(FONT,11,bold=True,color=W)
            c.fill=PatternFill("solid",fgColor=CHAR)
        h.row_dimensions[hr2].height=20; h.freeze_panes="A4"; rr=hr2+1
        for b in basis:
            q=b.get("qty")
            for col,v in enumerate([b.get("trade",""),b.get("csi",""),b.get("scope",""),
                                    (q if q is not None else "--"),b.get("unit",""),
                                    b.get("basis","")],1):
                c=h.cell(rr,col,v); c.font=Font(FONT,10)
                c.alignment=Alignment("left","top",wrap_text=True)
                c.fill=PatternFill("solid",fgColor=(RED if q is None else VLT))
            h.row_dimensions[rr].height=ch([(b.get("scope",""),54),(b.get("basis",""),70)],mn=24,lp=13)
            rr+=1
        paint(h,hr2,1,rr-1,6,MEDS,inner=HAIR)
        h.sheet_view.showGridLines=False; h.sheet_state="hidden"

    wb.save(out)
    nl=sum(len(t["lines"]) for t in d["trades"])
    nq=sum(1 for t in d["trades"] for l in t["lines"] if l.get("qty") is None)
    print("wrote %s | trades %d | lines %d | no-quantity %d | notes %d"
          % (out,len(d["trades"]),nl,nq,len(d.get("notes",[]))))

if __name__=="__main__":
    build(json.load(open(sys.argv[1])), sys.argv[2])
