"""
SOW Excel builder for the sow-generator skill.

Reads a sow_data.json (schema below) and produces a polished, professional
Scope of Work workbook with subcontractor leveling columns, per SKILL.md.

sow_data.json schema:
{
  "project": {"name": str, "address": str, "date_prepared": str, "prepared_by": str},
  "trades": [
    {
      "code": "09 30 13",
      "title": "Ceramic Tiling",
      "sub_names": ["Sub #1","Sub #2","Sub #3","Sub #4"],   # optional, defaults shown
      "items": [{"reference": str, "scope": str}, ...]
    }, ...
  ],
  "notes": {"standard": [str, ...], "drawing_specific": [str, ...]},
  "conflicts": [str, ...],
  "tbd": [str, ...]
}

Env vars:
  SOW_DATA_JSON   - path to input JSON (default: sow_data.json)
  SOW_OUTPUT_XLSX - path to output .xlsx (default: derived from project name + date)
"""
import json
import math
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

WHITE = "FFFFFF"
BLACK = "000000"
CHARCOAL = "3A3A3A"
MED_GRAY = "7A7A7A"
LIGHT_GRAY = "E8E8E8"
VLIGHT_GRAY = "F5F5F5"
FONT_NAME = "Calibri"
CURRENCY_FMT = '$#,##0.00;-$#,##0.00;"–"'


def calc_height(specs, min_h=30, line_pt=15):
    lines = max([1] + [math.ceil(len(str(t)) / cpl) for t, cpl in specs if t])
    return max(min_h, lines * line_pt)


def _fill(color):
    return PatternFill("solid", fgColor=color)


def add_section(ws, row, title, lines):
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    c = ws.cell(row=row, column=1, value=title)
    c.font = Font(name=FONT_NAME, size=14, bold=True, color=WHITE)
    for col in range(1, 7):
        cell = ws.cell(row=row, column=col)
        cell.fill = _fill(CHARCOAL)
        cell.border = Border(bottom=Side(style="thick", color=BLACK))
    ws.row_dimensions[row].height = 22
    row += 1
    if not lines:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        c = ws.cell(row=row, column=1, value="None")
        c.font = Font(name=FONT_NAME, size=10, italic=True, color=MED_GRAY)
        c.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[row].height = 30
        row += 1
        return row
    for line in lines:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        c = ws.cell(row=row, column=1, value=line)
        c.font = Font(name=FONT_NAME, size=10, color=BLACK)
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = calc_height([(line, 130)])
        row += 1
    return row


def build_workbook(data, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Scope of Work"

    project = data["project"]

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 60
    for col in ("C", "D", "E", "F"):
        ws.column_dimensions[col].width = 16
        ws.column_dimensions[col].outlineLevel = 1

    # Row 1: Title
    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = project["name"]
    c.font = Font(name=FONT_NAME, size=18, bold=True, color=BLACK)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30

    # Row 2: Address
    ws.merge_cells("A2:F2")
    c = ws["A2"]
    c.value = project.get("address", "")
    c.font = Font(name=FONT_NAME, size=11, color=BLACK)
    c.alignment = Alignment(horizontal="left", vertical="center")

    # Row 3: Date Prepared | Prepared By
    c = ws["A3"]
    c.value = f"Date Prepared: {project.get('date_prepared', '')}"
    c.font = Font(name=FONT_NAME, size=10, color=BLACK)
    c = ws["B3"]
    c.value = f"Prepared By: {project.get('prepared_by', '')}"
    c.font = Font(name=FONT_NAME, size=10, color=BLACK)

    # Row 4: spacer with bottom border
    for col in range(1, 7):
        ws.cell(row=4, column=col).border = Border(bottom=Side(style="thin", color=BLACK))

    # Row 5: Directive legend
    ws.merge_cells("A5:F5")
    c = ws["A5"]
    c.value = ("F/I = Furnish and Install  |  F/O = Furnish Only  |  I/O = Install Only  |  "
                "Remove = Demolition  |  Provide = Services")
    c.font = Font(name=FONT_NAME, size=9, italic=True, color=MED_GRAY)
    c.fill = _fill(LIGHT_GRAY)
    c.alignment = Alignment(horizontal="left", vertical="center")

    # Row 7: column headers
    headers = ["Reference", "Scope Item", "Sub #1", "Sub #2", "Sub #3", "Sub #4"]
    for idx, h in enumerate(headers, start=1):
        c = ws.cell(row=7, column=idx, value=h)
        c.font = Font(name=FONT_NAME, size=11, bold=True, color=WHITE)
        c.fill = _fill(CHARCOAL)
        c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[7].height = 22

    ws.freeze_panes = "A8"

    scope_start_row = 7
    row = 8

    for trade in data["trades"]:
        header_row1 = row
        header_row2 = row + 1
        ws.merge_cells(start_row=header_row1, start_column=1, end_row=header_row2, end_column=2)
        hcell = ws.cell(row=header_row1, column=1)
        hcell.value = f"{trade['code']} — {trade['title']}"
        hcell.font = Font(name=FONT_NAME, size=11, bold=True, color=WHITE)
        hcell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        for r in (header_row1, header_row2):
            for col in (1, 2):
                ws.cell(row=r, column=col).fill = _fill(MED_GRAY)

        sub_names = trade.get("sub_names") or ["Sub #1", "Sub #2", "Sub #3", "Sub #4"]
        for i, name in enumerate(sub_names[:4]):
            col = 3 + i
            c = ws.cell(row=header_row1, column=col, value=name)
            c.font = Font(name=FONT_NAME, size=10, bold=True, color=BLACK)
            c.fill = _fill(LIGHT_GRAY)
            c.alignment = Alignment(horizontal="center", vertical="center")

        lump_sum_row = header_row2
        for i in range(4):
            col = 3 + i
            c = ws.cell(row=lump_sum_row, column=col)
            c.number_format = CURRENCY_FMT
            c.font = Font(name=FONT_NAME, size=10, color=BLACK)
            c.fill = _fill(VLIGHT_GRAY)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = Border(bottom=Side(style="medium", color=BLACK))

        ws.row_dimensions[header_row1].height = 20
        ws.row_dimensions[header_row2].height = 20

        row = header_row2 + 1
        for idx, item in enumerate(trade["items"]):
            reference = item.get("reference", "")
            scope = item.get("scope", "")
            ref_cell = ws.cell(row=row, column=1, value=reference)
            scope_cell = ws.cell(row=row, column=2, value=scope)
            ref_cell.font = Font(name=FONT_NAME, size=10, color=CHARCOAL)
            ref_cell.fill = _fill(LIGHT_GRAY)
            ref_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            scope_cell.font = Font(name=FONT_NAME, size=10, color=BLACK)
            scope_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

            shade = (idx % 2 == 0)
            if shade:
                scope_cell.fill = _fill(VLIGHT_GRAY)
                for col in range(3, 7):
                    ws.cell(row=row, column=col).fill = _fill(VLIGHT_GRAY)

            for col in range(3, 7):
                c = ws.cell(row=row, column=col)
                c.number_format = CURRENCY_FMT
                c.font = Font(name=FONT_NAME, size=10, color=BLACK)
                c.alignment = Alignment(horizontal="right", vertical="center")
                if shade:
                    c.fill = _fill(VLIGHT_GRAY)

            for col in range(1, 7):
                cell = ws.cell(row=row, column=col)
                existing = cell.border
                cell.border = Border(
                    bottom=Side(style="thin", color=LIGHT_GRAY),
                    top=existing.top, left=existing.left, right=existing.right,
                )

            ws.row_dimensions[row].height = calc_height([(reference, 17), (scope, 60)])
            row += 1

        last_item_row = row - 1
        total_row = row
        ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=2)
        tcell = ws.cell(row=total_row, column=1, value="TRADE TOTAL")
        tcell.font = Font(name=FONT_NAME, size=10, bold=True, color=BLACK)
        tcell.alignment = Alignment(horizontal="left", vertical="center")
        for i in range(4):
            col = 3 + i
            col_letter = get_column_letter(col)
            c = ws.cell(row=total_row, column=col)
            c.value = f"=SUM({col_letter}{lump_sum_row}:{col_letter}{last_item_row})"
            c.font = Font(name=FONT_NAME, size=10, bold=True, color=BLACK)
            c.number_format = CURRENCY_FMT
            c.alignment = Alignment(horizontal="right", vertical="center")
        ws.row_dimensions[total_row].height = 20
        for col in range(1, 7):
            ws.cell(row=total_row, column=col).border = Border(bottom=Side(style="medium", color=CHARCOAL))

        row = total_row + 2  # blank spacer row before next trade

    scope_end_row = row - 2  # last Trade Total row

    # Outer medium-black border around the scope table (header row through last total row)
    for col in range(1, 7):
        top = ws.cell(row=scope_start_row, column=col)
        top.border = Border(top=Side(style="medium", color=BLACK), bottom=top.border.bottom,
                             left=top.border.left, right=top.border.right)
        bot = ws.cell(row=scope_end_row, column=col)
        bot.border = Border(bottom=Side(style="medium", color=BLACK), top=bot.border.top,
                             left=bot.border.left, right=bot.border.right)
    for r in range(scope_start_row, scope_end_row + 1):
        left = ws.cell(row=r, column=1)
        left.border = Border(left=Side(style="medium", color=BLACK), top=left.border.top,
                              bottom=left.border.bottom, right=left.border.right)
        right = ws.cell(row=r, column=6)
        right.border = Border(right=Side(style="medium", color=BLACK), top=right.border.top,
                               bottom=right.border.bottom, left=right.border.left)

    row = scope_end_row + 1

    notes = data.get("notes", {})
    combined_notes = list(notes.get("standard", [])) + list(notes.get("drawing_specific", []))
    row = add_section(ws, row, "Notes and Clarifications", combined_notes)
    row = add_section(ws, row, "Conflicts", data.get("conflicts", []))
    row = add_section(ws, row, "TBD Items", data.get("tbd", []))

    last_row = row - 1

    ws.sheet_properties.outlinePr.summaryRight = True
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.5, bottom=0.5, header=0.3, footer=0.3)
    ws.oddHeader.left.text = project["name"]
    ws.oddHeader.right.text = "Page &P of &N"
    ws.oddFooter.center.text = "Confidential - Montana Contracting"
    ws.oddFooter.right.text = project.get("date_prepared", "")
    ws.print_title_rows = "1:7"
    ws.print_area = f"A1:F{last_row}"

    wb.save(output_path)
    return output_path


def main():
    data_path = os.environ.get("SOW_DATA_JSON", "sow_data.json")
    with open(data_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    project_name = data["project"]["name"].replace(" ", "")
    date_str = data["project"].get("date_prepared", "")
    default_out = f"{project_name}_SOW_{date_str}.xlsx"
    output_path = os.environ.get("SOW_OUTPUT_XLSX", default_out)

    result = build_workbook(data, output_path)
    print(f"Wrote {result}")


if __name__ == "__main__":
    main()
