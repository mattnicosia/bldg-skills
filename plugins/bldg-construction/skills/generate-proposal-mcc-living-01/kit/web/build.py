#!/usr/bin/env python3
"""Website-format presentation build. Reference job: E260050 The Cooper Barn (v5.2, approved by Matt 2026-10-16).

Copy this folder to <project>/web/, then change the job-specific blocks marked JOB below: the element and allowance
row maps, the copy (PROSE, STEPS, BASIS, CLAR, EXCL), the schedule milestones and the zip name.

    python3 -I build.py

Every figure comes from the estimate's line items (Estimate sheet), to the cent, so the Elements and MasterFormat
views, the signature elements and the base budget all tie. Writes
    site/       index.html (Artifact page, no skeleton) + assets/
    netlify/    index.html (full document) + assets/ + robots.txt
    E260050_Presentation_Netlify_v5.zip
Stops if a figure does not tie, or client copy breaks the house rules.
"""
import base64, datetime as dt, glob, html, json, os, re, shutil, sys, zipfile
from urllib.parse import quote
from decimal import Decimal as D, ROUND_HALF_UP
import numpy as np
import openpyxl
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
CB = os.path.dirname(HERE)
KIT = os.environ.get('MCC_KIT') or os.path.join(os.path.dirname(CB), 'kit')
LENIS = 'https://cdn.jsdelivr.net/npm/lenis@1.1.13/dist/lenis.min.js'
E = lambda s: html.escape(str(s), quote=True)
C = lambda x: D(str(x)).quantize(D('0.01'), ROUND_HALF_UP)
money = lambda v: '${:,.2f}'.format(C(v))
pct = lambda v, of: '{:.1f}%'.format(float(D(v) / D(of) * 100))


def die(m):
    print('BUILD STOPPED:', m); sys.exit(1)


P = json.load(open(os.path.join(CB, 'project.json')))
SCH = json.load(open(os.path.join(CB, 'schedule.json')))

# ================================================================ money, from the line items
wb = openpyxl.load_workbook(os.path.join(CB, 'estimate.xlsx'), data_only=True)
ws = wb['Estimate']
LINES, trade = {}, None
for i, r in enumerate(ws.iter_rows(min_row=10, max_row=232, values_only=True), start=10):
    code, item, qty, unit, rate, tot = r[:6]
    if code and isinstance(code, str) and code.isupper() and not item and code not in ('TRADE TOTAL', 'SUBTOTAL'): trade = code; continue
    if code in ('TRADE TOTAL', 'SUBTOTAL') or not item: continue
    if isinstance(tot, (int, float)) and tot: LINES[i] = {'code': str(code).strip(), 'item': str(item).strip(), 'trade': trade, 'v': C(tot)}
DIRECT = sum(l['v'] for l in LINES.values())
OP = C(DIRECT * D('0.15')); BASE = DIRECT + OP
sheet_total = C(next(r[5] for r in ws.iter_rows(min_row=230, max_row=240, values_only=True) if r[0] == 'TOTAL -- BASE BID'))
if BASE != sheet_total: die(f'base {BASE} does not tie to the Estimate sheet total {sheet_total}')
ALT_ROWS = {}
for i, r in enumerate(ws.iter_rows(min_row=400, max_row=432, values_only=True), start=400):
    if r[0] and str(r[0]).startswith('## '): pass
    if r[0] == 'TOTAL' and isinstance(r[5], (int, float)): ALT_ROWS[len(ALT_ROWS) + 1] = C(r[5])

def clean(item):   # client wording for a line item: no bracketed estimator notes
    s = re.sub(r'\s*\[[^\]]*\]', '', item); s = s.replace(' (M+L)', '').replace('(M+L)', '')
    s = re.sub(r'\$(\d[\d,]*)(?![\d,]*\.\d)', r'$\1.00', s)   # every dollar figure carries cents
    s = re.sub(r'\s*\(Per (Week|Month|Page|Riser)\)', '', s)   # the line shows a total, so the unit basis would misread
    s = re.sub(r'\(\$[\d,]+\.\d\d/SF Mat\. Allowance\)', '(Material Allowance)', s)   # no unit rates in the line detail
    s = re.sub(r'\(\$[\d,]+\.\d\d Allowance\)', '(Allowance)', s)
    return re.sub(r'\s{2,}', ' ', s).strip()

trade_names = P.get('trade_names', {})
def tkey(name): return name.upper()
TR = {}
for l in LINES.values(): TR[l['trade']] = TR.get(l['trade'], D(0)) + l['v']
ROW_GROUP = {61: 'Building Envelope', 62: 'Building Envelope'}   # the cupola and vane are exterior work, carried under finish carpentry
if not ('Cupola' in LINES[61]['item'] and 'Weathervane' in LINES[62]['item']): die('cupola rows moved')
SCOPE_FIX = {'Building Envelope': 'Insulation, standing seam roof, siding, trim, the cupola, Marvin windows and bifold doors.',
             'Interior Finishes': 'Drywall, stairs, wainscot, reclaimed beams, the copper top bar, millwork, flooring, tile and paint.'}
def group_of(r, l):
    if r in ROW_GROUP: return ROW_GROUP[r]
    for g in P['groups']:
        if l['trade'] in [tkey(t) for t in g['trades']]: return g['name']
    die('line has no group: ' + l['item'])
GROUPS = []
for g in P['groups']:
    rows = [(r, l) for r, l in LINES.items() if group_of(r, l) == g['name']]
    GROUPS.append({'name': g['name'], 'scope': SCOPE_FIX.get(g['name'], g['scope']), 'v': sum(l['v'] for r, l in rows),
                   'lines': [{'item': clean(l['item']), 'v': l['v']} for r, l in rows]})
if sum(g['v'] for g in GROUPS) != DIRECT: die('groups do not tie to direct cost')

MF_TITLES = {'01': 'General Requirements', '02': 'Existing Conditions', '03': 'Concrete', '04': 'Masonry', '05': 'Metals',
             '06': 'Wood, Plastics, and Composites', '07': 'Thermal and Moisture Protection', '08': 'Openings', '09': 'Finishes',
             '10': 'Specialties', '12': 'Furnishings', '22': 'Plumbing', '23': 'Heating, Ventilating, and Air Conditioning',
             '26': 'Electrical', '27': 'Communications', '28': 'Electronic Safety and Security', '31': 'Earthwork', '33': 'Utilities'}
MF = {}
for l in LINES.values():
    dv = l['code'][:2]
    if dv not in MF_TITLES: die('no MasterFormat title for division ' + dv)
    MF.setdefault(dv, []).append({'code': l['code'], 'item': clean(l['item']), 'v': l['v']})
DIVS = [{'n': k, 'name': MF_TITLES[k], 'v': sum(x['v'] for x in MF[k]), 'lines': MF[k]} for k in sorted(MF)]
if sum(d['v'] for d in DIVS) != DIRECT: die('divisions do not tie to direct cost')

def el(rows):
    return C(sum(LINES[r]['v'] for r in rows) * D('1.15'))
# JOB: the signature elements, by Estimate row, each x 1.15 with O&P. Row words are checked so a moved row stops the build.
ELS = [
    ('bifolds', 'Bifolds and Transoms', el([125, 126, 45]), 'Four Marvin custom bifold units with transoms, 10 feet wide and 11 feet 8 inches high.'),
    ('deck', 'Roof Deck', el([39, 51, 79, 80, 107, 108, 109, 110, 128]), 'Cumaru over pitched EPDM, with the guardrail, the coping and the door from the mezzanine.'),
    ('bar', 'Whiskey Bar', el([86, 87, 88, 89]), 'Front bar, back bar, copper bar top and stone back counter.'),
    ('canopy', 'Canopies', el([54, 55]), 'Three canopies over the bifold doors, carried on brackets.'),
    ('cupola', 'Cupola and Vane', el([61, 62]), 'A glazed cupola on the ridge with a horse weathervane.'),
    ('stove', 'Stove and Flue', el([162]), 'Hearthstone Bari wood stove with a Class A flue and hearth pad.'),
]
for k, n, v, s in ELS:   # rows must be the lines we think they are
    pass
for r, word in [(125, 'Bifold'), (126, 'Bifold'), (45, 'LVL'), (39, 'Guardrail'), (51, 'Lounge Wing Roof'), (79, 'Cumaru'), (80, 'Drainage Mat'), (107, 'EPDM'),
                (108, 'Parapet'), (109, 'Coping'), (110, 'Scuppers'), (128, 'Roof Deck'), (86, 'Front Bar'), (87, 'Back Bar'), (88, 'Copper Bar'), (89, 'Back Bar Countertop'),
                (54, 'Canopies'), (55, 'Brackets'), (61, 'Cupola'), (62, 'Weathervane'), (162, 'Hearthstone')]:
    if word.lower() not in LINES[r]['item'].lower(): die(f'row {r} is not {word}: {LINES[r]["item"]}')
EL_TOTAL = sum(e[2] for e in ELS)
REST = BASE - EL_TOTAL
SITE_FDN = C((TR['SITEWORK & SITE UTILITIES'] + TR['CONCRETE & FOUNDATIONS']) * D('1.15'))

# JOB: allowances from the Notes sheet, by Estimate row.
ALLOW_ROWS = [('Marvin Bifold Doors', 'Four units with transoms, supply, $26,000.00 each', [125]), ('Roof Deck Guardrail', 'Type to be selected', [39]),
              ('House Electric Service Upgrade', 'To be set with the construction drawings', [211]), ('Lighting Fixtures', 'Fixtures throughout', [212]),
              ('Propane Tank and Piping', 'New underground tank', [182]), ('Back Bar Stone Countertop', '31 SF', [89]),
              ('Kohler Fixtures', 'Toilet, lav sink and bar sink with faucets', [179, 180, 181]), ('Powder Room Vanity', 'Vanity with top', [90]),
              ('Low Voltage', 'TV and data', [214])]
ALLOWS = [{'label': a, 'basis': b, 'v': money(sum(LINES[r]['v'] for r in rows))} for a, b, rows in ALLOW_ROWS]
ALLOW_TOTAL = sum(sum(LINES[r]['v'] for r in rows) for a, b, rows in ALLOW_ROWS)
ALLOWS += [{'label': 'Wood Flooring', 'basis': 'Material', 'v': '$13.00 per SF'}, {'label': 'Powder Room Tile', 'basis': 'Material', 'v': '$10.00 per SF'}]
if LINES[125]['v'] != D('104000.00') or 'Kohler' not in LINES[179]['item']: die('allowance rows moved')

ALTS = [{'id': str(k), 'title': P['alternates'][str(k)]['title'], 'note': P['alternates'][str(k)]['note'], 'v': ALT_ROWS[k]} for k in (1, 2, 3)]

# ================================================================ selections (kit v1.7): one card per element, prices from the workbook. Optional: no selections.json, no section.
sys.path.insert(0, os.path.join(KIT, 'tools', 'sel')); import selbuild
SEL_PATH = os.path.join(CB, 'selections.json')
SELX = {'C': C, 'die': die, 'money': money, 'E': E, 'ws': ws, 'LINES': LINES, 'ALT_ROWS': ALT_ROWS, 'BASE': BASE, 'CB': CB, 'OP': D('0.15')}
SEL = selbuild.resolve(json.load(open(SEL_PATH)), SELX) if os.path.exists(SEL_PATH) else None

# ================================================================ schedule, in weeks from mobilization, by phase
d = lambda s: dt.date.fromisoformat(s)
mob = d(SCH['mob'])
PHASES = [('Site and Foundation', ['mob', 'fdn'], '#001a8a'), ('Structure', ['frm', 'roof'], '#0028cc'), ('Envelope', ['win', 'rfg', 'sid', 'foam'], '#3d5bd9'),
          ('Systems', ['mep'], '#6a82e2'), ('Interiors', ['gwb', 'mill', 'pnt', 'flr'], '#93a5ec'), ('Closeout', ['trim', 'pun'], '#b9c5f3')]
phase_of = {k: (n, c) for n, ks, c in PHASES for k in ks}
SROWS = []
for r in SCH['rows']:
    n, c = phase_of[r['key']]
    SROWS.append({'name': r['name'].replace(', Gutters', ' and Gutters'), 'a': (d(r['s']) - mob).days / 7, 'b': ((d(r['f']) - mob).days + 1) / 7, 'phase': n, 'c': c})
SPAN = float(SCH['weeks'])   # the chart ends on the Work Complete line
for r in SROWS: r['b'] = min(r['b'], SPAN)   # the last finish date counts its own day, a fraction past week 26
PH_SPANS = [(n, c, min(r['a'] for r in SROWS if r['phase'] == n), max(r['b'] for r in SROWS if r['phase'] == n)) for n, ks, c in PHASES]
# JOB: schedule milestones from the working schedule.
MS = [(0, 'Mobilize'), (round((d('2027-05-14') - mob).days / 7), 'Foundation Complete'), (round((d('2027-07-06') - mob).days / 7), 'Dried In'), (round((d(SCH['keys']) - mob).days / 7), 'Work Complete')]
if MS[-1][0] != SCH['weeks']: die('schedule does not end at 26 weeks')

# ================================================================ model frames: one camera, cropped once for all
MD = os.path.join(HERE, 'model')
FR_C = ['c0_site', 'c0b_backfill', 'c1_workshop', 'c2_lounge', 'c3_exterior']
FR_W = ['w1_bifolds', 'w2_deck', 'w3_bar', 'w4_canopy', 'w5_cupola', 'w6_stove', 'w7_all']
FRAMES = FR_C + FR_W
box = None
for f in FRAMES:
    a = np.array(Image.open(os.path.join(MD, f + '.png')).split()[3])
    ys, xs = np.where(a > 200); b = (xs.min(), ys.min(), xs.max(), ys.max())
    box = b if box is None else (min(box[0], b[0]), min(box[1], b[1]), max(box[2], b[2]), max(box[3], b[3]))
CROP = (max(0, int(box[0]) - 40), max(0, int(box[1]) - 40), min(1600, int(box[2]) + 90), min(1200, int(box[3]) + 50))
FR = (CROP[2] - CROP[0], CROP[3] - CROP[1])
PTS = {k: [x - CROP[0], y - CROP[1]] for k, (x, y) in json.load(open(os.path.join(MD, 'c3_exterior.png.points.json'))).items()}
BB = {}
for (k, *_r), f in zip(ELS, FR_W):   # where each element shows in blue: the box the focus brackets draw around
    im = np.array(Image.open(os.path.join(MD, f + '.png')).convert('RGBA')).astype(int)
    blue = (im[..., 2] > im[..., 0] + 70) & (im[..., 2] > im[..., 1] + 50) & (im[..., 3] > 200)
    ys, xs = np.where(blue)
    BB[k] = [int(xs.min()) - CROP[0], int(ys.min()) - CROP[1], int(xs.max()) - CROP[0], int(ys.max()) - CROP[1]]

# ================================================================ copy
LOC = 'North Salem, NY'
# JOB: the project description, in the Project description voice (see the skill). Three paragraphs, 90 to 220 words.
PROSE = [
    'The Cooper Barn is a workshop and lounge barn at 194 Titicus Road in North Salem, New York, designed by Crisp Architects. A 40 by 30 foot workshop barn and a 16 by 26 foot lounge wing hold the owner\'s shop and a whiskey bar lounge.',
    'The barn is framed in 2x6 at a 15 foot plate under a 9:12 galvalume standing seam roof, clad in Hardie board and batten over ZIP sheathing and a rain screen, with spray foam insulation. Four Marvin custom bifold doors with transoms, each 10 feet wide and 11 feet 8 inches high, open the workshop on LVL headers. The lounge wing is built over a full-height mechanical space and carries a Cumaru roof deck over pitched EPDM.',
    'Inside, exposed beams cross the double-height workshop, with workbenches, desk areas, a 3D printer desk and a white oak stair to a mezzanine. The lounge has wainscot walls, reclaimed beams, wide plank flooring, a copper top whiskey bar and a Hearthstone Bari wood stove. Two ducted heat pump zones serve the workshop, the lounge and the mezzanine.',
]
TITLE_EL = 'Signature Elements'
# JOB: the design steps (frames from render_model.sh).
STEPS = [
    {'k': 'site', 'rail': 0, 'title': 'Site and Foundation', 'fig': money(SITE_FDN), 'cap': 'Site work and foundations, with overhead and profit.',
     'body': 'Excavation for 4 foot frost walls and a full-height mechanical space under the lounge wing. 12x24 footings, parged 10 inch walls and 4 inch slabs on 6 inch gravel. Electric feeder, water service and sanitary line each carried at 150 feet.',
     'img': 0, 'via': 1, 'note': 'Excavation, Footings and Frost Walls'},
    {'k': 'workshop', 'rail': 1, 'title': 'Workshop', 'body': 'A double-height workshop under a standing seam roof, with exposed beams, a cupola and a mezzanine. Workbenches, desk areas and a 3D printer desk are part of the budget.', 'img': 2, 'note': 'Workshop Barn'},
    {'k': 'lounge', 'rail': 2, 'title': 'Lounge', 'body': P['stops'][3]['body'][0], 'img': 3, 'note': 'Lounge Wing, Walls Cut Away'},
    {'k': 'exterior', 'rail': 3, 'title': 'Exterior', 'body': P['stops'][4]['body'][0], 'img': 4, 'note': 'Complete Building'},
]
for i, (k, name, v, spec) in enumerate(ELS):
    STEPS.append({'k': k, 'rail': 4, 'title': TITLE_EL, 'sub': f'<b>{i + 1:02d} / {len(ELS):02d}</b>{E(name)}', 'fig': money(v), 'cap': spec,
                  'img': 5 + i, 'el': i, 'note': 'Shown in Montana Blue'})
STEPS.append({'k': 'all', 'rail': 4, 'title': TITLE_EL, 'sub': 'Six Elements', 'fig': money(EL_TOTAL), 'cap': f'{pct(EL_TOTAL, BASE)} of the {money(BASE)} base budget.',
              'img': 11, 'all': True, 'note': 'Signature Elements in Montana Blue'})
for s in STEPS: s['label'] = '02 / Design'

# JOB: Basis of Budget, Clarifications and Exclusions, from the Notes sheet. Fragments are checked against it below.
BASIS = ['Crisp Architects SK-5 First and Second Floor Plans dated September 22, 2026 (Option 1) and the Cooper Barn Budget email of October 9, 2026.',
         'Quantities measured from SK-5 at 1/4" = 1\'-0".',
         'Structural, mechanical and site drawings are not yet issued. Final framing and systems follow those drawings.']
CLAR = ['Four Marvin custom exterior bifold doors with transoms, carried at a $26,000.00 supply allowance each plus installation, at 10\'-0" x 11\'-8". Each transom is a separate unit above its bifold.',
        'Wood framed 2x6 walls at a 15\' plate, a 9:12 roof, exposed beams at the workshop (100 LF) and LVL headers at the bifold openings. No structural steel is carried.',
        'Workshop on 4\' frost walls. Mechanical space under the full lounge wing footprint (16\' x 26\') with 9\' walls. 12x24 footings, parged 10" walls, 4" slabs on 6" gravel and 15 mil Stego.',
        'Electric feeder, water service and sanitary line each carried at 150 LF. House electric service upgrade carried as an $8,925.00 allowance. New propane tank and piping carried as a $6,825.00 allowance.',
        'Hearthstone Bari wood stove with Class A flue and hearth pad included at $11,550.00.',
        'Allowance amounts are carried before overhead and profit. Overhead and profit of 15% applies to the base budget and the alternates.',
        '2" XPS below grade is priced as Add Alternate 1. None of the three alternates is in the base budget.']
EXCL = ['Workshop equipment is furnished and set by the owner. Dedicated 240V circuits are included.',
        'Building permit, health department and utility fees by owner.',
        'Structural, MEP and civil engineering and site survey by others. Building stakeout and foundation as-built survey is Add Alternate 2.',
        'Upgrades to the existing septic system, if required by the county, are excluded.',
        'Landscaping, topsoil, seed, driveway and paving are excluded. Rough and finish grading around the barn is included.',
        'Rock removal, dewatering and unsuitable soil replacement are excluded.',
        'Dust collection ductwork and the PM2200 collector are excluded.']
notes_src = ' '.join(n for n in [str(r[2]) for r in wb['Notes & Clarifications'].iter_rows(min_row=5, values_only=True) if r[2]])
for frag in ['240V', 'utility fees', 'Add Alternate 2', 'septic', 'driveway and paving', 'Rock removal', 'PM2200', '26,000', '15\' plate', '100 LF', 'LVL headers',
             'No structural steel', '4\' frost walls', '16\' x 26\'', '15 mil Stego', '150 LF', '8,925', '6,825', 'Hearthstone Bari', '11,550', 'XPS']:
    if frag not in notes_src: die('note drifted from the estimate: ' + frag)

# JOB: principals and the Approach copy. Confirm titles.
TEAM = [('Joe Montana', 'Founder and CEO', 'joe-montana-sr-800.webp'), ('Matt Nicosia', 'COO', 'matt-nicosia-800.webp'), ('Chris Lange', 'Director of Construction', 'chris-lange-800.webp')]
# kit v1.10: the team as numbers (1, 1, 3), three short lines each. JOB: the architect and the long-lead item named in the lines.
TIERS = [
    ('1', 'Site Superintendent', 'On site, start to finish', ['Runs the site from mobilization to the punch list.',
        'Checks the work against the drawings before it is covered.', 'Files daily field reports and photos in Montana OS.']),
    ('1', 'Project Manager', 'On every project', ['Runs the schedule, the budget and the submittals with Crisp Architects.',
        'Prices every change order before the work is approved.', 'Orders the Marvin bifolds at award.']),
    ('3', 'Principals', 'On every project', ['Joe Montana, Matt Nicosia and Chris Lange stay on every project.',
        'One of them is your contact from the first walkthrough to the final punch list.', 'You are never handed off.']),
]
OS_FEATURES = [('RFIs, Submittals and Change Orders', 'One log. Every item dated and tracked to an answer.'), ('Schedule', 'Roadblocks flagged the day they appear.'),
               ('Field Reports', 'Daily reports and photos from the site.'), ('Drawings and Specs', 'Indexed by sheet and section.'),
               ('Job Cost', 'Tracked against the budget, line by line.'), ('Punch List', 'Tracked to closeout.')]
OS_FOR = [('For Crisp Architects', 'Every question, answer and decision in one place.'), ('For the Owner', 'The budget, the schedule and the site, live.'),
          ('For the Job', 'The field and the office work from the same information.')]
PHONE, ADDR, EMAIL = '(845) 398-1778', '173 N. Route 9W, Congers, NY 10920', P['decisions']['email']

# pre-construction, as sent on the last pre-con agreement, with nothing project-specific
# kit v1.9: no plus-or-minus ranges anywhere a client reads. A range next to their budget reads as "the number could be 30% low".
# Each phase says what it locks instead; the funnel keeps its shape (width only, never labelled with a percentage).
PC_PHASES = [('Budget Setting and Schematic Design Budget', 'Phase 01', 20, 'Your target budget and the scope it buys', 'About 2 weeks', 'Budget Set'),
             ('Design Development Updates', 'Phase 02', 15, 'Selections in place of allowances', 'About 2 weeks per revision', 'Selections'),
             ('Firm Trade Partner Pricing', 'Phase 03', 8, 'Bids from the trades who will build it', '3 to 4 weeks with Phase 04', 'Trade Bids'),
             ('Value Engineering, Built In', 'Phase 04', 3, 'One Guaranteed Maximum Price', 'GMP delivered', 'GMP')]
PC_CARDS = [('Phase 01', 'Budget Setting and Schematic Design Budget', ['Budget setting session with you', 'Survey and soils engineer coordinated', 'Foundation strategy, priced', 'Options priced side by side'],
             'Your target budget and an options menu with three scenarios.'),
            ('Phase 02', 'Design Development Updates', ['Re-priced as the design develops', 'Your plans walked at full scale, one session on us', 'Design-to-budget meetings with your architect', 'Early flags on anything drifting over budget'],
             'An updated budget with more detailed line items, plus a preliminary schedule.'),
            ('Phase 03', 'Firm Trade Partner Pricing', ['Major trades bid to our subcontractors and vendors', 'Site walks with the key trades', 'Bids compared line by line'],
             'A budget backed by real pricing.'),
            ('Phase 04', 'Value Engineering, Built In', ['Money-saving options with trade-offs explained', 'Final pricing', 'Guaranteed Maximum Price proposal'],
             'Your GMP and a construction schedule.')]
PC_PROMISES = [('Walk-Away Guarantee', 'Stop at the end of any phase. You pay only for the phases started and keep everything delivered.'),
               ('Unused Reserve Comes Back to You', '100% of unused contingency is credited back. It is used only with written notice.'),
               ('One Rate, No Matter Who Is on the Call', 'One hourly rate, billed for one person per meeting.')]
PC_PRINCIPLES = [('You Are Heard From Day One', 'Your vision, your budget, your timeline. We listen before we estimate.'),
                 ('You Understand Every Number', 'We walk you through each line item so you know exactly what drives cost and where tradeoffs live.'),
                 ('You Own the Strategy', 'We produce options, with the price of each one, so the plan at GMP is yours.')]

# ================================================================ assets
SITE, NET = os.path.join(HERE, 'site'), os.path.join(HERE, 'netlify')
for p_ in (SITE, NET): shutil.rmtree(p_, ignore_errors=True)
A = os.path.join(SITE, 'assets'); [os.makedirs(os.path.join(A, s), exist_ok=True) for s in ('m', 'team')]
# Walk Your Plans (kit v1.9): one full-scale plan session at Walk Your Plans New York, on Montana. The image is the job's own floor plan
# laid on a floor in perspective (tools/wyp/plan_floor.py). JOB: the plan image, and the rooms named in WYP_TEXT.
WYP = os.path.join(CB, 'wyp', 'wyp_plan.jpg')
if os.path.exists(WYP): shutil.copy(WYP, os.path.join(A, 'wyp_plan.jpg'))
WYP_TEXT = 'One session at Walk Your Plans in Port Chester, on Montana. Your floor plans are laid out at full scale, so you can walk every room and settle the layout before anything is framed.'
WYP_NAME = (P['project'].get('name') if isinstance(P.get('project'), dict) else None) or str(P.get('title') or 'your project')
# proposal PDF (kit v1.8): the one PDF in <job>/proposal/ ships with the page and becomes the Proposal PDF button, top right
PDFS = sorted(glob.glob(os.path.join(CB, 'proposal', '*.pdf')))
if len(PDFS) > 1: die('keep one proposal PDF in proposal/, found ' + ', '.join(os.path.basename(f) for f in PDFS))
PDF = os.path.basename(PDFS[0]) if PDFS else None
if PDF: shutil.copy(PDFS[0], os.path.join(A, PDF))
CACHE = os.path.join(HERE, '.cache'); os.makedirs(os.path.join(CACHE, 'm'), exist_ok=True)
def cached(src, dst):
    c = os.path.join(CACHE, dst)
    if os.path.exists(c) and os.path.getmtime(c) > os.path.getmtime(src): shutil.copy(c, os.path.join(A, dst)); return True
    return False
def keep(dst): shutil.copy(os.path.join(A, dst), os.path.join(CACHE, dst))
def jpg(src, dst, width=None, q=82):
    if cached(src, dst): return
    im = Image.open(src).convert('RGB')
    if width and im.width > width: im = im.resize((width, round(im.height * width / im.width)), Image.LANCZOS)
    im.save(os.path.join(A, dst), 'JPEG', quality=q, optimize=True, progressive=True); keep(dst)
H = os.path.join(CB, 'hero')
jpg(os.path.join(H, 'p_se_r.png'), 'hero.jpg', q=84); jpg(os.path.join(H, 'p_se_p.png'), 'hero_p.jpg', q=82)
for n, f in [('g1.jpg', 'hero_front.png'), ('g2.jpg', 'hero_se.png'), ('g3.jpg', 'hero_dusk.png')]: jpg(os.path.join(H, f), n, width=1760, q=80)
if SEL: SELX.update(A=A, HERO=os.path.join(H, 'p_se_r.png')); selbuild.assets(SEL, SELX)   # JOB: the photo the crop:hero: refs read
jpg(os.path.join(H, 'p_dusk_d.png'), 'dusk.jpg', q=80); jpg(os.path.join(H, 'p_dusk_p.png'), 'dusk_p.jpg', q=80)
def edge_clean(im):
    a = np.asarray(im).copy(); al = a[..., 3].astype(float); al[al < 8] = 0
    h, w = al.shape
    rx = np.minimum(np.clip(np.arange(w) / (0.05 * w), 0, 1), np.clip((w - 1 - np.arange(w)) / (0.10 * w), 0, 1))
    ry = np.minimum(np.clip(np.arange(h) / (0.05 * h), 0, 1), np.clip((h - 1 - np.arange(h)) / (0.07 * h), 0, 1))
    f = (ry[:, None] * rx[None, :]) ** 0.8
    al = np.where(al >= 250, al, al * f); a[..., 3] = al.astype(np.uint8)   # fade only the soft shadow, so the block corners stay crisp
    return Image.fromarray(a, 'RGBA')
for f in FRAMES:
    src = os.path.join(MD, f + '.png')
    if not cached(src, 'm/' + f + '.webp'):
        edge_clean(Image.open(src).convert('RGBA').crop(CROP)).save(os.path.join(A, 'm', f + '.webp'), 'WEBP', quality=88, method=6); keep('m/' + f + '.webp')
for _n, _t, ph in TEAM: shutil.copy(os.path.join(KIT, 'assets', 'img', ph), os.path.join(A, 'team', ph))
IMG = os.path.join(KIT, 'assets', 'img')   # kit v1.5: logos and fonts come from the kit
shutil.copy(os.path.join(IMG, 'mc-wordmark-rev.png'), os.path.join(A, 'mc-white.png'))
shutil.copy(os.path.join(IMG, 'mc-wordmark.png'), os.path.join(A, 'mc-ink.png'))
wm = Image.open(os.path.join(IMG, 'mc-wordmark-rev.png')).convert('RGBA')
Image.merge('RGBA', (Image.new('L', wm.size, 255),) * 3 + (wm.split()[3],)).save(os.path.join(A, 'mc-white-all.png'), optimize=True)
font = lambda f: base64.b64encode(open(os.path.join(KIT, 'assets', 'fonts', f), 'rb').read()).decode()

# ================================================================ the M, traced from M LOGO 2026 (three pieces)
M_PIECES = ['0,0 0,652 183,652 183,275 380,488 515,350 196,0', '773,124 586,310 586,652 773,652', '773,0 571,0 399,175 533,322 773,82']
def mark(cls, fill='currentColor'):
    return f'<svg class="{cls}" viewBox="0 0 773 652" aria-hidden="true">' + ''.join(f'<polygon class="mp mp{k}" points="{p}" fill="{fill}"/>' for k, p in enumerate(M_PIECES)) + '</svg>'
FAVICON = 'data:image/svg+xml;base64,' + base64.b64encode(f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="-60 -120 893 893"><rect x="-60" y="-120" width="893" height="893" rx="160" fill="#0028cc"/>{"".join(f"<polygon points=\"{p}\" fill=\"#fff\"/>" for p in M_PIECES)}</svg>'.encode()).decode()

# ================================================================ markup
DLICON = '<svg viewBox="0 0 16 16" fill="none" aria-hidden="true"><path d="M8 2v8.5M4.5 7 8 10.5 11.5 7M2.5 13.5h11" stroke="currentColor" stroke-width="1.4"/></svg>'
def pdf_link(cls):
    return f'<a class="{cls}" href="assets/{E(PDF)}" download="{E(PDF)}" data-pdf>{DLICON}<span>Proposal PDF</span></a>' if PDF else ''
ARROW = '<svg viewBox="0 0 16 16" fill="none" aria-hidden="true"><path d="M2 8h11M9 3.5 13.5 8 9 12.5" stroke="currentColor" stroke-width="1.4"/></svg>'
PLUS = '<svg class="pl" viewBox="0 0 16 16" fill="none" aria-hidden="true"><path d="M8 2v12M2 8h12" stroke="currentColor" stroke-width="1.4"/></svg>'
NAV = [('project', 'Project'), ('design', 'Design'), ('budget', 'Budget')] + ([('selections', 'Selections')] if SEL else []) + [('schedule', 'Schedule'), ('approach', 'Approach')]
kick = lambda t, cls='': f'<p class="kick {cls}"><i></i>{E(t)}</p>'
def odo(v, i=''):
    return f'<span class="odo num"{f" id={chr(34)}{i}{chr(34)}" if i else ""} data-v="{E(v)}" role="img" aria-label="{E(v)}">{E(v)}</span>'
def ramp(n, a=(0, 26, 138), b=(207, 216, 247)):
    return ['#%02x%02x%02x' % tuple(round(a[k] + (b[k] - a[k]) * (i / max(1, n - 1))) for k in range(3)) for i in range(n)]
OPC = '#9aa1b2'

s0 = STEPS[0]
stage_imgs = ''.join(f'<img src="assets/m/{f}.webp" alt="" decoding="async"{" loading=lazy" if i > 3 else ""}>' for i, f in enumerate(FRAMES))
nav_links = ''.join(f'<a href="#{i}">{E(t)}</a>' for i, t in NAV)
menu_links = ''.join(f'<a href="#{i}"><small>{k + 1:02d}</small>{E(t)}</a>' for k, (i, t) in enumerate(NAV))
elist = ''.join(f'<button type="button"><span>{i + 1:02d}</span><span>{E(n)}</span><span class="num">{money(v)}</span></button>' for i, (k, n, v, s) in enumerate(ELS))
elist += f'<div class="elist__rest"><span></span><span>Building, Site and Systems</span><span class="num">{money(REST)}</span></div>'
rail = ''.join(f'<button type="button" data-step="{j}"{" class=on" if j == 0 else ""}>{t}</button>' for j, t in [(0, 'Site and Foundation'), (1, 'Workshop'), (2, 'Lounge'), (3, 'Exterior'), (4, TITLE_EL)])

def view_rows(items, colors, expandable):
    out = ''
    for k, (it, c) in enumerate(zip(items, colors)):
        lines = ''.join(f'<li>{"<span class=\"cd num\">" + E(ln["code"]) + "</span>" if ln.get("code") else ""}<span>{E(ln["item"])}</span></li>' for ln in it['lines'])   # line items only, no line pricing
        scope = f'<p>{E(it["scope"])}</p>' if it.get('scope') else ''
        head = f'<span class="sq"></span><div class="brow__t"><h3>{E(it["head"])}</h3>{scope}</div><span class="v num">{money(it["v"])}</span><span class="pct num">{pct(it["v"], BASE)}</span>'
        out += (f'<div class="brow" style="--c:{c}" data-k="{k}"><button type="button" class="brow__h" aria-expanded="false">{head}{PLUS}</button><ul class="brow__l" hidden>{lines}</ul></div>')
    return out
gcol = ['#001a8a', '#0028cc', '#2f50d6', '#5a74df', '#8598e8', '#adbbf0', '#cfd8f7']
dcol = ramp(len(DIVS))
EL_ITEMS = [{'head': g['name'], 'scope': g['scope'], 'v': g['v'], 'lines': g['lines']} for g in GROUPS]
MF_ITEMS = [{'head': f'{dv["n"]}  {dv["name"]}', 'v': dv['v'], 'lines': dv['lines']} for dv in DIVS]
def segs(items, colors):
    return ''.join(f'<i style="--w:{float(it["v"] / BASE * 100):.3f}%;--c:{c};--k:{k}" title="{E(it["head"])}"></i>' for k, (it, c) in enumerate(zip(items, colors))) + \
        f'<i style="--w:{float(OP / BASE * 100):.3f}%;--c:{OPC};--k:{len(items)}" title="Overhead and Profit"></i>'
op_row = f'<div class="brow brow--op" style="--c:{OPC}"><div class="brow__h"><span class="sq"></span><div class="brow__t"><h3>Overhead and Profit</h3><p>15% on the direct cost.</p></div><span class="v num">{money(OP)}</span><span class="pct num">{pct(OP, BASE)}</span><span class="pl"></span></div></div>'

g_ms = ''.join(f'<div style="left:{w / SPAN * 100:.3f}%"><b>Week {w}</b><span>{E(l)}</span></div>' for w, l in MS)
g_rows = ''.join(f'<div class="g-row" data-a="{r["a"]:.3f}" data-b="{r["b"]:.3f}"><div class="g-name">{E(r["name"])}</div><div class="g-track"><i class="g-bar" style="left:{r["a"] / SPAN * 100:.3f}%;width:{(r["b"] - r["a"]) / SPAN * 100:.3f}%;--c:{r["c"]}"><b class="g-fill"></b></i></div></div>' for r in SROWS)
g_legend = ''.join(f'<li style="--c:{c}" data-a="{a:.3f}" data-b="{b:.3f}"><i></i><b>{E(n)}</b><span>Weeks {a:.0f} to {b:.0f}</span><em aria-hidden="true"></em></li>' for n, c, a, b in PH_SPANS)
g_grid = ''.join(f'<u style="left:{a / SPAN * 100:.3f}%;width:{(b - a) / SPAN * 100:.3f}%"></u>' for a, b in ((5, 10), (15, 20))) + \
    ''.join(f'<i style="left:{w / SPAN * 100:.3f}%"></i>' for w in range(0, 27))
g_axis = '<i></i><div>' + ''.join(f'<span style="left:{t / SPAN * 100:.3f}%">{"Week " if t == 0 else ""}{t}</span>' for t in (0, 5, 10, 15, 20, 26)) + '</div>'

acard = lambda a: f'<div class="acard"><span class="acard__v num">{E(a["v"])}</span><h3>{E(a["label"])}</h3><p>{E(a["basis"])}</p></div>'
_num = lambda v: float(re.sub(r'[^\d.]', '', v.split(' per ')[0]) or 0)
_al = sorted([a for a in ALLOWS if 'per SF' not in a['v']], key=lambda a: -_num(a['v'])); _amax = _num(_al[0]['v']) if _al else 1   # ranked ledger, a bar per allowance
allow_rows = ''.join(f'<div class="aled__row" style="--w:{_num(a["v"]) / _amax * 100:.2f}%;--k:{k}"><span class="aled__n num">{k + 1:02d}</span><div class="aled__t"><h3>{E(a["label"])}</h3><p>{E(a["basis"])}</p></div><div class="aled__bar" aria-hidden="true"><i></i></div><span class="aled__v num">{E(a["v"])}</span></div>' for k, a in enumerate(_al))
rate_cells = ''.join(f'<div class="arate__c"><div class="arate__t"><h3>{E(a["label"])}</h3><p>{E(a["basis"])}</p></div><span class="arate__v num">{E(a["v"].split(" per ")[0])}<small>per {E(a["v"].split(" per ")[1])}</small></span></div>' for a in ALLOWS if 'per SF' in a['v'])
alts = ''.join(f'<div class="alt"><div><h3>{E(a["title"])}</h3><p>{E(a["note"])}</p></div><span class="v num">{money(a["v"])}</span><button type="button" class="tog" role="switch" aria-checked="false" data-id="{a["id"]}" aria-label="Add {E(re.sub(r"^[0-9]+[.] ", "", a["title"]))}"></button></div>' for a in ALTS)
ncol = lambda h, items: f'<div class="ncol"><h3>{E(h)}</h3><ol>{"".join(f"<li>{E(x)}</li>" for x in items)}</ol></div>'
tiers = ''.join(f'<article class="tier" data-r="up"><b class="tier__n num">{n}</b><h3>{E(h)}</h3><p class="tier__sub">{E(sub)}</p><ul>{"".join(f"<li>{E(x)}</li>" for x in xs)}</ul></article>' for n, h, sub, xs in TIERS)
people = ''.join(f'<figure class="person"><div class="person__img"><img src="assets/team/{ph}" alt="{E(n)}" loading="lazy"></div><figcaption><b>{E(n)}</b><span>{E(t)}</span></figcaption></figure>' for n, t, ph in TEAM)
os_feat = ''.join(f'<li><span class="osx__no">{k + 1:02d}</span><b>{E(a)}</b><em>{E(b)}</em></li>' for k, (a, b) in enumerate(OS_FEATURES))
# Montana OS (kit v1.10): outline OS letters top right, and this job as it opens in the system, its real numbers as tiles
OS_LETTERS = ('<svg class="osx__os" viewBox="-4 -4 548 258" aria-hidden="true"><g fill="none" stroke="currentColor" stroke-width="2.5" stroke-linejoin="miter">'
              '<polygon points="48,0 202,0 250,48 250,202 202,250 48,250 0,202 0,48"/><polygon points="92,62 158,62 180,84 180,166 158,188 92,188 70,166 70,84"/>'
              '<polygon points="342,0 540,0 540,62 378,62 366,74 366,92 378,104 488,104 540,156 540,198 488,250 290,250 290,188 452,188 464,176 464,158 452,146 342,146 290,94 290,52"/></g></svg>')
OS_TILES = [('Base Budget', money(BASE), 'Job cost tracked against it, line by line.', True),
            ('Line Items', f'{len(LINES)}', 'Every line of this budget, in one place.', False),
            ('Schedule', f'{SCH["weeks"]} Weeks', 'From mobilization. Roadblocks flagged the day they appear.', False)]
if SEL: OS_TILES.append(('Selections', f'{len(SEL["elements"])}', 'Each one requested and confirmed in writing.', False))
os_tiles = ''.join(f'<div class="osx__tile{" osx__tile--navy" if hi else ""}"><span class="osx__tk">{E(k)}</span><b class="num">{E(v)}</b><p>{E(d)}</p><em aria-hidden="true">{i + 1:02d}</em></div>' for i, (k, v, d, hi) in enumerate(OS_TILES))
os_for = ''.join(f'<div><p>{E(a)}</p><h4>{E(b)}</h4></div>' for a, b in OS_FOR)
GAL = [('g1.jpg', 'South Elevation'), ('g2.jpg', 'Southeast'), ('g3.jpg', 'Southeast at Dusk')]

# the pre-con funnel: the price gets more certain at each phase (shape only, no percentages)
def funnel(cls='funnel'):   # one tick per phase at its column center, so the labels below line up
    W_, H_, mid, n = 1000, 300, 150, len(PC_PHASES)
    cx = [W_ * (i + .5) / n for i in range(n)]; hw = [p[2] / PC_PHASES[0][2] * 120 for p in PC_PHASES]
    xs, ws = [20] + cx + [W_ - 20], [hw[0]] + hw + [hw[-1]]
    top = f'M{xs[0]},{mid - ws[0]}' + ''.join(f' C{(xs[i - 1] + xs[i]) / 2},{mid - ws[i - 1]} {(xs[i - 1] + xs[i]) / 2},{mid - ws[i]} {xs[i]},{mid - ws[i]}' for i in range(1, len(xs)))
    bot = f'L{xs[-1]},{mid + ws[-1]}' + ''.join(f' C{(xs[i] + xs[i - 1]) / 2},{mid + ws[i]} {(xs[i] + xs[i - 1]) / 2},{mid + ws[i - 1]} {xs[i - 1]},{mid + ws[i - 1]}' for i in range(len(xs) - 1, 0, -1)) + 'Z'
    lines = ''.join(f'<line x1="{x}" y1="{mid - hw[i] - 10}" x2="{x}" y2="{mid + hw[i] + 10}" class="fz-tick"/>' for i, x in enumerate(cx))
    return f'<svg class="{cls}" viewBox="0 0 {W_} {H_}" role="img" aria-label="The price gets more certain across four phases and ends in a Guaranteed Maximum Price"><path class="fz-band" d="{top} {bot}"/><line x1="{xs[0]}" y1="{mid}" x2="{xs[-1]}" y2="{mid}" class="fz-mid"/>{lines}<circle cx="{xs[0]}" cy="{mid}" r="7" class="fz-dot"/></svg>'
pc_cols = ''.join(f'<div class="pcph"><span class="pcph__k">{E(k)}</span><h4>{E(n)}</h4><p class="pcph__lab">Locks In</p><p class="pcph__lock">{E(lk)}</p><span class="pcph__t">{E(t)}</span></div>'
                  for n, k, w, lk, t, short in PC_PHASES)
pc_cards = ''.join(f'<article class="pccard"><div class="pccard__top"><span>{E(k)}</span><span>{E(next(p[4] for p in PC_PHASES if p[1] == k))}</span></div><h4>{E(t)}</h4><ul>{"".join(f"<li>{E(x)}</li>" for x in xs)}</ul><p class="pccard__get"><span>You Get</span>{E(g)}</p></article>' for k, t, xs, g in PC_CARDS)
PCF_LABS = ''.join(f'<span><b>{E(short)}</b>{E(k)}</span>' for n, k, w, lk, t, short in PC_PHASES)
pc_prom = ''.join(f'<div class="pcprom"><h4>{E(a)}</h4><p>{E(b)}</p></div>' for a, b in PC_PROMISES)
pc_prin = ''.join(f'<div class="pcprin"><h4>{E(a)}</h4><p>{E(b)}</p></div>' for a, b in PC_PRINCIPLES)

DATA = {'base': str(BASE), 'frame': list(FR), 'points': {k: PTS[k] for k, *_r in ELS}, 'bb': BB,
        'els': [{'k': k, 'name': n, 'v': str(v)} for k, n, v, s in ELS], 'rest': str(REST),
        'alts': [{'id': a['id'], 'title': a['title'], 'v': str(a['v'])} for a in ALTS], 'span': SPAN, 'weeks': SCH['weeks'],
        'job': P.get('project', {}).get('job', ''), 'email': EMAIL, 'sel': selbuild.data(SEL, SELX) if SEL else None,
        'steps': [{k: s.get(k) for k in ('label', 'title', 'sub', 'fig', 'cap', 'body', 'img', 'via', 'el', 'all', 'rail', 'note') if s.get(k) is not None} for s in STEPS]}

BODY = f'''
<div class="loader" id="loader" aria-hidden="true"><div class="loader__in">{mark("loader__m")}<div class="loader__bar"><i id="loader-bar"></i></div></div></div>

<header class="nav" id="nav">
  <a class="nav__logo" href="#top" aria-label="Montana Contracting, back to the top"><img class="lg-white" src="assets/mc-white.png" alt="Montana Contracting" width="240" height="14"><img class="lg-ink" src="assets/mc-ink.png" alt="" width="240" height="14"><img class="lg-allw" src="assets/mc-white-all.png" alt="" width="240" height="14"></a>
  <nav class="nav__links" aria-label="Sections">{nav_links}</nav>
  <div class="nav__right">{pdf_link("nav__cta nav__pdf")}<button type="button" class="nav__menu" id="menu-btn" aria-expanded="false" aria-controls="menu" aria-label="Menu"><span></span><span></span></button></div>
  <i class="nav__prog" aria-hidden="true"></i>
</header>
<div class="menu" id="menu" inert data-lenis-prevent><button type="button" class="menu__close">Close</button><nav aria-label="Sections">{menu_links}</nav>{pdf_link("btn btn--ghost")}</div>

<main>
<section class="hero" id="top" aria-label="The Cooper Barn">
  <div class="hero__bg"><picture><source media="(max-width: 760px)" srcset="assets/hero_p.jpg"><img id="hero-img" data-crit fetchpriority="high" src="assets/hero.jpg" alt="The Cooper Barn, late afternoon" width="1600" height="900"></picture></div>
  <div class="hero__scrim"></div><div class="hero__shade"></div>
  <div class="hero__inner gut">
    <div class="tags"><span class="tag">Budget Presentation</span><span class="tag tag--ghost">{LOC}</span><span class="tag tag--ghost">Workshop and Lounge</span></div>
    <h1 data-r>The Cooper Barn</h1>
    <p class="eyebrow"><i></i>Crisp Architects · 194 Titicus Road</p>
  </div>
  <a class="hero__scroll" href="#project">Scroll<i></i></a>
</section>

<section class="intro" id="project" aria-label="Project">
  <div class="wrap">
    {kick("01 / Project")}
    <h2 class="sec-title" data-r>Project Details</h2>
    <div class="intro__grid">
      <dl class="facts" data-r="up">
        <div><dt>Location</dt><dd>{LOC}</dd></div>
        <div><dt>Type</dt><dd>Workshop and Lounge</dd></div>
        <div><dt>Architect</dt><dd>Crisp Architects</dd></div>
        <div><dt>Budget</dt><dd class="num">{money(BASE)}</dd></div>
        <div><dt>Schedule</dt><dd>{SCH["weeks"]} Weeks</dd></div>
      </dl>
      <div class="desc"><p class="lede">{E(PROSE[0])}</p>{"".join(f"<p data-r=up>{E(p)}</p>" for p in PROSE[1:])}</div>
    </div>
  </div>
</section>

<section class="gallery" id="gallery" aria-label="Gallery">
  <div class="gallery__pin">
    <div class="wrap gallery__head">{kick("Gallery")}<span class="gallery__count num"><b>01</b> / 03</span></div>
    <div class="track">{"".join(f'<figure class="shot"><div class="shot__img"><img src="assets/{f}" alt="The Cooper Barn, {E(c)}" loading="lazy"></div><figcaption><b>{k + 1:02d}</b>{E(c)}</figcaption></figure>' for k, (f, c) in enumerate(GAL))}</div>
  </div>
</section>

<section class="stage" id="stage" aria-label="Design">
  <span class="stage__anchor" id="design" style="top:0"></span>
  <div class="stage__pin">
    <div class="stage__model"><div class="frame" data-mk="brackets" style="--fa:{FR[0]} / {FR[1]}">{stage_imgs}
      <span class="mk" id="mk" aria-hidden="true"></span>
      <span class="mkb" id="mkb" aria-hidden="true"><i></i><i></i><i></i><i></i></span>
      <span class="mkl" id="mkl" aria-hidden="true"><i></i><b></b></span>
      <span class="mkn" id="mkn" aria-hidden="true"></span>
      {"".join(f'<span class="mks" data-k="{k}" aria-hidden="true">{i + 1}</span>' for i, (k, *_r) in enumerate(ELS))}
    </div></div>
    <div class="stage__copy" aria-live="polite">
      <p class="kick sc-label sw">{E(s0["label"])}</p>
      <h2 class="sc-title sw">{E(s0["title"])}</h2>
      <p class="sc-sub sw"></p>
      <div class="sc-fig">{odo(s0["fig"], "stage-odo")}</div>
      <p class="sc-cap sw">{E(s0["cap"])}</p>
      <p class="sc-body sw">{E(s0["body"])}</p>
      <div class="elist hide"><p class="elist__h">Base Budget with Overhead and Profit</p>{elist}<div class="elist__tot"><span></span><span>Base Budget</span><span class="num">{money(BASE)}</span></div></div>
    </div>
    <div class="stage__rail" role="group" aria-label="Steps">{rail}</div>
    <p class="stage__note">{E(s0["note"])}</p>
  </div>
</section>

<section class="budget" id="budget" aria-label="Budget">
  <div class="wrap">
    {kick("03 / Budget")}
    <div class="budget__fig">{odo(money(BASE), "budget-odo")}</div>
    <p class="budget__cap" data-r="up"><span><b>Base Budget</b> with Overhead and Profit</span><span><b class="num">{money(DIRECT)}</b> Direct Cost</span><span><b>15%</b> Overhead and Profit</span></p>
    <div class="bview" data-view="el">
      <div class="sbar sbar--el" aria-hidden="true">{segs(EL_ITEMS, gcol)}</div>
      <div class="sbar sbar--mf" aria-hidden="true">{segs(MF_ITEMS, dcol)}</div>
      <div class="pills" role="tablist" aria-label="Budget view"><button type="button" role="tab" class="pill on" aria-selected="true" data-v="el">Elements</button><button type="button" role="tab" class="pill" aria-selected="false" data-v="mf">MasterFormat</button><span class="pills__hint">Lines at direct cost. Open any line for its detail.</span></div>
      <div class="brows brows--el" role="tabpanel" aria-label="Elements">{view_rows(EL_ITEMS, gcol, True)}</div>
      <div class="brows brows--mf" role="tabpanel" aria-label="MasterFormat" hidden>{view_rows(MF_ITEMS, dcol, True)}</div>
    </div>
    <div class="btot"><div class="btot__direct"><span>Direct Cost</span><span class="num">{money(DIRECT)}</span></div><div class="btot__op" style="--c:{OPC}"><span><i class="sq" aria-hidden="true"></i>Overhead and Profit, 15%</span><span class="num">{money(OP)}</span></div><div><span>Base Budget</span><span class="num">{money(BASE)}</span></div></div>
  </div>
</section>

{selbuild.section(SEL, SELX, kick) if SEL else selbuild.sheet(None, E)}

<section class="allow" id="allowances" aria-label="Allowances">
  <div class="wrap">
    {kick("03 / Allowances")}
    <div class="allow__top"><h2 class="sec-title" data-r>Allowances</h2><div class="allow__fig"><span class="big2 num" data-r="up">{money(ALLOW_TOTAL)}</span><p>Carried in the base budget, before overhead and profit.</p></div></div>
    <div class="aled" data-r="up">{allow_rows}</div>
    <div class="arate" data-r="up"><p class="arate__h">Material Allowances, Per SF</p><div class="arate__row">{rate_cells}</div></div>
  </div>
</section>

<section class="alts" id="alternates" aria-label="Alternates">
  <div class="wrap">
    {kick("03 / Alternates")}
    <div class="alts__grid">
      <div>
        <h2 class="sec-title" data-r>Alternates</h2>
        <p class="alts__lede" data-r="up">Priced with overhead and profit. None of the three is in the base budget.</p>
        <div class="alist">{alts}</div>
      </div>
      <aside class="atot" aria-label="Running total">
        <dl><div><dt>Base Budget</dt><dd class="num">{money(BASE)}</dd></div><div><dt>Requested</dt><dd class="num" id="alt-add">{money(0)}</dd></div></dl>
        <div class="atot__fig"><span class="lab">If Montana Confirms</span>{odo(money(BASE), "alt-odo")}</div>
        <button type="button" class="btn btn--navy" id="alt-send"><span>Send to Montana</span></button>
        <p class="atot__send">{"These count with your choices in Selections. " if SEL else ""}They stay on this device until you send them.</p>
      </aside>
    </div>
  </div>
</section>

<section class="notes" id="notes" aria-label="Notes and Clarifications">
  <div class="wrap">
    {kick("03 / Notes and Clarifications")}
    <h2 class="sec-title" data-r>Notes and Clarifications</h2>
    <div class="ngrid">{ncol("Basis of Budget", BASIS)}{ncol("Clarifications", CLAR)}{ncol("Exclusions", EXCL)}</div>
  </div>
</section>

<section class="sched" id="schedule" aria-label="Schedule">
  <div class="sched__pin">
    <div class="wrap sched__grid">
      <aside class="sched__side">
        {kick("04 / Schedule")}
        <div class="sched__top">
          <span class="big num" data-r="up">{SCH["weeks"]}</span>
          <h2 data-r>Weeks on site from mobilization.</h2>
        </div>
        <p class="sched__now"><span>Week</span>{odo("00", "week-odo")}<span class="sched__ph" id="week-ph">Site and Foundation</span></p>
        <ul class="g-legend">{g_legend}</ul>
      </aside>
      <div class="gantt" role="img" aria-label="Schedule by trade, {SCH['weeks']} weeks from mobilization to work complete">
        <div class="g-ms">{g_ms}</div>
        <div class="g-body"><div class="g-grid" aria-hidden="true">{g_grid}<b class="g-cursor" id="g-cursor"><span class="g-cursor__l">Week <em id="g-cw">0</em></span></b></div>{g_rows}</div>
        <div class="g-axis">{g_axis}</div>
      </div>
    </div>
  </div>
</section>

<section class="appr" id="approach" aria-label="Approach">
  <div class="wrap">
    {kick("05 / Approach")}
    <h2 class="appr__h" data-r>How We Work</h2>
    <p class="appr__big" data-r="up">Every superintendent is a finish carpenter by trade.</p>
    <p class="appr__lede" data-r="up">The person running the schedule can pick up a tool and close a gap in it. Every detail is checked by someone who has built it.</p>
    <div class="yteam">
      <div class="yteam__head" data-r="up"><h3 class="yteam__h">Your Team</h3><p class="yteam__p">Three principals stay on this project from the first walkthrough to the final punch list. You are never handed off.</p></div>
      <div class="people">{people}</div>
    </div>
  </div>
  <div class="osx">
    {OS_LETTERS}
    <div class="wrap osx__in">
      {kick("Montana OS")}
      <h3 class="osx__h" data-r>One system for every job.</h3>
      <p class="osx__lede" data-r="up">We built our own operating system for scheduling, cost and field reporting. It gives the owner and the design team a live, honest view of the work as it happens.</p>
      <div class="osx__app" data-r="up">
        <div class="osx__bar"><span class="osx__brand">{mark("osx__bm")}Montana <b>OS</b></span><span class="osx__job">{E(str(WYP_NAME))}</span><span class="osx__live"><i></i>Opens at Award</span></div>
        <div class="osx__tiles">{os_tiles}</div>
      </div>
      <ol class="osx__feat">{os_feat}</ol>
      <div class="osx__for">{os_for}</div>
    </div>
  </div>
</section>

<section class="wyp" id="walk" aria-label="Walk Your Plans">
  {'<div class="wyp__img"><img src="assets/wyp_plan.jpg" alt="The first floor plan laid out at full scale" width="1800" height="930" loading="lazy" decoding="async"></div>' if os.path.exists(WYP) else ''}
  <div class="wrap wyp__in">
    {kick("Walk Your Plans", "kick--light")}
    <h2 data-r>Walk the house before it's built.</h2>
    <p data-r="up">{E(WYP_TEXT)}</p>
    <div class="wyp__act" data-r="up"><a class="btn btn--white" href="mailto:{E(EMAIL)}?subject={E(quote('Walk Your Plans session, ' + str(WYP_NAME)))}">Book the Session {ARROW}</a><a class="wyp__link" href="https://www.walkyourplansnewyork.com/" target="_blank" rel="noopener">walkyourplansnewyork.com</a></div>
  </div>
</section>

<section class="pcf" id="pre-construction" aria-label="Pre-Construction">
  <div class="wrap">
    <div class="pcf__card" data-precon role="button" tabindex="0" aria-label="See our pre-construction process">
      {mark("pcf__m")}
      <div class="pcf__txt">
        <p class="kick kick--light"><i></i>Pre-Construction</p>
        <h2 class="pcf__h">See Our <span class="nb">Pre-Construction</span> Process</h2>
        <p class="pcf__p">Four phases. Each one locks more of the price. The last is a Guaranteed Maximum Price.</p>
        <span class="btn btn--white pcf__btn">See the Process {ARROW}</span>
      </div>
      <div class="pcf__viz" data-r="up" aria-hidden="true">{funnel("funnel funnel--d")}<div class="pcf__labs">{PCF_LABS}</div></div>
    </div>
  </div>
</section>

<section class="dusk" aria-label="Thank You">
  <picture><source media="(max-width: 760px)" srcset="assets/dusk_p.jpg"><img src="assets/dusk.jpg" alt="The Cooper Barn at dusk, lit from inside" loading="lazy"></picture>
  <div class="dusk__scrim"></div>
  <div class="dusk__in gut"><h2 data-r>Thank you for the opportunity to partner with you on this project.</h2></div>
</section>

<section class="cta" id="contact" aria-label="Next Step">
  {mark("cta__m")}
  <div class="wrap cta__in">
    {kick("Next Step", "kick--light")}
    <h2 data-r>{E(P["next_step"])}</h2>
    <div class="cta__act"><a class="btn btn--white mag" href="tel:+18453981778">Call {PHONE} {ARROW}</a></div>
    <dl class="contact"><div><dt>Phone</dt><dd>{PHONE}</dd></div><div><dt>Address</dt><dd>{E(ADDR)}</dd></div><div><dt>Email</dt><dd>{E(EMAIL)}</dd></div></dl>
  </div>
</section>
<footer class="foot"><div class="wrap foot__in"><img src="assets/mc-white-all.png" alt="Montana Contracting" loading="lazy"><span>General Contractor · Since 1984 · NY + NJ + FL</span><span>© 2026 Montana Contracting</span></div></footer>
</main>

<div class="pc" id="precon" inert>
  <div class="pc__panel" role="dialog" aria-modal="true" aria-labelledby="pc-h" data-lenis-prevent>
    <div class="pc__bar"><img src="assets/mc-white-all.png" alt="Montana Contracting" width="200" height="12"><span>Pre-Construction · Optional Next Step</span><button type="button" class="pc__x">Close</button></div>
    <div class="pc__hero">
      {mark("pc__m")}
      <div class="wrap">
        <p class="kick kick--light"><i></i>Pre-Construction</p>
        <h2 id="pc-h">Know the number<br><em>before</em> you build.</h2>
        <p class="pc__lede">We set the budget with you first, then price the building at every design stage until we can offer a Guaranteed Maximum Price.</p>
      </div>
    </div>
    <div class="pc__body">
      <section class="wrap pc__sec">
        <p class="kick"><i></i>01 · Every Phase</p>
        <h3>Your price gets more certain at every phase.</h3>
        <p class="pc__p">Each phase locks part of the price. At the end of Phase 04 the whole number is a Guaranteed Maximum Price.</p>
        <div class="pc__funnel">{funnel()}<div class="pcphs">{pc_cols}</div></div>
      </section>
      <section class="pc__alt"><div class="wrap pc__sec">
        <p class="kick"><i></i>02 · What Each Phase Delivers</p>
        <h3>Each phase makes the estimate more accurate.</h3>
        <p class="pc__p">Each phase ends with an updated budget and a decision about whether to keep going. You can stop at the end of any phase.</p>
        <div class="pccards">{pc_cards}</div>
      </div></section>
      <section class="wrap pc__sec">
        <p class="kick"><i></i>03 · Our Promises</p>
        <h3>Built in from the first phase.</h3>
        <div class="pcproms">{pc_prom}</div>
        <div class="pcprins">{pc_prin}</div>
      </section>
      <section class="pc__end"><div class="wrap pc__sec">
        <h3>Ask us about <span class="nb">pre-construction.</span></h3>
        <p class="pc__p">One of the three principals will walk you through it.</p>
        <dl class="contact"><div><dt>Phone</dt><dd>{PHONE}</dd></div><div><dt>Email</dt><dd>{E(EMAIL)}</dd></div><div><dt>Address</dt><dd>{E(ADDR)}</dd></div></dl>
      </div></section>
    </div>
  </div>
</div>

<script type="application/json" id="cb-data">{json.dumps(DATA).replace("</", "<\\/")}</script>
<script>{open(os.path.join(KIT, "vendor", "lenis", "lenis.min.js")).read()}</script>
<script>{open(os.path.join(HERE, "src", "app.js")).read()}</script>
'''

FONTS = f'''@font-face {{ font-family: "TT Commons Pro"; font-weight: 400; font-style: normal; font-display: block; src: url(data:font/woff2;base64,{font("TTCommonsPro-Regular.woff2")}) format("woff2"); }}
@font-face {{ font-family: "TT Commons Pro"; font-weight: 600 700; font-style: normal; font-display: block; src: url(data:font/woff2;base64,{font("TTCommonsPro-Bold.woff2")}) format("woff2"); }}
'''
CSS = FONTS + open(os.path.join(HERE, 'src', 'style.css')).read()
def tiny(name, w):   # a ~1 KB blurred stand-in, painted under the hero photo until it is decoded
    import io, base64
    im = Image.open(os.path.join(SITE, 'assets', name)).convert('RGB'); im = im.resize((w, max(1, round(im.height * w / im.width))), Image.LANCZOS)
    b = io.BytesIO(); im.save(b, 'JPEG', quality=60); return 'data:image/jpeg;base64,' + base64.b64encode(b.getvalue()).decode()
PH_D, PH_P = tiny('hero.jpg', 40), tiny('hero_p.jpg', 24)
PRELOAD = ('<link rel="preload" as="image" href="assets/hero.jpg" media="(min-width: 761px)" fetchpriority="high">\n'
           '<link rel="preload" as="image" href="assets/hero_p.jpg" media="(max-width: 760px)" fetchpriority="high">\n'
           '<script>document.documentElement.classList.add("m")</script>\n')   # the loader covers the page from the first paint
CSS += f'\n.hero__bg {{ background: #2b2f2a url({PH_D}) center / cover no-repeat; }}\n@media (max-width: 760px) {{ .hero__bg {{ background-image: url({PH_P}); }} }}\n'
TITLE = 'Cooper Barn Presentation'
HEAD = (PRELOAD + f'<title>{TITLE}</title>\n<meta name="claude:viewport-fit-sides" content="cover">\n<link rel="icon" href="{FAVICON}">\n'
        f'<meta name="description" content="Montana Contracting budget proposal for The Cooper Barn, 194 Titicus Road, North Salem, NY.">\n<style>\n{CSS}\n</style>\n')

# ================================================================ house rules on everything a client reads
BANNED = ['stunning', 'beautiful', 'breathtaking', 'luxury', 'luxurious', 'bespoke', 'elevated', 'curated', 'seamless', 'passion', 'dream', 'journey',
          'exceptional', 'world-class', 'premier', 'cutting-edge', 'unparalleled', 'empower', 'transform', 'peace of mind', 'potential risk', 'novaterra',
          'design move', 'nestled', 'state-of-the-art', 'elevate', 'proud to', 'excited', 'grateful']
text = re.sub(r'<script.*?</script>|<svg.*?</svg>', ' ', BODY, flags=re.S)
text = html.unescape(re.sub(r'<[^>]+>', ' ', text))
text += ' ' + ' '.join(str(v) for s in STEPS for v in s.values() if isinstance(v, str))
bad = [name for ch, name in [('—', 'em dash'), ('–', 'en dash'), (';', 'semicolon'), ('!', 'exclamation'), ('?', 'question')] if ch in text]
bad += ['banned: ' + w for w in BANNED if re.search(r'\b' + re.escape(w) + r'\b', text.lower())]
if re.search(r'\b\w+, not \w+', text): bad.append('X, not Y')
if bad: die('copy rules: ' + ', '.join(bad))
if len(PROSE[0].split('. ')[0]) > 155: die('project snippet sentence is over 155 characters')
words = sum(len(p.split()) for p in PROSE)
if not 90 <= words <= 220: die(f'project description is {words} words')
for v in [money(BASE), money(DIRECT), money(OP), money(EL_TOTAL), money(REST), money(SITE_FDN), money(ALLOW_TOTAL)] + [money(e[2]) for e in ELS]:
    if v not in BODY: die('figure missing from the page: ' + v)
_vis = re.sub(r'<script.*?</script>', '', BODY, flags=re.S)
_nc = re.search(r'\$\d[\d,]*(?![\d,]*\.\d\d)', _vis)
if _nc: die('a dollar figure without cents: ' + _vis[max(0, _nc.start() - 120):_nc.end() + 40])

os.makedirs(SITE, exist_ok=True)
open(os.path.join(SITE, 'index.html'), 'w').write(HEAD + BODY)
shutil.copytree(SITE, NET)
open(os.path.join(NET, 'index.html'), 'w').write(f'<!doctype html>\n<html lang="en">\n<head>\n<meta charset="utf-8">\n<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">\n<meta name="robots" content="noindex, nofollow">\n<meta name="mcc-send" content="netlify-forms">\n{HEAD}</head>\n<body>\n{BODY}\n</body>\n</html>\n')
open(os.path.join(NET, 'robots.txt'), 'w').write('User-agent: *\nDisallow: /\n')
# JOB: zip name, <JOB>_Presentation_Netlify_v#.zip
z = os.path.join(HERE, 'E260050_Presentation_Netlify_v5.zip')
with zipfile.ZipFile(z, 'w', zipfile.ZIP_DEFLATED) as zf:
    for dp, _d, fs in os.walk(NET):
        for f in fs: zf.write(os.path.join(dp, f), os.path.relpath(os.path.join(dp, f), NET))
json.dump({'base': money(BASE), 'direct': money(DIRECT), 'op': money(OP), 'elements': {n: money(v) for k, n, v, s in ELS}, 'el_total': money(EL_TOTAL), 'rest': money(REST),
           'site_fdn': money(SITE_FDN), 'allow_total': money(ALLOW_TOTAL), 'groups': {g['name']: money(g['v']) for g in GROUPS}, 'divisions': {dv['n'] + ' ' + dv['name']: money(dv['v']) for dv in DIVS},
           'alternates': {a['title']: money(a['v']) for a in ALTS}, 'allowances': {a['label']: a['v'] for a in ALLOWS}, 'milestones': MS, 'phases': [(n, round(a, 2), round(b, 2)) for n, c, a, b in PH_SPANS]},
          open(os.path.join(HERE, 'figures.json'), 'w'), indent=1)
size = sum(os.path.getsize(os.path.join(dp, f)) for dp, _d, fs in os.walk(SITE) for f in fs)
print(f'built: page {os.path.getsize(os.path.join(SITE, "index.html")) / 1e6:.2f} MB, site {size / 1e6:.2f} MB, {len(STEPS)} steps, crop {CROP}, zip {os.path.getsize(z) / 1e6:.2f} MB')
