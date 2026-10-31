#!/usr/bin/env python3
"""Website-format presentation build. Job: E260050 The Cooper Barn and Pool House, one presentation for the whole property
(kit v1.10, from the E26049 example). Three workbooks feed one page: the Barn, the Pool House and the Shared General Requirements.

    python3 -I build.py

Every figure comes from the estimates' line items (Estimate sheets), to the cent. Each workbook's base ties to its own
TOTAL -- BASE BID row, and the page base is their sum. Writes site/, netlify/, the Netlify zip, figures.json and
proposal_lists.json. Stops if a figure does not tie, or client copy breaks the house rules.
"""
import base64, datetime as dt, glob, html, json, os, re, shutil, sys, zipfile
from urllib.parse import quote
from decimal import Decimal as D, ROUND_HALF_UP
import numpy as np
import openpyxl
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
CB = os.path.dirname(HERE)
KIT = os.environ.get('MCC_KIT') or os.path.join(os.path.dirname(CB), 'kit')
E = lambda s: html.escape(str(s), quote=True)
C = lambda x: D(str(x)).quantize(D('0.01'), ROUND_HALF_UP)
money = lambda v: '${:,.2f}'.format(C(v))
X = lambda x: D(str(round(x, 6)))   # a line extension as the workbook holds it (Qty x Rate can run past the cent); figures round once, at display
pct = lambda v, of: '{:.1f}%'.format(float(D(v) / D(of) * 100))


def die(m):
    print('BUILD STOPPED:', m); sys.exit(1)


P = json.load(open(os.path.join(CB, 'project.json')))
SCH = json.load(open(os.path.join(CB, 'schedule.json')))

# ================================================================ money, from the line items of three workbooks
BOOKS = [('shared', 'estimate_shared.xlsx', 'Barn and Pool House', 200000), ('barn', 'estimate_barn.xlsx', 'The Barn', 0), ('ph', 'estimate_ph.xlsx', 'The Pool House', 100000)]
LINES, WB, BK = {}, {}, {}
for key, fn, label, off in BOOKS:
    wb = openpyxl.load_workbook(os.path.join(CB, fn), data_only=True); ws = wb['Estimate']; WB[key] = wb
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))
    rt = next(i for i, r in enumerate(rows, start=1) if r[0] == 'TOTAL -- BASE BID')
    trade, lines = None, {}
    for i, r in enumerate(rows[9:rt], start=10):
        code, item, qty, unit, rate, tot = r[:6]
        if code and isinstance(code, str) and code.isupper() and not item and code not in ('TRADE TOTAL', 'SUBTOTAL'): trade = code; continue
        if code in ('TRADE TOTAL', 'SUBTOTAL') or not item: continue
        if isinstance(tot, (int, float)) and tot:
            lines[off + i] = {'code': str(code).strip(), 'item': str(item).strip(), 'trade': trade, 'v': X(tot), 'bld': key, 'qty': qty, 'rate': rate}
    direct = sum(l['v'] for l in lines.values()); base = C(direct * D('1.15')); op = base - C(direct)
    sheet = C(ws.cell(rt, 6).value)
    if base != sheet: die(f'{key}: base {direct + op} does not tie to its Estimate sheet total {sheet}')
    alts, cur = {}, None
    for i, r in enumerate(rows[rt:], start=rt + 1):
        h = str(r[0] or '')
        m_ = re.match(r'(?:ADD )?ALTERNATE (\d+) --', h)
        if m_: cur = int(m_.group(1))
        if r[0] == 'TOTAL' and isinstance(r[5], (int, float)) and cur is not None: alts[cur] = C(r[5]); cur = None
    BK[key] = {'label': label, 'direct': C(direct), 'x': direct, 'op': op, 'base': base, 'alts': alts, 'off': off}
    LINES.update(lines)
DIRECT = sum(b['direct'] for b in BK.values()); OP = sum(b['op'] for b in BK.values()); BASE = DIRECT + OP
if BASE != C(sum(b['base'] for b in BK.values())): die('base')


class _Cell:
    def __init__(self, v): self.value = v
class _WS:   # one Estimate sheet view over the three workbooks, for selbuild (row keys carry the workbook offset)
    def cell(self, r, c):
        key = 'ph' if 100000 <= r < 200000 else 'shared' if r >= 200000 else 'barn'
        return _Cell(WB[key]['Estimate'].cell(r - BK[key]['off'], c).value)
ws = _WS()

# JOB: client wording for lines whose estimator notes should not reach the page
ITEM_FIX = {'Sales Tax 8.375% on $126': 'Sales Tax on Window Supply', 'Sales Tax 8.375% on $56': 'Sales Tax on Door Supply',
            'Sales Tax 8.375% on $6,500': 'Sales Tax on Generator Supply', 'Sales Tax 8.375% on $109': 'Sales Tax on Bifold Supply',
            'Outdoor Shower Supply, Valve & Drain': 'Outdoor Shower Supply, Valve and Drain (Allowance)',
            'Glass Package Upgrade to U-0.27': 'Glass Package to U-0.27, 2025 NYS Energy Code (Allowance)',
            'Tempered Glazing @ Hazardous Locations': 'Tempered Glass at Doors, Stair and Large Low Units',
            'LVL Headers @ Ganged Window Openings': 'LVL Headers at the Ganged Windows',
            '10" Frost Walls, 48" Deep': '10" Frost Walls, 48" Deep, East and North Walls Carried Up at the Grade',
            'Gas Trench from Barn Propane Tank': 'Gas Trench from the Barn Propane Tank (Allowance)',
            'Roof Deck Guardrail Allowance': 'Roof Deck Guardrail (Allowance)'}
def clean(item):   # client wording for a line item: no bracketed estimator notes, no drawing references, no unit rates
    for k, v in ITEM_FIX.items():
        if item.startswith(k): return plain(v)
    s = re.sub(r'\s*\[Allowance[^\]]*\]', ' (Allowance)', item)
    s = re.sub(r'\s*\[[^\]]*Allowance[^\]]*\]', ' (Allowance)', s)
    s = re.sub(r'\s*\[[^\]]*\]', '', s); s = s.replace(' (M+L)', '').replace('(M+L)', '')
    s = re.sub(r'\(\$[\d,.]+/(?:SF|Ea|EA|LF)[^)]*Material[^)]*\)', '(Material Allowance)', s)
    s = re.sub(r'\(\$[\d,.]+/(?:SF|Ea|EA|LF)[^)]*\)', '(Allowance)', s)
    s = re.sub(r'\$(\d[\d,]*)(?![\d,]*\.\d)', r'$\1.00', s)
    s = re.sub(r'\s*\(Per (Week|Month|Page|Riser|Hour|Day)\)', '', s)
    s = re.sub(r'\(\$[\d,]+\.\d\d/SF Mat\. Allowance\)', '(Material Allowance)', s)
    s = re.sub(r'\(\$[\d,]+\.\d\d Allowance\)', '(Allowance)', s)
    s = s.replace('(~', '(About ').replace('~', 'About ').replace(' @ ', ' at ').replace(' & ', ' and ').replace('&', 'and')
    s = re.sub(r'\(Allowance\)\s*\(Allowance\)', '(Allowance)', s)
    return plain(re.sub(r'\s{2,}', ' ', s).strip())

def plain(s):
    for k, v in (('F/I ', 'Furnish and Install '), ('F/O ', 'Furnish Only '), ('I/O ', 'Install Only '), (' w/ ', ' with '), (' W/ ', ' with ')):
        s = s.replace(k, v)
    return s

def find_rows(*starts, n=None, bld=None):
    out = [r for r, l in LINES.items() if l['item'].startswith(starts) and (bld is None or l['bld'] == bld)]
    if not out or (n is not None and len(out) != n): die(f'expected {n} rows for {starts} in {bld}, found {len(out)}')
    return out
def code_rows(*prefixes, bld=None): return [r for r, l in LINES.items() if l['code'].startswith(prefixes) and (bld is None or l['bld'] == bld)]
def trade_rows(bld, *trades): return [r for r, l in LINES.items() if l['bld'] == bld and l['trade'] in trades]
tot = lambda rows: sum(LINES[r]['v'] for r in rows)

# ---------------------------------------------------------------- Elements view: groups by building
G_DEF = [
    ('Shared General Requirements', 'Barn and Pool House. One superintendent and one project manager for both buildings, temporary facilities, protection and final cleaning.', 'shared', ['GENERAL REQUIREMENTS']),
    ('Barn Site and Foundations', 'Excavation, footings, frost walls, the mechanical level, slabs, utilities and waste removal.', 'barn', ['WASTE REMOVAL', 'SITEWORK & SITE UTILITIES', 'CONCRETE & FOUNDATIONS']),
    ('Barn Structure and Roof Deck', 'Framing, exposed beams, canopies, the Cumaru roof deck and its guardrail.', 'barn', ['ROUGH CARPENTRY & FRAMING', 'DECKING', 'MISCELLANEOUS METALS']),
    ('Barn Building Envelope', 'Insulation, standing seam roof, siding, trim, the cupola, Marvin windows and bifold doors.', 'barn', ['WATERPROOFING & INSULATION', 'ROOFING, GUTTERS & LEADERS', 'SIDING & EXTERIOR TRIM', 'WINDOWS, DOORS & HARDWARE']),
    ('Barn Mechanical and Electrical', 'Plumbing, two zone heating and cooling, ventilation, power, lighting and low voltage.', 'barn', ['PLUMBING & FUEL GAS', 'HVAC', 'ELECTRICAL, LOW VOLTAGE & DETECTION', 'EQUIPMENT']),
    ('Barn Interior Finishes', 'Drywall, stairs, wainscot, reclaimed beams, the copper top bar, millwork, flooring, tile, paint and the wood stove.', 'barn', ['DRYWALL', 'FINISH CARPENTRY, STAIRS & RAILINGS', 'CUSTOM MILLWORK & COUNTERTOPS', 'WOOD FLOORING', 'TILE', 'PAINTING & STAINING', 'SPECIALTIES']),
    ('Pool House Site and Foundations', 'Removal of the existing pavilion, excavation, frost walls, the slab, waterproofing, drainage, utilities and bluestone steps.', 'ph', ['DEMOLITION & WASTE REMOVAL', 'CAST-IN-PLACE CONCRETE', 'FOUNDATION WATERPROOFING', 'SITEWORK', 'HARDSCAPING', 'SITE UTILITIES', 'DRAINAGE & STORMWATER']),
    ('Pool House Structure', 'Framing, the ridge beam and the steel frame at the south bifold wall.', 'ph', ['WOOD FRAMING', 'STRUCTURAL STEEL']),
    ('Pool House Building Envelope', 'Insulation, cedar shingle roof with copper, Hardie lap siding, Claymark trim, Marvin Ultimate windows and bifold doors.', 'ph', ['THERMAL & SOUND INSULATION', 'SIDING & EXTERIOR TRIM', 'ROOFING', 'EXTERIOR DOORS', 'WINDOWS']),
    ('Pool House Mechanical and Electrical', 'Plumbing, the outdoor shower, heat pumps, ventilation, power, lighting and the standby generator.', 'ph', ['PLUMBING', 'HVAC', 'ELECTRICAL']),
    ('Pool House Audio Visual and Shades', 'Sonos audio, Lutron lighting control and motorized shades.', 'ph', ['AUDIO VISUAL & SHADES']),
    ('Pool House Interior Finishes', 'Trim, the bar and vanity, the white oak stair, doors, drywall, tile, the stone and oak floors, paint, appliances and countertops.', 'ph', ['FINISH CARPENTRY', 'CUSTOM CABINETRY', 'WOOD STAIRS & RAILINGS', 'INTERIOR DOORS & HARDWARE', 'GYPSUM BOARD', 'TILE', 'STONE FLOORING', 'WOOD FLOORING', 'PAINTING', 'BATH ACCESSORIES', 'APPLIANCES', 'COUNTERTOPS']),
]
ROW_GROUP = {r: 'Barn Building Envelope' for r in find_rows('Custom Cupola', 'Copper Horse Weathervane', n=2, bld='barn')}   # the cupola and vane are exterior work
def group_of(r, l):
    if r in ROW_GROUP: return ROW_GROUP[r]
    for name, scope, b, trades in G_DEF:
        if l['bld'] == b and l['trade'] in trades: return name
    die(f'line has no group: {l["bld"]} {l["trade"]} {l["item"]}')
GROUPS = []
for name, scope, b, trades in G_DEF:
    rows = [(r, l) for r, l in LINES.items() if group_of(r, l) == name]
    GROUPS.append({'name': name, 'scope': scope, 'v': C(sum(l['v'] for r, l in rows)), 'lines': [{'item': clean(l['item']), 'v': l['v']} for r, l in rows]})
if sum(g['v'] for g in GROUPS) != DIRECT: die('groups do not tie to direct cost')

# ---------------------------------------------------------------- MasterFormat view = each workbook's Schedule of Values, row for row
SOV_NAMES = {'Sitework & Site Utilities': 'Site Work and Utilities', 'Concrete & Foundations': 'Concrete and Foundations', 'Cast-in-Place Concrete': 'Concrete',
             'Demolition & Waste Removal': 'Demolition and Waste Removal', 'Thermal & Sound Insulation': 'Insulation', 'Electrical, Low Voltage & Detection': 'Electrical and Low Voltage',
             'Drainage & Stormwater': 'Drainage', 'Audio Visual & Shades': 'Audio Visual and Shades'}
sov_name = lambda n: SOV_NAMES.get(n, n).replace(' & ', ' and ')
DIVS = []
for key, fn, label, off in BOOKS:
    sv, on, rows = WB[key]['SOV'], False, []
    for r in sv.iter_rows(values_only=True):
        if r[0] == 'CSI' and r[1] == 'Trade': on = True; continue
        if on and r[0] == 'SUBTOTAL': break
        if on and r[0]: rows.append((str(r[0]), r[1], r[2]))
    trades = [t for t in dict.fromkeys(l['trade'] for l in LINES.values() if l['bld'] == key)]
    by_trade = {}
    for r_, l in LINES.items():
        if l['bld'] == key: by_trade.setdefault(l['trade'], []).append(l)
    for codes, name, v in rows:
        tl = by_trade.get(name.upper(), [])
        if not tl: continue
        got = C(sum(l['v'] for l in tl))
        if got != C(v): die(f'SOV row {key} {name} {got} vs workbook {C(v)}')
        first = codes.split(',')[0].strip().replace('.', ' ')
        DIVS.append({'n': first, 'name': sov_name(name) if key != 'shared' else 'Shared General Requirements', 'b': label, 'v': got,
                     'lines': [{'item': clean(l['item']), 'code': l['code'].replace('.', ' ')} for l in tl]})
    if sum(d['v'] for d in DIVS if d['b'] == label) != BK[key]['direct']: die(f'SOV rows do not tie for {key}')
if sum(d['v'] for d in DIVS) != DIRECT: die('SOV rows do not tie to direct cost')

# ---------------------------------------------------------------- signature elements, by building, each x 1.15
def el(rows): return C(tot(rows) * D('1.15'))
fb = lambda *a, n=None: find_rows(*a, n=n, bld='barn'); fp = lambda *a, n=None: find_rows(*a, n=n, bld='ph')
ELS = [   # (key, building, name, value, spec)
    ('bifolds', 'barn', 'Bifolds and Transoms', el(fb('Marvin Custom Exterior Bifold', 'Install Bifold Door Units', 'LVL Headers @ Bifold', n=3)),
     'Four Marvin custom bifold units with transoms, 10 feet wide and 11 feet 8 inches high.'),
    ('deck', 'barn', 'Roof Deck', el(fb('Roof Deck Guardrail', 'Lounge Wing Roof Framing', 'Cumaru Decking', 'Integrated Drainage Mat', 'Pitched EPDM', 'EPDM @ Parapet', 'Metal Coping', 'Scuppers', 'Marvin Exterior Swing Door to Roof Deck', n=9)),
     'Cumaru over pitched EPDM, with the guardrail, the coping and the door to the deck.'),
    ('bar', 'barn', 'Whiskey Bar', el(fb('Custom Front Bar', 'Back Bar Cabinets', 'Copper Bar Top', 'Back Bar Countertop', n=4)), 'Front bar, back bar, copper bar top and stone back counter.'),
    ('canopy', 'barn', 'Canopies', el(fb('Canopies over Bifolds', 'Heavy Timber Brackets', n=2)), 'Three canopies over the bifold doors, carried on brackets.'),
    ('cupola', 'barn', 'Cupola and Vane', el(fb('Custom Cupola', 'Copper Horse Weathervane', n=2)), 'A glazed cupola on the ridge with a horse weathervane.'),
    ('stove', 'barn', 'Stove and Flue', el(fb('Hearthstone', n=1)), 'Hearthstone Bari wood stove with a Class A flue and hearth pad.'),
    ('bifold', 'ph', 'Bifold Doors and Bar Window', el(fp('Marvin Ultimate Bifold Door 2', 'Marvin Ultimate Bifold Door 3', 'Install Bifold Door Units', 'Marvin Ultimate Bifold Window 108', 'LVL/PSL Headers', n=5)),
     'Marvin Ultimate bifold doors across the south gable and on the west wall, and a 9 foot bifold window over the outside bar counter.'),
    ('win', 'ph', 'Marvin Ultimate Windows', el(fp('Marvin Ultimate Clad Casement', 'Marvin Ultimate Picture', 'Marvin Ultimate Shaped', 'Install Windows', 'Install Large & Bifold', 'Tempered Glazing', 'Glass Package Upgrade', n=9)),
     '30 white clad windows, including the glass gable on the south end and the shaped units in both gables.'),
    ('roof', 'ph', 'Cedar Roof and Copper', el(code_rows('07.50.00', bld='ph')), 'Cedar shingles to match the main house on a Cedar Breather mat, with copper drip edge, flashing, half-round gutters, leaders and the canopy at the north door.'),
    ('trim', 'ph', 'Mitered Claymark Trim', el(fp('Claymark', 'Cornice Returns', n=5)), 'Claymark Solid Gold casings, corner boards, pilasters, belt, frieze and soffits, with built-up cornice returns.'),
    ('stair', 'ph', 'White Oak Stair', el(fp('White Oak Stair', 'Stair Turns', 'Wood Guardrail', 'Wood Handrails', n=4)), 'From the entry landing, one run down to the bar and one up to the Gym and Yoga room, 16 risers with a wood guardrail at the stair opening.'),
    ('pbar', 'ph', 'Bar, Counters and Appliances', el(fp('Bar Island & Back Bar Cabinets', '1-1/4" Stone Countertops', '3" Stone Exterior Counter', 'Template, Fabricate', 'F/O 15"', 'F/O 24"', 'I/O Appliances', n=9)),
     'Painted inset Shaker bar cabinets with stone tops, the 3 inch stone counter outside the bar window and four undercounter appliances.'),
]
B_ELS = [e for e in ELS if e[1] == 'barn']; P_ELS = [e for e in ELS if e[1] == 'ph']
EL_T = {b: sum(e[3] for e in ELS if e[1] == b) for b in ('barn', 'ph')}
REST = {b: BK[b]['base'] - EL_T[b] for b in ('barn', 'ph')}
SITE_FDN = {'barn': el(trade_rows('barn', 'SITEWORK & SITE UTILITIES', 'CONCRETE & FOUNDATIONS') + fb('2" XPS')),   # below-grade insulation counts with the foundation
            'ph': el(trade_rows('ph', 'DEMOLITION & WASTE REMOVAL', 'CAST-IN-PLACE CONCRETE', 'FOUNDATION WATERPROOFING', 'SITEWORK', 'HARDSCAPING', 'SITE UTILITIES', 'DRAINAGE & STORMWATER'))}

# ---------------------------------------------------------------- allowances (from the Notes sheets), by building
A_ROWS = [
    ('Barn Bifold Doors and Glass', 'Four Marvin bifold units at $26,000.00 each, and the U-0.27 glass package', fb('Marvin Custom Exterior Bifold', 'Glass Package Upgrade', n=2)),
    ('Barn Roof Deck Guardrail', 'Type to be selected', fb('Roof Deck Guardrail', n=1)),
    ('House Electric Service Upgrade', 'To be set with the construction drawings', fb('House Electric Service Upgrade', n=1)),
    ('Barn Lighting Fixtures', 'Fixtures throughout', fb('Lighting Fixtures (Allowance)', n=1)),
    ('Propane Tank and Piping', 'New underground tank at the Barn', fb('New Propane Tank', n=1)),
    ('Barn Fixtures and Vanity', 'Kohler toilet, lav and bar sink with faucets, and the powder room vanity', fb('F/O Kohler', 'Powder Room Vanity', n=4)),
    ('Barn Back Bar Countertop', 'Stone, 31 SF', fb('Back Bar Countertop', n=1)),
    ('Barn Low Voltage', 'TV and data', fb('TV / Data', n=1)),
    ('Pool House Bifolds and Glass', 'Bifold doors 2 and 3, the bar window and the U-0.27 glass package', fp('Marvin Ultimate Bifold Door 2', 'Marvin Ultimate Bifold Door 3', 'Marvin Ultimate Bifold Window 108', 'Glass Package Upgrade', n=4)),
    ('Pool House Structure', 'Ridge beam and the steel frame at the south bifold wall', fp('LVL Ridge Beam', 'Steel Moment Frame', n=2)),
    ('Pool House Audio Visual and Shades', 'Sonos, Lutron, 26 motorized shades and the pre-wire', fp('Sonos', 'Lutron Lighting', 'Lutron Motorized', 'Low Voltage Pre-Wire', n=4)),
    ('Pool House Appliances', 'Ice maker, undercounter and beverage refrigerators, dishwasher', fp('F/O 15"', 'F/O 24"', n=4)),
    ('Pool House Stone Floor', 'To match the pool terrace, 415 SF installed', fp('Stone Floor to Match', n=1)),
    ('Pool House Countertops', 'Stone material, 75 SF', fp('1-1/4" Stone Countertops', '3" Stone Exterior Counter', n=2)),
    ('Pool House Recessed Lights', '27 four inch LED fixtures', fp('4" LED Recessed', n=1)),
    ('Pool House Doors and Hardware', 'Interior doors, door and cabinet hardware, the glass shower door', fp('Solid-Core 2-Panel', 'Door Hardware', 'Cabinet Hardware', 'Frameless Glass Shower Door', n=4)),
    ('Pool House Outdoor Shower', 'Location to be set', fp('Outdoor Shower', n=1)),
    ('Pool House Storage and Bath', 'Built-in storage shelving and bath accessories', fp('Built-In Storage Shelving', 'Bath Accessories Budget', n=2)),
    ('Barn Interior Doors and Hardware', 'Solid-core wood doors and door hardware', fb('Solid-Core Wood Interior Doors', 'Door Hardware', n=2)),
    ('Barn Utility Runs', 'Feeder, water and sanitary at 150 LF each, with trenching', fb('Feeder from House', 'Utility Trenching', 'Water Service from House', 'Sanitary Line to Existing Septic', n=4)),
    ('Pool House Utility Runs', 'Feeder, water and sanitary at 100 LF each, the gas trench from the Barn tank and roof drain piping', fp('Feeder from House', 'Gas Trench from Barn Propane', 'Water Service from House', 'Sanitary Line to Existing Septic', '4" Roof Drain Piping', n=5)),
    ('Pool House Standby Generator', '14 kW propane unit with automatic transfer switch', fp('14 kW Air-Cooled', n=1)),
    ('Pool House TV and Data', 'Low voltage', fp('TV / Data', n=1)),
    ('Existing Pavilion Utilities', 'Disconnect and cap', fp('Disconnect & Cap Existing Utilities', n=1)),
    ('Temporary Power', 'Both buildings', find_rows('Temporary Power', n=1, bld='shared')),
]
ALLOWS = [{'label': a, 'basis': b, 'v': money(tot(rows_))} for a, b, rows_ in A_ROWS]
ALLOW_TOTAL = sum(C(tot(rows_)) for a, b, rows_ in A_ROWS)
if len(set(r for a, b, rr in A_ROWS for r in rr)) != sum(len(rr) for a, b, rr in A_ROWS): die('an allowance row is counted twice')
RATES = [{'label': 'Pool House Tile', 'basis': 'Material', 'v': '$20.00 per SF'}, {'label': 'Pool House Countertops', 'basis': 'Material', 'v': '$150.00 per SF'},
         {'label': 'Pool House White Oak', 'basis': 'Material, site finished', 'v': '$13.00 per SF'}, {'label': 'Pool House Stone Floor', 'basis': 'Installed', 'v': '$42.00 per SF'},
         {'label': 'Barn Wood Flooring', 'basis': 'Material', 'v': '$13.00 per SF'}, {'label': 'Barn Powder Room Tile', 'basis': 'Material', 'v': '$10.00 per SF'}]

# ---------------------------------------------------------------- alternates, numbered across the property
ALT_SRC = {1: ('ph', 1), 2: ('barn', 1), 3: ('ph', 2), 4: ('barn', 2), 5: ('ph', 3)}
ALT_ROWS = {k: BK[b]['alts'][n] for k, (b, n) in ALT_SRC.items()}
ALTS = [{'id': str(k), 'title': P['alternates'][str(k)]['title'], 'note': P['alternates'][str(k)]['note'], 'v': ALT_ROWS[k]} for k in sorted(ALT_ROWS)]
if ALT_ROWS[5] != D('3450.00'): die('rock day rate moved')

# ---------------------------------------------------------------- selections (kit v1.7)
sys.path.insert(0, os.path.join(KIT, 'tools', 'sel')); import selbuild
SEL_PATH = os.path.join(CB, 'selections.json')
SELX = {'C': C, 'die': die, 'money': money, 'E': E, 'ws': ws, 'LINES': LINES, 'find_rows': lambda *a: find_rows(*a), 'code_rows': lambda *a: code_rows(*a),
        'ALT_ROWS': ALT_ROWS, 'VE_ROWS': {}, 'BASE': BASE, 'CB': CB, 'OP': D('0.15')}
SEL = selbuild.resolve(json.load(open(SEL_PATH)), SELX) if os.path.exists(SEL_PATH) else None

# ---------------------------------------------------------------- schedule, in weeks from mobilization, by phase
d = lambda s: dt.date.fromisoformat(s)
mob = d(SCH['mob'])
PHASES = [('Site and Foundation', ['mob', 'fdn', 'phdem', 'phfdn'], '#001a8a'), ('Structure', ['frm', 'roof', 'phfrm'], '#0028cc'),
          ('Envelope', ['win', 'rfg', 'sid', 'foam', 'phenv'], '#3d5bd9'), ('Systems', ['mep', 'phmep'], '#6a82e2'),
          ('Interiors', ['gwb', 'mill', 'pnt', 'flr', 'phins', 'phint'], '#93a5ec'), ('Closeout', ['trim', 'pun'], '#b9c5f3')]
phase_of = {k: (n, c) for n, ks, c in PHASES for k in ks}
SROWS = []
for r in SCH['rows']:
    n, c = phase_of[r['key']]
    SROWS.append({'name': r['name'], 'a': (d(r['s']) - mob).days / 7, 'b': ((d(r['f']) - mob).days + 1) / 7, 'phase': n, 'c': c})
SPAN = float(SCH['weeks'])
for r in SROWS: r['b'] = min(r['b'], SPAN)
PH_SPANS = [(n, c, min(r['a'] for r in SROWS if r['phase'] == n), max(r['b'] for r in SROWS if r['phase'] == n)) for n, ks, c in PHASES]
MS = [tuple(m) for m in SCH['milestones']]
if MS[-1][0] != SCH['weeks']: die(f"schedule does not end at {SCH['weeks']} weeks")
_gr = WB['shared']['Estimate']
if not any(r[1] and str(r[1]).startswith('Supervision') and r[2] == SCH['weeks'] for r in _gr.iter_rows(values_only=True)): die('shared supervision weeks differ from the schedule')

# ================================================================ model frames: two buildings, one frame size
MD = os.path.join(HERE, 'model')
FR_B = ['b_c0_site', 'b_c0b_backfill', 'b_c1_workshop', 'b_c2_lounge', 'b_c3_exterior', 'b_w1_bifolds', 'b_w2_deck', 'b_w3_bar', 'b_w4_canopy', 'b_w5_cupola', 'b_w6_stove', 'b_w7_all']
FR_P = ['p_c0_site', 'p_c0b_backfill', 'p_c1_first', 'p_c2_second', 'p_c3_exterior', 'p_w1_bifold', 'p_w2_win', 'p_w3_roof', 'p_w4_trim', 'p_w5_stair', 'p_w6_bar', 'p_w7_all']
FRAMES = FR_B + FR_P
def bbox(files):
    box = None
    for f in files:
        a = np.array(Image.open(os.path.join(MD, f + '.png')).split()[3]); ys, xs = np.where(a > 200); b = (xs.min(), ys.min(), xs.max(), ys.max())
        box = b if box is None else (min(box[0], b[0]), min(box[1], b[1]), max(box[2], b[2]), max(box[3], b[3]))
    return box
bb_ = bbox(FR_B)
CROPS = {'b': (max(0, int(bb_[0]) - 40), max(0, int(bb_[1]) - 40), min(1600, int(bb_[2]) + 90), min(1200, int(bb_[3]) + 50))}
FR = (CROPS['b'][2] - CROPS['b'][0], CROPS['b'][3] - CROPS['b'][1])
pb = bbox(FR_P); cx, cy = (pb[0] + pb[2]) / 2 + 25, (pb[1] + pb[3]) / 2 + 5   # the pool house frames: the same size, centred on the pool house
if pb[2] - pb[0] > FR[0] - 40 or pb[3] - pb[1] > FR[1] - 30: die('pool house frames are larger than the barn frame size')
x0 = int(min(max(0, cx - FR[0] / 2), 1600 - FR[0])); y0 = int(min(max(0, cy - FR[1] / 2), 1200 - FR[1]))
CROPS['p'] = (x0, y0, x0 + FR[0], y0 + FR[1])
crop_of = lambda f: CROPS[f[0]]
PTS = {}
for pre, keys in (('b', [e[0] for e in B_ELS]), ('p', [e[0] for e in P_ELS])):
    raw = json.load(open(os.path.join(MD, f'{pre}_c3_exterior.png.points.json'))); cr = CROPS[pre]
    for k in keys:
        rk = {'pbar': 'bar'}.get(k, k) if pre == 'p' else k
        PTS[k] = [raw[rk][0] - cr[0], raw[rk][1] - cr[1]]
BB = {}
for (k, *_r), f in zip(B_ELS + P_ELS, FR_B[5:11] + FR_P[5:11]):
    im = np.array(Image.open(os.path.join(MD, f + '.png')).convert('RGBA')).astype(int)
    blue = (im[..., 2] > im[..., 0] + 70) & (im[..., 2] > im[..., 1] + 50) & (im[..., 3] > 200)
    ys, xs = np.where(blue); cr = crop_of(f)
    if not len(xs): die('no blue in ' + f)
    BB[k] = [int(xs.min()) - cr[0], int(ys.min()) - cr[1], int(xs.max()) - cr[0], int(ys.max()) - cr[1]]

# ================================================================ copy
LOC = 'North Salem, NY'
PNAME = 'The Cooper Barn and Pool House'
PROSE = [
    'A workshop barn and a pool house on one property at 194 Titicus Road in North Salem, New York, designed by Crisp Architects. The two buildings go up together, on one 26 week schedule with one site team.',
    'The barn is a 40 by 30 foot workshop with a 16 by 26 foot lounge wing, framed in 2x6 at a 15 foot plate under a galvalume standing seam roof, with four Marvin bifold doors. The pool house is 35 feet 6 inches by 16 feet over two floors, in 6 inch smooth Hardie lap siding and mitered Claymark trim under a cedar shingle roof with copper gutters.',
    'Inside the barn, exposed beams cross the double-height workshop, and the lounge has a copper top whiskey bar and a Hearthstone wood stove. The pool house opens to the terrace through Marvin Ultimate bifold doors, with a bar, a bath and storage on a stone floor. Upstairs, the white oak Gym and Yoga room has a glazed south gable.',
]
STEPS = [
    {'k': 'b_site', 'bld': 'barn', 'rail': 0, 'title': 'Site and Foundation', 'sub': 'The Barn', 'fig': money(SITE_FDN['barn']), 'cap': 'Barn site work, foundations and below-grade insulation, with overhead and profit.',
     'body': 'Excavation for 4 foot frost walls and a full-height mechanical space under the lounge wing. 12x24 footings, parged 10 inch walls and 4 inch slabs on 6 inch gravel, with 2 inch XPS below grade.',
     'img': 0, 'via': 1, 'note': 'The Barn, Excavation, Footings and Frost Walls'},
    {'k': 'b_shop', 'bld': 'barn', 'rail': 1, 'title': 'Workshop', 'sub': 'The Barn', 'body': 'A double-height workshop under a standing seam roof, with exposed beams, a cupola and a mezzanine. Workbenches, desk areas and a 3D printer desk are part of the budget.', 'img': 2, 'note': 'The Barn, Workshop'},
    {'k': 'b_lounge', 'bld': 'barn', 'rail': 1, 'title': 'Lounge', 'sub': 'The Barn', 'body': 'A whiskey bar and lounge with wainscot walls, reclaimed beams, a copper top bar and a Hearthstone wood stove. Above it, a Cumaru roof deck over pitched EPDM.', 'img': 3, 'note': 'The Barn, Lounge Wing, Walls Cut Away'},
    {'k': 'b_ext', 'bld': 'barn', 'rail': 1, 'title': 'Exterior', 'sub': 'The Barn', 'body': 'Hardie board and batten, ZIP sheathing with a rain screen, spray foam insulation and a galvalume standing seam roof over high temperature ice and water shield.', 'img': 4, 'note': 'The Barn, Complete Building'},
]
nB, nP = len(B_ELS), len(P_ELS)
for i, (k, b, name, v, spec) in enumerate(B_ELS):
    STEPS.append({'k': k, 'bld': 'barn', 'rail': 2, 'title': 'Barn Elements', 'sub': f'<b>{i + 1:02d} / {nB:02d}</b>{E(name)}', 'fig': money(v), 'cap': spec, 'img': 5 + i, 'el': i, 'no': i + 1, 'note': 'Shown in Montana Blue'})
STEPS.append({'k': 'b_all', 'bld': 'barn', 'rail': 2, 'title': 'Barn Elements', 'sub': 'Six Elements', 'fig': money(EL_T['barn']), 'cap': f'{pct(EL_T["barn"], BK["barn"]["base"])} of the {money(BK["barn"]["base"])} Barn budget.', 'img': 11, 'all': True, 'note': 'The Barn, Signature Elements in Montana Blue'})
STEPS += [
    {'k': 'p_site', 'bld': 'ph', 'rail': 3, 'title': 'Site and Foundation', 'sub': 'The Pool House', 'fig': money(SITE_FDN['ph']), 'cap': 'Pool House demolition, site work, utilities and foundations, with overhead and profit.',
     'body': 'The existing pavilion comes down. 10 inch frost walls 48 inches deep on 24 inch footings, with the east and north walls carried up against the higher grade, and a 4 inch slab on R-10 rigid and 6 inch gravel.',
     'img': 12, 'via': 13, 'note': 'The Pool House, Excavation, Footings and Frost Walls'},
    {'k': 'p_first', 'bld': 'ph', 'rail': 4, 'title': 'First Floor', 'sub': 'The Pool House', 'body': 'The bar, a bath and storage on a stone floor, open to the terrace through bifold doors on the south and west and a bifold window over an outside bar counter.', 'img': 14, 'note': 'The Pool House, First Floor'},
    {'k': 'p_second', 'bld': 'ph', 'rail': 4, 'title': 'Second Floor', 'sub': 'The Pool House', 'body': 'A Gym and Yoga room in white oak under a vaulted ceiling, with Marvin Ultimate casements along both long walls and a glass gable to the south.', 'img': 15, 'note': 'The Pool House, Second Floor'},
    {'k': 'p_ext', 'bld': 'ph', 'rail': 4, 'title': 'Exterior', 'sub': 'The Pool House', 'body': 'A cedar shingle roof with copper gutters and leaders, 6 inch smooth Hardie lap siding, mitered Claymark trim, and bluestone steps with stone veneer at the north entry.', 'img': 16, 'note': 'The Pool House, Complete Building'},
]
for i, (k, b, name, v, spec) in enumerate(P_ELS):
    STEPS.append({'k': k, 'bld': 'ph', 'rail': 5, 'title': 'Pool House Elements', 'sub': f'<b>{i + 1:02d} / {nP:02d}</b>{E(name)}', 'fig': money(v), 'cap': spec, 'img': 17 + i, 'el': nB + i, 'no': i + 1, 'note': 'Shown in Montana Blue'})
STEPS.append({'k': 'p_all', 'bld': 'ph', 'rail': 5, 'title': 'Pool House Elements', 'sub': 'Six Elements', 'fig': money(EL_T['ph']), 'cap': f'{pct(EL_T["ph"], BK["ph"]["base"])} of the {money(BK["ph"]["base"])} Pool House budget.', 'img': 23, 'all': True, 'note': 'The Pool House, Signature Elements in Montana Blue'})
for s in STEPS: s['label'] = '02 / Design'
RAIL = [(0, 'Barn Site'), (1, 'Barn'), (2, 'Barn Elements'), (3, 'Pool House Site'), (4, 'Pool House'), (5, 'Pool House Elements')]

BASIS = ['Crisp Architects SK-5 First and Second Floor Plans dated September 22, 2026 (Option 1) and the Cooper Barn Budget email of October 9, 2026, for the Barn.',
         'Crisp Architects Cooper Pool House sketches A1.1 and A1.2 (Pool House SK4) dated October 14, 2026, for the Pool House.',
         'Structural, mechanical and site drawings are not yet issued. Final framing and systems follow those drawings.']
CLAR = ['The Barn and the Pool House are built in parallel. General Requirements are carried once for both buildings, on one 26 week duration with one superintendent and one project manager.',
        'If only one building is awarded, General Requirements for that building alone are $100,303.00 for the Barn or $88,027.21 for the Pool House, with overhead and profit.',
        'Both buildings are carried to the 2025 Energy Conservation Construction Code of New York State, each with a blower door test and an ERV.',
        'Heating, cooling and hot water are all electric. The Pool House standby generator runs on propane from the Barn tank.',
        'Four Marvin custom exterior bifold doors with transoms at the Barn, carried at a $26,000.00 supply allowance each plus installation.',
        'Pool House windows and doors are Marvin Ultimate clad units, white clad exterior and white primed interior, Low-E argon, with Reeb Select storage doors.',
        'The Pool House includes demolition and removal of the existing pool house pavilion, with protection of the pool, stone terrace and existing retaining wall.',
        'Pool House plumbing fixtures, decorative light fixtures, sconces and pendants are furnished by owner and installed by Montana.',
        'The Pool House includes Sonos audio, Lutron lighting control and 26 motorized shades at the bar and Gym/Yoga, with low voltage pre-wire.',
        'House electric service upgrade carried as an $8,925.00 allowance. Electric feeder, water service and sanitary line carried at 150 LF each at the Barn and 100 LF each at the Pool House.',
        'Allowance amounts are carried before overhead and profit. Overhead and profit of 15% applies to the base budget and the alternates.',
        'None of the five alternates is in the base budget.']
EXCL = ['Building permit, filing and utility fees, and all planning, zoning and health department fees and approvals.',
        'Architectural, structural, MEP and civil engineering and surveys. Building stakeout is priced as an alternate for each building.',
        'Septic design and testing, and any septic system upgrade required by the Westchester County Department of Health.',
        'Rock removal, priced per day as Alternate 5. Dewatering and unsuitable soils.',
        'Temporary heat, snow removal and winter protection.',
        'Landscaping, topsoil, seed, driveway and paving. Pool barrier fencing, gates and pool door alarms.',
        'Workshop equipment is furnished and set by the owner. Dedicated 240V circuits are included.',
        'Dust collection ductwork and collector at the Barn.', 'SWPPP preparation and inspections.', 'Furniture, TVs and owner equipment.']
notes_src = ' '.join(str(r[2]) for key in WB for r in WB[key]['Notes & Clarifications'].iter_rows(min_row=5, values_only=True) if r[2]) + ' ' + ' '.join(l['item'] for l in LINES.values())
for frag in ['SK-5 First & Second Floor Plans dated 22 September 2026 (Option 1)', 'Cooper Barn Budget email of 9 October 2026', 'A1.1 and A1.2 (Pool House SK4) dated 10/14/2026',
             'No structural, MEP or site drawings issued', 'built in parallel', 'one superintendent and one project manager', 'Barn $100,303.00 or Pool House $88,027.21',
             '2025 Energy Conservation Construction Code of New York State', 'a blower door test and an ERV', 'Heating, cooling and hot water are all electric', 'fed from the Barn propane tank',
             'Four Marvin custom exterior bifold doors with transoms', '$26,000 supply allowance each plus installation', 'Marvin Ultimate clad units, white clad exterior and white primed interior, Low-E argon',
             'Reeb Select storage doors', 'demolition and removal of the existing pool house pavilion', 'protection of the pool, stone terrace and existing retaining wall',
             'furnished by owner and installed by Montana', 'Sonos audio, Lutron lighting control and 26 motorized shades', 'House electric service upgrade carried as an $8,925 allowance',
             'each carried at 150 LF', '100 LF runs each', 'the amounts carried on the Estimate sheet, before Overhead & Profit', 'Overhead & Profit of 15% applies to base and alternates',
             'Building permit, health department and utility fees', 'planning, zoning and health department approvals', 'Architectural, structural, MEP and civil engineering and surveys',
             'Building stakeout and foundation as-built survey is Add Alternate 1', 'Building stakeout and foundation as-built survey is Alternate 2', 'Westchester County Department of Health',
             'Rock hammering and removal', 'unsuitable soils', 'dewatering', 'Temporary heat, snow removal and winter protection', 'Landscaping, driveway and paving are excluded',
             'Pool barrier fencing, gates and pool door alarms', 'Workshop equipment is furnished and set by the owner. Dedicated 240V circuits are included.',
             '2" XPS', 'Hearthstone Bari', 'Dust collection ductwork', 'SWPPP preparation and inspections', 'furniture, TV and equipment', 'Building permit, filing and utility fees']:
    if frag not in notes_src: die('note drifted from the estimate: ' + frag)

TEAM = [('Joe Montana', 'Founder and CEO', 'joe-montana-sr-800.webp'), ('Matt Nicosia', 'COO', 'matt-nicosia-800.webp'), ('Chris Lange', 'Director of Construction', 'chris-lange-800.webp')]
TIERS = [
    ('1', 'Site Superintendent', 'On site, start to finish', ['Runs both buildings from mobilization to the punch list.',
        'Checks the work against the drawings before it is covered.', 'Files daily field reports and photos in Montana OS.']),
    ('1', 'Project Manager', 'On every project', ['Runs the schedule, the budget and the submittals with Crisp Architects.',
        'Prices every change order before the work is approved.', 'Orders the Marvin bifold doors and windows at award.']),
    ('3', 'Principals', 'On every project', ['Joe Montana, Matt Nicosia and Chris Lange stay on every project.',
        'One of them is your contact from the first walkthrough to the final punch list.', 'You are never handed off.']),
]
OS_FEATURES = [('RFIs, Submittals and Change Orders', 'One log. Every item dated and tracked to an answer.'), ('Schedule', 'Roadblocks flagged the day they appear.'),
               ('Field Reports', 'Daily reports and photos from the site.'), ('Drawings and Specs', 'Indexed by sheet and section.'),
               ('Job Cost', 'Tracked against the budget, line by line.'), ('Punch List', 'Tracked to closeout.')]
OS_FOR = [('For Crisp Architects', 'Every question, answer and decision in one place.'), ('For the Owners', 'The budget, the schedule and both sites, live.'),
          ('For the Job', 'The field and the office work from the same information.')]
PHONE, ADDR, EMAIL = '(845) 398-1778', '173 N. Route 9W, Congers, NY 10920', P['decisions']['email']

PC_PHASES = [('Budget Setting and Schematic Design Budget', 'Phase 01', 20, 'Your target budget and the scope it buys', 'About 2 weeks', 'Budget Set'),
             ('Design Development Updates', 'Phase 02', 15, 'Selections in place of allowances', 'About 2 weeks per revision', 'Selections'),
             ('Firm Trade Partner Pricing', 'Phase 03', 8, 'Bids from the trades who will build it', '3 to 4 weeks with Phase 04', 'Trade Bids'),
             ('Value Engineering, Built In', 'Phase 04', 3, 'One Guaranteed Maximum Price', 'GMP delivered', 'GMP')]
PC_CARDS = [('Phase 01', 'Budget Setting and Schematic Design Budget', ['Budget setting session with you', 'Survey and soils engineer coordinated', 'Foundation strategy, priced', 'Options priced side by side'],
             'Your target budget and an options menu with three scenarios.'),
            ('Phase 02', 'Design Development Updates', ['Re-priced as the design develops', 'Design-to-budget meetings with your architect', 'Early flags on anything drifting over budget'],
             'An updated budget with more detailed line items, plus a preliminary schedule.'),
            ('Phase 03', 'Firm Trade Partner Pricing', ['Major trades bid to our subcontractors and vendors', 'Site walks with the key trades', 'Bids compared line by line'],
             'A budget backed by real pricing.'),
            ('Phase 04', 'Value Engineering, Built In', ['Money-saving options with trade-offs explained', 'Final pricing', 'Guaranteed Maximum Price proposal'],
             'Your GMP and a construction schedule.')]
PC_PROMISES = [('Walk-Away Guarantee', 'Stop at the end of any phase. You pay only for the phases started and keep everything delivered.'),
               ('Unused Reserve Comes Back to You', '100% of unused contingency is credited back. It is used only with written notice.'),
               ('One Rate, No Matter Who Is on the Call', 'One hourly rate, billed for one person per meeting.')]
PC_PRINCIPLES = [('You Are Heard From Day One', 'Your vision, your budget, your timeline. We listen before we estimate.'),
                 ('You Understand Every Number', 'We walk you through each line item so you know exactly what drives cost and where tradeoffs live.'),
                 ('You Own the Strategy', 'We produce options, with the price of each one, so the plan at GMP is yours.')]

# ================================================================ assets
SITE, NET = os.path.join(HERE, 'site'), os.path.join(HERE, 'netlify')
for p_ in (SITE, NET): shutil.rmtree(p_, ignore_errors=True)
A = os.path.join(SITE, 'assets'); [os.makedirs(os.path.join(A, s), exist_ok=True) for s in ('m', 'team')]
PDFS = sorted(glob.glob(os.path.join(CB, 'proposal', '*.pdf')))
if len(PDFS) > 1: die('keep one proposal PDF in proposal/')
PDF = os.path.basename(PDFS[0]) if PDFS else None
if PDF: shutil.copy(PDFS[0], os.path.join(A, PDF))
CACHE = os.path.join(HERE, '.cache'); os.makedirs(os.path.join(CACHE, 'm'), exist_ok=True)
def cached(src, dst):
    c = os.path.join(CACHE, dst)
    if os.path.exists(c) and os.path.getmtime(c) > os.path.getmtime(src): shutil.copy(c, os.path.join(A, dst)); return True
    return False
def keep(dst): shutil.copy(os.path.join(A, dst), os.path.join(CACHE, dst))
def jpg(src, dst, width=None, q=82):
    if cached(src, dst): return
    im = Image.open(src).convert('RGB')
    if width and im.width > width: im = im.resize((width, round(im.height * width / im.width)), Image.LANCZOS)
    im.save(os.path.join(A, dst), 'JPEG', quality=q, optimize=True, progressive=True); keep(dst)
H = os.path.join(CB, 'hero')   # Montana Barn renders for the hero, the gallery and the close; Crisp Architects renderings of the Pool House
jpg(os.path.join(H, 'p_se_r.png'), 'hero.jpg', q=84); jpg(os.path.join(H, 'p_se_p.png'), 'hero_p.jpg', q=82)
GAL = [('g1.jpg', 'hero_front.png', 'The Barn, South Elevation'), ('g2.jpg', 'hero_se.png', 'The Barn, Southeast'), ('g3.jpg', 'crisp_ph_ext.png', 'The Pool House, Southwest'),
       ('g4.jpg', 'crisp_ph_gym.png', 'The Pool House, Gym and Yoga Room'), ('g5.jpg', 'hero_dusk.png', 'The Barn at Dusk')]
for n, f, c in GAL: jpg(os.path.join(H, f), n, width=1760, q=80)
jpg(os.path.join(H, 'p_dusk_d.png'), 'dusk.jpg', q=80); jpg(os.path.join(H, 'p_dusk_p.png'), 'dusk_p.jpg', q=80)
def edge_clean(im):
    a = np.asarray(im).copy(); al = a[..., 3].astype(float); al[al < 8] = 0
    h, w = al.shape
    rx = np.minimum(np.clip(np.arange(w) / (0.05 * w), 0, 1), np.clip((w - 1 - np.arange(w)) / (0.10 * w), 0, 1))
    ry = np.minimum(np.clip(np.arange(h) / (0.05 * h), 0, 1), np.clip((h - 1 - np.arange(h)) / (0.07 * h), 0, 1))
    f = (ry[:, None] * rx[None, :]) ** 0.8
    al = np.where(al >= 250, al, al * f); a[..., 3] = al.astype(np.uint8)
    return Image.fromarray(a, 'RGBA')
for f in FRAMES:
    src = os.path.join(MD, f + '.png')
    if not cached(src, 'm/' + f + '.webp'):
        edge_clean(Image.open(src).convert('RGBA').crop(crop_of(f))).save(os.path.join(A, 'm', f + '.webp'), 'WEBP', quality=88, method=6); keep('m/' + f + '.webp')
for _n, _t, ph in TEAM: shutil.copy(os.path.join(KIT, 'assets', 'img', ph), os.path.join(A, 'team', ph))
IMG = os.path.join(KIT, 'assets', 'img')
shutil.copy(os.path.join(IMG, 'mc-wordmark-rev.png'), os.path.join(A, 'mc-white.png'))
shutil.copy(os.path.join(IMG, 'mc-wordmark.png'), os.path.join(A, 'mc-ink.png'))
wm = Image.open(os.path.join(IMG, 'mc-wordmark-rev.png')).convert('RGBA')
Image.merge('RGBA', (Image.new('L', wm.size, 255),) * 3 + (wm.split()[3],)).save(os.path.join(A, 'mc-white-all.png'), optimize=True)
if SEL: SELX.update(A=A, HERO=os.path.join(H, 'p_se_r.png')); selbuild.assets(SEL, SELX)
font = lambda f: base64.b64encode(open(os.path.join(KIT, 'assets', 'fonts', f), 'rb').read()).decode()

M_PIECES = ['0,0 0,652 183,652 183,275 380,488 515,350 196,0', '773,124 586,310 586,652 773,652', '773,0 571,0 399,175 533,322 773,82']
def mark(cls, fill='currentColor'):
    return f'<svg class="{cls}" viewBox="0 0 773 652" aria-hidden="true">' + ''.join(f'<polygon class="mp mp{k}" points="{p}" fill="{fill}"/>' for k, p in enumerate(M_PIECES)) + '</svg>'
FAVICON = 'data:image/svg+xml;base64,' + base64.b64encode(f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="-60 -120 893 893"><rect x="-60" y="-120" width="893" height="893" rx="160" fill="#0028cc"/>{"".join(f"<polygon points=\"{p}\" fill=\"#fff\"/>" for p in M_PIECES)}</svg>'.encode()).decode()

# ================================================================ markup
DLICON = '<svg viewBox="0 0 16 16" fill="none" aria-hidden="true"><path d="M8 2v8.5M4.5 7 8 10.5 11.5 7M2.5 13.5h11" stroke="currentColor" stroke-width="1.4"/></svg>'
def pdf_link(cls):
    return f'<a class="{cls}" href="assets/{E(PDF)}" download="{E(PDF)}" data-pdf>{DLICON}<span>Proposal PDF</span></a>' if PDF else ''
ARROW = '<svg viewBox="0 0 16 16" fill="none" aria-hidden="true"><path d="M2 8h11M9 3.5 13.5 8 9 12.5" stroke="currentColor" stroke-width="1.4"/></svg>'
PLUS = '<svg class="pl" viewBox="0 0 16 16" fill="none" aria-hidden="true"><path d="M8 2v12M2 8h12" stroke="currentColor" stroke-width="1.4"/></svg>'
NAV = [('project', 'Project'), ('design', 'Design'), ('budget', 'Budget')] + ([('selections', 'Selections')] if SEL else []) + [('schedule', 'Schedule'), ('approach', 'Approach')]
kick = lambda t, cls='': f'<p class="kick {cls}"><i></i>{E(t)}</p>'
def odo(v, i=''):
    return f'<span class="odo num"{f" id={chr(34)}{i}{chr(34)}" if i else ""} data-v="{E(v)}" role="img" aria-label="{E(v)}">{E(v)}</span>'
def ramp(n, a=(0, 26, 138), b=(207, 216, 247)):
    return ['#%02x%02x%02x' % tuple(round(a[k] + (b[k] - a[k]) * (i / max(1, n - 1))) for k in range(3)) for i in range(n)]
OPC = '#9aa1b2'

s0 = STEPS[0]
stage_imgs = ''.join(f'<img src="assets/m/{f}.webp" alt="" decoding="async"{" loading=lazy" if i > 3 else ""}>' for i, f in enumerate(FRAMES))
nav_links = ''.join(f'<a href="#{i}">{E(t)}</a>' for i, t in NAV)
menu_links = ''.join(f'<a href="#{i}"><small>{k + 1:02d}</small>{E(t)}</a>' for k, (i, t) in enumerate(NAV))
def elist(b, els, lab):
    rows = ''.join(f'<button type="button"><span>{i + 1:02d}</span><span>{E(n)}</span><span class="num">{money(v)}</span></button>' for i, (k, bb, n, v, s) in enumerate(els))
    rows += f'<div class="elist__rest"><span></span><span>{lab} Building, Site and Systems</span><span class="num">{money(REST[b])}</span></div>'
    return f'<div class="elist hide" data-b="{b}"><p class="elist__h">{lab} Budget with Overhead and Profit</p>{rows}<div class="elist__tot"><span></span><span>{lab} Budget</span><span class="num">{money(BK[b]["base"])}</span></div></div>'
rail = ''.join(f'<button type="button" data-step="{j}"{" class=on" if j == 0 else ""}>{E(t)}</button>' for j, t in RAIL)

def view_rows(items, colors):
    out = ''
    for k, (it, c) in enumerate(zip(items, colors)):
        lines = ''.join(f'<li>{"<span class=\"cd num\">" + E(ln["code"]) + "</span>" if ln.get("code") else ""}<span>{E(ln["item"])}</span></li>' for ln in it['lines'])
        scope = f'<p>{E(it["scope"])}</p>' if it.get('scope') else ''
        head = f'<span class="sq"></span><div class="brow__t"><h3>{E(it["head"])}</h3>{scope}</div><span class="v num">{money(it["v"])}</span><span class="pct num">{pct(it["v"], BASE)}</span>'
        out += (f'<div class="brow" style="--c:{c}" data-k="{k}"><button type="button" class="brow__h" aria-expanded="false">{head}{PLUS}</button><ul class="brow__l" hidden>{lines}</ul></div>')
    return out
gcol = ramp(len(GROUPS))
dcol = ramp(len(DIVS))
EL_ITEMS = [{'head': g['name'], 'scope': g['scope'], 'v': g['v'], 'lines': g['lines']} for g in GROUPS]
MF_ITEMS = [{'head': f'{dv["n"]}  {dv["name"]}', 'scope': dv['b'], 'v': dv['v'], 'lines': dv['lines']} for dv in DIVS]
def segs(items, colors):
    return ''.join(f'<i style="--w:{float(it["v"] / BASE * 100):.3f}%;--c:{c};--k:{k}" title="{E(it["head"])}"></i>' for k, (it, c) in enumerate(zip(items, colors))) + \
        f'<i style="--w:{float(OP / BASE * 100):.3f}%;--c:{OPC};--k:{len(items)}" title="Overhead and Profit"></i>'

g_ms = ''.join(f'<div style="left:{w / SPAN * 100:.3f}%"><b>Week {w}</b><span>{E(l)}</span></div>' for w, l in MS)
g_rows = ''.join(f'<div class="g-row" data-a="{r["a"]:.3f}" data-b="{r["b"]:.3f}"><div class="g-name">{E(r["name"])}</div><div class="g-track"><i class="g-bar" style="left:{r["a"] / SPAN * 100:.3f}%;width:{(r["b"] - r["a"]) / SPAN * 100:.3f}%;--c:{r["c"]}"><b class="g-fill"></b></i></div></div>' for r in SROWS)
g_legend = ''.join(f'<li style="--c:{c}" data-a="{a:.3f}" data-b="{b:.3f}"><i></i><b>{E(n)}</b><span>Weeks {a:.0f} to {b:.0f}</span><em aria-hidden="true"></em></li>' for n, c, a, b in PH_SPANS)
_band = 8 if SPAN > 30 else 5
g_grid = ''.join(f'<u style="left:{a / SPAN * 100:.3f}%;width:{(min(a + _band, SPAN) - a) / SPAN * 100:.3f}%"></u>' for a in range(_band, int(SPAN), 2 * _band)) + \
    ''.join(f'<i style="left:{w / SPAN * 100:.3f}%"></i>' for w in range(0, int(SPAN) + 1, 2 if SPAN > 30 else 1))
_ticks = sorted(set(list(range(0, int(SPAN), _band)) + [int(SPAN)]))
g_axis = '<i></i><div>' + ''.join(f'<span style="left:{t / SPAN * 100:.3f}%">{"Week " if t == 0 else ""}{t}</span>' for t in _ticks) + '</div>'

_al = sorted(zip(ALLOWS, [tot(rr) for a_, b_, rr in A_ROWS]), key=lambda t: -t[1]); _amax = _al[0][1]   # ranked ledger, a bar per allowance
allow_rows = ''.join(f'<div class="aled__row" style="--w:{float(v / _amax * 100):.2f}%;--k:{k}"><span class="aled__n num">{k + 1:02d}</span><div class="aled__t"><h3>{E(a["label"])}</h3><p>{E(a["basis"])}</p></div><div class="aled__bar" aria-hidden="true"><i></i></div><span class="aled__v num">{E(a["v"])}</span></div>' for k, (a, v) in enumerate(_al))
rate_cells = ''.join(f'<div class="arate__c"><div class="arate__t"><h3>{E(a["label"])}</h3><p>{E(a["basis"])}</p></div><span class="arate__v num">{E(a["v"].split(" per ")[0])}<small>per {E(a["v"].split(" per ")[1])}</small></span></div>' for a in RATES)
alts = ''.join(f'<div class="alt"><div><h3>{E(a["title"])}</h3><p>{E(a["note"])}</p></div><span class="v num">{money(a["v"])}</span><button type="button" class="tog" role="switch" aria-checked="false" data-id="{a["id"]}" aria-label="Add {E(re.sub(r"^[0-9]+[.] ", "", a["title"]))}"></button></div>' for a in ALTS)
ncol = lambda h, items: f'<div class="ncol"><h3>{E(h)}</h3><ol>{"".join(f"<li>{E(x)}</li>" for x in items)}</ol></div>'
tiers = ''.join(f'<article class="tier" data-r="up"><b class="tier__n num">{n}</b><h3>{E(h)}</h3><p class="tier__sub">{E(sub)}</p><ul>{"".join(f"<li>{E(x)}</li>" for x in xs)}</ul></article>' for n, h, sub, xs in TIERS)
people = ''.join(f'<figure class="person"><div class="person__img"><img src="assets/team/{ph}" alt="{E(n)}" loading="lazy" decoding="async"></div><figcaption><b>{E(n)}</b><span>{E(t)}</span></figcaption></figure>' for n, t, ph in TEAM)
os_feat = ''.join(f'<li><span class="osx__no">{k + 1:02d}</span><b>{E(a)}</b><em>{E(b)}</em></li>' for k, (a, b) in enumerate(OS_FEATURES))
OS_LETTERS = ('<svg class="osx__os" viewBox="-4 -4 548 258" aria-hidden="true"><g fill="none" stroke="currentColor" stroke-width="2.5" stroke-linejoin="miter">'
              '<polygon points="48,0 202,0 250,48 250,202 202,250 48,250 0,202 0,48"/><polygon points="92,62 158,62 180,84 180,166 158,188 92,188 70,166 70,84"/>'
              '<polygon points="342,0 540,0 540,62 378,62 366,74 366,92 378,104 488,104 540,156 540,198 488,250 290,250 290,188 452,188 464,176 464,158 452,146 342,146 290,94 290,52"/></g></svg>')
OS_TILES = [('Base Budget', money(BASE), 'Job cost tracked against it, line by line.', True),
            ('Line Items', f'{len(LINES)}', 'Every line of both buildings, in one place.', False),
            ('Schedule', f'{SCH["weeks"]} Weeks', 'From mobilization. Roadblocks flagged the day they appear.', False)]
if SEL: OS_TILES.append(('Selections', f'{len(SEL["elements"])}', 'Each one requested and confirmed in writing.', False))
os_tiles = ''.join(f'<div class="osx__tile{" osx__tile--navy" if hi else ""}"><span class="osx__tk">{E(k)}</span><b class="num">{E(v)}</b><p>{E(d_)}</p><em aria-hidden="true">{i + 1:02d}</em></div>' for i, (k, v, d_, hi) in enumerate(OS_TILES))
os_for = ''.join(f'<div><p>{E(a)}</p><h4>{E(b)}</h4></div>' for a, b in OS_FOR)

def funnel(cls='funnel'):
    W_, H_, mid, n = 1000, 300, 150, len(PC_PHASES)
    cx_ = [W_ * (i + .5) / n for i in range(n)]; hw = [p[2] / PC_PHASES[0][2] * 120 for p in PC_PHASES]
    xs, ws_ = [20] + cx_ + [W_ - 20], [hw[0]] + hw + [hw[-1]]
    top = f'M{xs[0]},{mid - ws_[0]}' + ''.join(f' C{(xs[i - 1] + xs[i]) / 2},{mid - ws_[i - 1]} {(xs[i - 1] + xs[i]) / 2},{mid - ws_[i]} {xs[i]},{mid - ws_[i]}' for i in range(1, len(xs)))
    bot = f'L{xs[-1]},{mid + ws_[-1]}' + ''.join(f' C{(xs[i] + xs[i - 1]) / 2},{mid + ws_[i]} {(xs[i] + xs[i - 1]) / 2},{mid + ws_[i - 1]} {xs[i - 1]},{mid + ws_[i - 1]}' for i in range(len(xs) - 1, 0, -1)) + 'Z'
    lines = ''.join(f'<line x1="{x}" y1="{mid - hw[i] - 10}" x2="{x}" y2="{mid + hw[i] + 10}" class="fz-tick"/>' for i, x in enumerate(cx_))
    return f'<svg class="{cls}" viewBox="0 0 {W_} {H_}" role="img" aria-label="The price gets more certain across four phases and ends in a Guaranteed Maximum Price"><path class="fz-band" d="{top} {bot}"/><line x1="{xs[0]}" y1="{mid}" x2="{xs[-1]}" y2="{mid}" class="fz-mid"/>{lines}<circle cx="{xs[0]}" cy="{mid}" r="7" class="fz-dot"/></svg>'
pc_cols = ''.join(f'<div class="pcph"><span class="pcph__k">{E(k)}</span><h4>{E(n)}</h4><p class="pcph__lab">Locks In</p><p class="pcph__lock">{E(lk)}</p><span class="pcph__t">{E(t)}</span></div>' for n, k, w, lk, t, short in PC_PHASES)
pc_cards = ''.join(f'<article class="pccard"><div class="pccard__top"><span>{E(k)}</span><span>{E(next(p[4] for p in PC_PHASES if p[1] == k))}</span></div><h4>{E(t)}</h4><ul>{"".join(f"<li>{E(x)}</li>" for x in xs)}</ul><p class="pccard__get"><span>You Get</span>{E(g)}</p></article>' for k, t, xs, g in PC_CARDS)
PCF_LABS = ''.join(f'<span><b>{E(short)}</b>{E(k)}</span>' for n, k, w, lk, t, short in PC_PHASES)
pc_prom = ''.join(f'<div class="pcprom"><h4>{E(a)}</h4><p>{E(b)}</p></div>' for a, b in PC_PROMISES)
pc_prin = ''.join(f'<div class="pcprin"><h4>{E(a)}</h4><p>{E(b)}</p></div>' for a, b in PC_PRINCIPLES)

DATA = {'base': str(BASE), 'frame': list(FR), 'points': {k: PTS[k] for k, *_r in ELS}, 'bb': BB,
        'els': [{'k': k, 'name': n, 'v': str(v), 'b': b} for k, b, n, v, s in ELS], 'rest': str(REST['barn'] + REST['ph']),
        'alts': [{'id': a['id'], 'title': a['title'], 'v': str(a['v'])} for a in ALTS], 'span': SPAN, 'weeks': SCH['weeks'], 'name': PNAME, 'job': 'E260050', 'email': EMAIL,
        'sel': selbuild.data(SEL, SELX) if SEL else None,
        'steps': [{k: s.get(k) for k in ('label', 'title', 'sub', 'fig', 'cap', 'body', 'img', 'via', 'el', 'no', 'all', 'rail', 'note', 'bld') if s.get(k) is not None} for s in STEPS]}

BODY = f'''
<div class="loader" id="loader" aria-hidden="true"><div class="loader__in">{mark("loader__m")}<div class="loader__bar"><i id="loader-bar"></i></div></div></div>

<header class="nav" id="nav">
  <a class="nav__logo" href="#top" aria-label="Montana Contracting, back to the top"><img class="lg-white" src="assets/mc-white.png" alt="Montana Contracting" width="240" height="14"><img class="lg-ink" src="assets/mc-ink.png" alt="" width="240" height="14"><img class="lg-allw" src="assets/mc-white-all.png" alt="" width="240" height="14"></a>
  <nav class="nav__links" aria-label="Sections">{nav_links}</nav>
  <div class="nav__right">{pdf_link("nav__cta nav__pdf")}<button type="button" class="nav__menu" id="menu-btn" aria-expanded="false" aria-controls="menu" aria-label="Menu"><span></span><span></span></button></div>
  <i class="nav__prog" aria-hidden="true"></i>
</header>
<div class="menu" id="menu" inert data-lenis-prevent><button type="button" class="menu__close">Close</button><nav aria-label="Sections">{menu_links}</nav>{pdf_link("btn btn--ghost")}</div>

<main>
<section class="hero" id="top" aria-label="{E(PNAME)}">
  <div class="hero__bg"><picture><source media="(max-width: 760px)" srcset="assets/hero_p.jpg"><img id="hero-img" data-crit fetchpriority="high" src="assets/hero.jpg" alt="The Cooper Barn from the southeast, Montana rendering" width="2400" height="1198"></picture></div>
  <div class="hero__scrim"></div><div class="hero__shade"></div>
  <div class="hero__inner gut">
    <div class="tags"><span class="tag">Budget Presentation</span><span class="tag tag--ghost">{LOC}</span><span class="tag tag--ghost">Barn and Pool House</span></div>
    <h1 data-r>{E(PNAME)}</h1>
    <p class="eyebrow"><i></i>Crisp Architects · 194 Titicus Road</p>
  </div>
  <a class="hero__scroll" href="#project">Scroll<i></i></a>
</section>

<section class="intro" id="project" aria-label="Project">
  <div class="wrap">
    {kick("01 / Project")}
    <h2 class="sec-title" data-r>Project Details</h2>
    <div class="intro__grid">
      <dl class="facts" data-r="up">
        <div><dt>Location</dt><dd>{LOC}</dd></div>
        <div><dt>Type</dt><dd>Workshop Barn and Pool House</dd></div>
        <div><dt>Architect</dt><dd>Crisp Architects</dd></div>
        <div><dt>Budget</dt><dd class="num">{money(BASE)}</dd></div>
        <div><dt>Schedule</dt><dd>{SCH["weeks"]} Weeks</dd></div>
      </dl>
      <div class="desc"><p class="lede">{E(PROSE[0])}</p>{"".join(f"<p data-r=up>{E(p)}</p>" for p in PROSE[1:])}</div>
    </div>
  </div>
</section>

<section class="gallery" id="gallery" aria-label="Gallery">
  <div class="gallery__pin">
    <div class="wrap gallery__head">{kick("Gallery")}<span class="gallery__count num"><b>01</b> / {len(GAL):02d}</span></div>
    <div class="track">{"".join(f'<figure class="shot"><div class="shot__img"><img src="assets/{n}" alt="{E(c)}" loading="lazy" decoding="async"></div><figcaption><b>{k + 1:02d}</b>{E(c)}</figcaption></figure>' for k, (n, f, c) in enumerate(GAL))}</div>
  </div>
</section>

<section class="stage" id="stage" aria-label="Design">
  <span class="stage__anchor" id="design" style="top:0"></span>
  <div class="stage__pin">
    <div class="stage__model"><div class="frame" data-mk="brackets" style="--fa:{FR[0]} / {FR[1]}">{stage_imgs}
      <span class="mk" id="mk" aria-hidden="true"></span>
      <span class="mkb" id="mkb" aria-hidden="true"><i></i><i></i><i></i><i></i></span>
      <span class="mkl" id="mkl" aria-hidden="true"><i></i><b></b></span>
      <span class="mkn" id="mkn" aria-hidden="true"></span>
      {"".join(f'<span class="mks" data-k="{k}" data-b="{b}" aria-hidden="true">{(i if b == "barn" else i - nB) + 1}</span>' for i, (k, b, *_r) in enumerate(ELS))}
    </div></div>
    <div class="stage__copy" aria-live="polite">
      <p class="kick sc-label sw">{E(s0["label"])}</p>
      <h2 class="sc-title sw">{E(s0["title"])}</h2>
      <p class="sc-sub sw">{E(s0["sub"])}</p>
      <div class="sc-fig">{odo(s0["fig"], "stage-odo")}</div>
      <p class="sc-cap sw">{E(s0["cap"])}</p>
      <p class="sc-body sw">{E(s0["body"])}</p>
      <div class="elists">{elist("barn", B_ELS, "Barn")}{elist("ph", P_ELS, "Pool House")}</div>
    </div>
    <div class="stage__rail" role="group" aria-label="Steps">{rail}</div>
    <p class="stage__note">{E(s0["note"])}</p>
  </div>
</section>

<section class="budget" id="budget" aria-label="Budget">
  <div class="wrap">
    {kick("03 / Budget")}
    <div class="budget__fig">{odo(money(BASE), "budget-odo")}</div>
    <p class="budget__cap" data-r="up"><span><b>Base Budget</b> with Overhead and Profit</span><span><b class="num">{money(DIRECT)}</b> Direct Cost</span><span><b>15%</b> Overhead and Profit</span></p>
    <p class="budget__split" data-r="up"><span><b class="num">{money(BK["barn"]["base"])}</b> The Barn</span><span><b class="num">{money(BK["ph"]["base"])}</b> The Pool House</span><span><b class="num">{money(BK["shared"]["base"])}</b> Shared General Requirements</span></p>
    <div class="bview" data-view="el">
      <div class="sbar sbar--el" aria-hidden="true">{segs(EL_ITEMS, gcol)}</div>
      <div class="sbar sbar--mf" aria-hidden="true">{segs(MF_ITEMS, dcol)}</div>
      <div class="pills" role="tablist" aria-label="Budget view"><button type="button" role="tab" class="pill on" aria-selected="true" data-v="el">Elements</button><button type="button" role="tab" class="pill" aria-selected="false" data-v="mf">MasterFormat</button><span class="pills__hint">Lines at direct cost. Open any line for its detail.</span></div>
      <div class="brows brows--el" role="tabpanel" aria-label="Elements">{view_rows(EL_ITEMS, gcol)}</div>
      <div class="brows brows--mf" role="tabpanel" aria-label="MasterFormat" hidden>{view_rows(MF_ITEMS, dcol)}</div>
    </div>
    <div class="btot"><div class="btot__direct"><span>Direct Cost</span><span class="num">{money(DIRECT)}</span></div><div class="btot__op" style="--c:{OPC}"><span><i class="sq" aria-hidden="true"></i>Overhead and Profit, 15%</span><span class="num">{money(OP)}</span></div><div><span>Base Budget</span><span class="num">{money(BASE)}</span></div></div>
  </div>
</section>

{selbuild.section(SEL, SELX, kick) if SEL else selbuild.sheet(None, E)}

<section class="allow" id="allowances" aria-label="Allowances">
  <div class="wrap">
    {kick("03 / Allowances")}
    <div class="allow__top"><h2 class="sec-title" data-r>Allowances</h2><div class="allow__fig"><span class="big2 num" data-r="up">{money(ALLOW_TOTAL)}</span><p>Carried in the base budget, before overhead and profit.</p></div></div>
    <div class="aled" data-r="up">{allow_rows}</div>
    <div class="arate" data-r="up"><p class="arate__h">Material Allowances, Per SF</p><div class="arate__row">{rate_cells}</div></div>
  </div>
</section>

<section class="alts" id="alternates" aria-label="Alternates">
  <div class="wrap">
    {kick("03 / Alternates")}
    <div class="alts__grid">
      <div>
        <h2 class="sec-title" data-r>Alternates</h2>
        <p class="alts__lede" data-r="up">Priced with overhead and profit. None of the five is in the base budget.</p>
        <div class="alist">{alts}</div>
      </div>
      <aside class="atot" aria-label="Running total">
        <dl><div><dt>Base Budget</dt><dd class="num">{money(BASE)}</dd></div><div><dt>Requested</dt><dd class="num" id="alt-add">{money(0)}</dd></div></dl>
        <div class="atot__fig"><span class="lab">If Montana Confirms</span>{odo(money(BASE), "alt-odo")}</div>
        <button type="button" class="btn btn--navy" id="alt-send"><span>Send to Montana</span></button>
        <p class="atot__send">Your choices stay on this device until you send them.</p>
      </aside>
    </div>
  </div>
</section>

<section class="notes" id="notes" aria-label="Notes and Clarifications">
  <div class="wrap">
    {kick("03 / Notes and Clarifications")}
    <h2 class="sec-title" data-r>Notes and Clarifications</h2>
    <div class="ngrid">{ncol("Basis of Budget", BASIS)}{ncol("Clarifications", CLAR)}{ncol("Exclusions", EXCL)}</div>
  </div>
</section>

<section class="sched" id="schedule" aria-label="Schedule">
  <div class="sched__pin">
    <div class="wrap sched__grid">
      <aside class="sched__side">
        {kick("04 / Schedule")}
        <div class="sched__top">
          <span class="big num" data-r="up">{SCH["weeks"]}</span>
          <h2 data-r>Weeks on site from mobilization, both buildings.</h2>
        </div>
        <p class="sched__now"><span>Week</span>{odo("00", "week-odo")}<span class="sched__ph" id="week-ph">Site and Foundation</span></p>
        <ul class="g-legend">{g_legend}</ul>
      </aside>
      <div class="gantt" role="img" aria-label="Schedule by trade for the barn and the pool house, {SCH['weeks']} weeks from mobilization to work complete">
        <div class="g-ms">{g_ms}</div>
        <div class="g-body"><div class="g-grid" aria-hidden="true">{g_grid}<b class="g-cursor" id="g-cursor"><span class="g-cursor__l">Week <em id="g-cw">0</em></span></b></div>{g_rows}</div>
        <div class="g-axis">{g_axis}</div>
      </div>
    </div>
  </div>
</section>

<section class="appr" id="approach" aria-label="Approach">
  <div class="wrap">
    {kick("05 / Approach")}
    <h2 class="appr__h" data-r>How We Work</h2>
    <p class="appr__big" data-r="up">Every superintendent is a finish carpenter by trade.</p>
    <p class="appr__lede" data-r="up">The person running the schedule can pick up a tool and close a gap in it. Every detail is checked by someone who has built it.</p>
    <div class="yteam">
      <div class="yteam__head" data-r="up"><h3 class="yteam__h">Your Team</h3><p class="yteam__p">Three principals stay on this project from the first walkthrough to the final punch list. You are never handed off.</p></div>
      <div class="people">{people}</div>
    </div>
  </div>
  <div class="osx">
    {OS_LETTERS}
    <div class="wrap osx__in">
      {kick("Montana OS")}
      <h3 class="osx__h" data-r>One system for every job.</h3>
      <p class="osx__lede" data-r="up">We built our own operating system for scheduling, cost and field reporting. It gives the owner and the design team a live, honest view of the work as it happens.</p>
      <div class="osx__app" data-r="up">
        <div class="osx__bar"><span class="osx__brand">{mark("osx__bm")}Montana <b>OS</b></span><span class="osx__job">{E(PNAME)}</span><span class="osx__live"><i></i>Opens at Award</span></div>
        <div class="osx__tiles">{os_tiles}</div>
      </div>
      <ol class="osx__feat">{os_feat}</ol>
      <div class="osx__for">{os_for}</div>
    </div>
  </div>
</section>

<section class="pcf" id="pre-construction" aria-label="Pre-Construction">
  <div class="wrap">
    <div class="pcf__card" data-precon role="button" tabindex="0" aria-label="See our pre-construction process">
      {mark("pcf__m")}
      <div class="pcf__txt">
        <p class="kick kick--light"><i></i>Pre-Construction</p>
        <h2 class="pcf__h">See Our <span class="nb">Pre-Construction</span> Process</h2>
        <p class="pcf__p">Four phases. Each one locks more of the price. The last is a Guaranteed Maximum Price.</p>
        <span class="btn btn--white pcf__btn">See the Process {ARROW}</span>
      </div>
      <div class="pcf__viz" data-r="up" aria-hidden="true">{funnel("funnel funnel--d")}<div class="pcf__labs">{PCF_LABS}</div></div>
    </div>
  </div>
</section>

<section class="dusk" aria-label="Thank You">
  <picture><source media="(max-width: 760px)" srcset="assets/dusk_p.jpg"><img src="assets/dusk.jpg" alt="The Cooper Barn at dusk, lit from inside" loading="lazy" decoding="async"></picture>
  <div class="dusk__scrim"></div>
  <div class="dusk__in gut"><h2 data-r>Thank you for the opportunity to partner with you on this project.</h2></div>
</section>

<section class="cta" id="contact" aria-label="Next Step">
  {mark("cta__m")}
  <div class="wrap cta__in">
    {kick("Next Step", "kick--light")}
    <h2 data-r>{E(P["next_step"])}</h2>
    <div class="cta__act"><a class="btn btn--white mag" href="tel:+18453981778">Call {PHONE} {ARROW}</a></div>
    <dl class="contact"><div><dt>Phone</dt><dd>{PHONE}</dd></div><div><dt>Address</dt><dd>{E(ADDR)}</dd></div><div><dt>Email</dt><dd>{E(EMAIL)}</dd></div></dl>
  </div>
</section>
<footer class="foot"><div class="wrap foot__in"><img src="assets/mc-white-all.png" alt="Montana Contracting" loading="lazy"><span>General Contractor · Since 1984 · NY + NJ + FL</span><span>© 2026 Montana Contracting</span></div></footer>
</main>

<div class="pc" id="precon" inert>
  <div class="pc__panel" role="dialog" aria-modal="true" aria-labelledby="pc-h" data-lenis-prevent>
    <div class="pc__bar"><img src="assets/mc-white-all.png" alt="Montana Contracting" width="200" height="12"><span>Pre-Construction · Optional Next Step</span><button type="button" class="pc__x">Close</button></div>
    <div class="pc__hero">
      {mark("pc__m")}
      <div class="wrap">
        <p class="kick kick--light"><i></i>Pre-Construction</p>
        <h2 id="pc-h">Know the number<br><em>before</em> you build.</h2>
        <p class="pc__lede">We set the budget with you first, then price the building at every design stage until we can offer a Guaranteed Maximum Price.</p>
      </div>
    </div>
    <div class="pc__body">
      <section class="wrap pc__sec">
        <p class="kick"><i></i>01 · Every Phase</p>
        <h3>Your price gets more certain at every phase.</h3>
        <p class="pc__p">Each phase locks part of the price. At the end of Phase 04 the whole number is a Guaranteed Maximum Price.</p>
        <div class="pc__funnel">{funnel()}<div class="pcphs">{pc_cols}</div></div>
      </section>
      <section class="pc__alt"><div class="wrap pc__sec">
        <p class="kick"><i></i>02 · What Each Phase Delivers</p>
        <h3>Each phase makes the estimate more accurate.</h3>
        <p class="pc__p">Each phase ends with an updated budget and a decision about whether to keep going. You can stop at the end of any phase.</p>
        <div class="pccards">{pc_cards}</div>
      </div></section>
      <section class="wrap pc__sec">
        <p class="kick"><i></i>03 · Our Promises</p>
        <h3>Built in from the first phase.</h3>
        <div class="pcproms">{pc_prom}</div>
        <div class="pcprins">{pc_prin}</div>
      </section>
      <section class="pc__end"><div class="wrap pc__sec">
        <h3>Ask us about <span class="nb">pre-construction.</span></h3>
        <p class="pc__p">One of the three principals will walk you through it.</p>
        <dl class="contact"><div><dt>Phone</dt><dd>{PHONE}</dd></div><div><dt>Email</dt><dd>{E(EMAIL)}</dd></div><div><dt>Address</dt><dd>{E(ADDR)}</dd></div></dl>
      </div></section>
    </div>
  </div>
</div>

<script type="application/json" id="cb-data">{json.dumps(DATA).replace("</", "<\\/")}</script>
<script>{open(os.path.join(KIT, "vendor", "lenis", "lenis.min.js")).read()}</script>
<script>{open(os.path.join(HERE, "src", "app.js")).read()}</script>
'''

FONTS = f'''@font-face {{ font-family: "TT Commons Pro"; font-weight: 400; font-style: normal; font-display: block; src: url(data:font/woff2;base64,{font("TTCommonsPro-Regular.woff2")}) format("woff2"); }}
@font-face {{ font-family: "TT Commons Pro"; font-weight: 600 700; font-style: normal; font-display: block; src: url(data:font/woff2;base64,{font("TTCommonsPro-Bold.woff2")}) format("woff2"); }}
'''
CSS = FONTS + open(os.path.join(HERE, 'src', 'style.css')).read()
CSS += '''
.budget__split { display: flex; flex-wrap: wrap; gap: 6px 28px; margin: 6px 0 26px; color: var(--mute, #5b6173); font-size: 15px; }
.budget__split b { color: var(--ink, #0b1020); font-weight: 600; margin-right: 6px; }
.brows--mf .brow__t p, .brows--el .brow__t p { margin-top: 2px; }
.yteam { margin-top: clamp(3.5rem, 7vw, 6rem); padding-top: clamp(1.6rem, 3vw, 2.4rem); border-top: 2px solid var(--ink); }
.yteam__head { display: grid; grid-template-columns: minmax(0, 1fr) minmax(0, 1.2fr); gap: clamp(1rem, 3vw, 3rem); align-items: end; }
.yteam__h { font-size: clamp(2rem, 4.2vw, 3.8rem); font-weight: 600; letter-spacing: -.035em; line-height: 1; }
.yteam__p { color: var(--muted); font-size: clamp(1rem, 1.3vw, 1.2rem); line-height: 1.5; max-width: 44ch; }
.appr .yteam .people { margin-top: clamp(1.8rem, 4vw, 3rem); }
@media (max-width: 760px) { .yteam__head { grid-template-columns: 1fr; gap: .7rem; } }
.elists { display: grid; } .elists > .elist { grid-area: 1 / 1; }   /* one list per building, stacked in one place */
@media (max-width: 760px) { .g-ms div:not(:first-child):not(:last-child) { display: none; } }
@media (max-width: 600px) {   /* narrow windows and side panels keep the PDF button, as a compact pill */
  .nav__cta.nav__pdf { display: inline-flex; padding: .45rem .8rem; margin-right: .5rem; }
  .nav__pdf span { display: none; }
  .nav__pdf::after { content: "PDF"; }
}   /* five milestones: the phone keeps the first and last */
@media (min-width: 1181px) {   /* six rail steps: the rail stays under the model, clear of the copy column */
  .stage__rail { left: calc(var(--edge) + var(--copy-w) + var(--gap)); gap: 0 1.3rem; flex-wrap: wrap; }
  .stage__rail button { font-size: .66rem; letter-spacing: .12em; }
}
'''
TITLE = 'Cooper Barn and Pool House'
def tiny(name, w):
    import io
    im = Image.open(os.path.join(SITE, 'assets', name)).convert('RGB'); im = im.resize((w, max(1, round(im.height * w / im.width))), Image.LANCZOS)
    b = io.BytesIO(); im.save(b, 'JPEG', quality=60); return 'data:image/jpeg;base64,' + base64.b64encode(b.getvalue()).decode()
PH_D, PH_P = tiny('hero.jpg', 40), tiny('hero_p.jpg', 24)
PRELOAD = ('<link rel="preload" as="image" href="assets/hero.jpg" media="(min-width: 761px)" fetchpriority="high">\n'
           '<link rel="preload" as="image" href="assets/hero_p.jpg" media="(max-width: 760px)" fetchpriority="high">\n')
CSS += f'\n.hero__bg {{ background: #2b2f2a url({PH_D}) center / cover no-repeat; }}\n@media (max-width: 760px) {{ .hero__bg {{ background-image: url({PH_P}); }} }}\n'
HEAD = (PRELOAD + '<script>document.documentElement.classList.add("m")</script>\n' + f'<title>{TITLE}</title>\n<meta name="claude:viewport-fit-sides" content="cover">\n<link rel="icon" href="{FAVICON}">\n'
        f'<meta name="description" content="Montana Contracting budget presentation for The Cooper Barn and Pool House, 194 Titicus Road, North Salem, NY.">\n<style>\n{CSS}\n</style>\n')

# ================================================================ house rules on everything a client reads
BANNED = ['stunning', 'beautiful', 'breathtaking', 'luxury', 'luxurious', 'bespoke', 'elevated', 'curated', 'seamless', 'passion', 'dream', 'journey',
          'exceptional', 'world-class', 'premier', 'cutting-edge', 'unparalleled', 'empower', 'transform', 'peace of mind', 'potential risk', 'novaterra',
          'design move', 'nestled', 'state-of-the-art', 'elevate', 'proud to', 'excited', 'grateful', 'tbd']
text = re.sub(r'<script.*?</script>|<svg.*?</svg>', ' ', BODY, flags=re.S)
text = html.unescape(re.sub(r'<[^>]+>', ' ', text))
text += ' ' + ' '.join(str(v) for s in STEPS for v in s.values() if isinstance(v, str))
if SEL: text += ' ' + json.dumps(DATA['sel'])
bad = [name for ch, name in [('—', 'em dash'), ('–', 'en dash'), (';', 'semicolon'), ('!', 'exclamation'), ('?', 'question')] if ch in text]
bad += ['banned: ' + w for w in BANNED if re.search(r'\b' + re.escape(w) + r'\b', text.lower())]
if re.search(r'\b\w+, not \w+', text): bad.append('X, not Y')
if bad: die('copy rules: ' + ', '.join(bad))
if len(PROSE[0].split('. ')[0]) > 155: die('project snippet sentence is over 155 characters')
words = sum(len(p.split()) for p in PROSE)
if not 90 <= words <= 220: die(f'project description is {words} words')
for v in [money(BASE), money(DIRECT), money(OP), money(ALLOW_TOTAL)] + [money(e[3]) for e in ELS] + [money(BK[b]['base']) for b in BK] + [money(REST[b]) for b in REST] + [money(SITE_FDN[b]) for b in SITE_FDN]:
    if v not in BODY + json.dumps(DATA): die('figure missing from the page: ' + v)
_vis = re.sub(r'<script.*?</script>', '', BODY, flags=re.S)
_nc = re.search(r'\$\d[\d,]*(?![\d,]*\.\d\d)', _vis)
if _nc: die('a dollar figure without cents: ' + _vis[max(0, _nc.start() - 120):_nc.end() + 40])

os.makedirs(SITE, exist_ok=True)
open(os.path.join(SITE, 'index.html'), 'w').write(HEAD + BODY)
shutil.copytree(SITE, NET)
open(os.path.join(NET, 'index.html'), 'w').write(f'<!doctype html>\n<html lang="en">\n<head>\n<meta charset="utf-8">\n<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">\n<meta name="robots" content="noindex, nofollow">\n<meta name="mcc-send" content="netlify-forms">\n{HEAD}</head>\n<body>\n{BODY}\n</body>\n</html>\n')
open(os.path.join(NET, 'robots.txt'), 'w').write('User-agent: *\nDisallow: /\n')
z = os.path.join(HERE, 'E260050_Presentation_Netlify_v7.zip')
with zipfile.ZipFile(z, 'w', zipfile.ZIP_DEFLATED) as zf:
    for dp, _d, fs in os.walk(NET):
        for f in fs: zf.write(os.path.join(dp, f), os.path.relpath(os.path.join(dp, f), NET))
json.dump({'base': money(BASE), 'direct': money(DIRECT), 'op': money(OP), 'books': {b: {'base': money(BK[b]['base']), 'direct': money(BK[b]['direct'])} for b in BK},
           'elements': {f'{b} {n}': money(v) for k, b, n, v, s in ELS}, 'el_total': {b: money(EL_T[b]) for b in EL_T}, 'rest': {b: money(REST[b]) for b in REST},
           'site_fdn': {b: money(SITE_FDN[b]) for b in SITE_FDN}, 'allow_total': money(ALLOW_TOTAL), 'groups': {g['name']: money(g['v']) for g in GROUPS},
           'sov': {f"{dv['b']} {dv['n']} {dv['name']}": money(dv['v']) for dv in DIVS}, 'alternates': {a['title']: money(a['v']) for a in ALTS},
           'allowances': {a['label']: a['v'] for a in ALLOWS + RATES}, 'milestones': MS, 'phases': [(n, round(a, 2), round(b, 2)) for n, c, a, b in PH_SPANS]},
          open(os.path.join(HERE, 'figures.json'), 'w'), indent=1)
size = sum(os.path.getsize(os.path.join(dp, f)) for dp, _d, fs in os.walk(SITE) for f in fs)
print(f'built: page {os.path.getsize(os.path.join(SITE, "index.html")) / 1e6:.2f} MB, site {size / 1e6:.2f} MB, {len(STEPS)} steps, frame {FR}, crops {CROPS}, zip {os.path.getsize(z) / 1e6:.2f} MB')
print('base', money(BASE), {b: money(BK[b]['base']) for b in BK})
json.dump({'allowances': ALLOWS + RATES, 'allowance_total': str(ALLOW_TOTAL),
           'alternates': [{'id': a['id'], 'title': a['title'], 'note': a['note'], 'v': str(a['v'])} for a in ALTS], 've': []}, open(os.path.join(HERE, 'proposal_lists.json'), 'w'), indent=1)
