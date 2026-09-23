#!/usr/bin/env python3
"""Internal estimate workbook -- Estimate (base + alternates from row 401), SOV (base + alternates),
Notes & Clarifications, Estimate Review, Quantity Basis (hidden).

Formula rules carried from generate-estimate-from-unit-cost-report:
  * SUM() wraps every chained reference so "-" text never raises #VALUE!
  * markups compound on the running total, one percent per row, entered once on Estimate
  * Final Price (SOV) is an input column: it references nothing, so nothing can loop through it
"""
import json, math, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.pagebreak import Break

W = "FFFFFFFF"; BLACK = "FF000000"; CHAR = "FF3A3A3A"; MED = "FF7A7A7A"
LT = "FFE8E8E8"; VLT = "FFF5F5F5"; RED = "FFFFD6D6"; FINAL = "FFFFF6D5"; GRN = "FFD6F0D6"
FONT = "Calibri"; CUR = '$#,##0.00'
HAIR = Side("thin", color="FFB0B0B0"); MEDS = Side("medium", color=CHAR); BLKS = Side("medium", color=BLACK)
NONE = Side(style=None)

# Estimate grid: A CSI | B Scope | C Qty | D Unit | E Rate | F Total | G gutter | H:K Sub #1-4
COLW = [11, 56, 10, 9, 13, 14, 2.5, 15, 15, 15, 15]
NCOL = 11; GUT = 7; SUB1 = 8; TOT = 6; RATE = 5
MONEY = [TOT] + list(range(SUB1, NCOL + 1))
HEADS = ["CSI", "Scope Item", "Qty", "Unit", "Unit Rate", "Total", "", "Sub #1", "Sub #2", "Sub #3", "Sub #4"]
ALT_START = 401          # default; pushed to the next row ending in 01 if the base bid runs long


def fill(c, color): c.fill = PatternFill("solid", fgColor=color)


def ch(spec, mn=22, lp=13):
    n = max([1] + [math.ceil(len(str(t)) / c) for t, c in spec if t])
    return max(mn, n * lp)


def paint(ws, r1, c1, r2, c2, edge, inner=HAIR):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cur = ws.cell(r, c).border
            ws.cell(r, c).border = Border(
                left=edge if c == c1 else (inner or cur.left), right=edge if c == c2 else (inner or cur.right),
                top=edge if r == r1 else (inner or cur.top), bottom=edge if r == r2 else (inner or cur.bottom))


def erow(ws, r, c1, c2, top=None, bottom=None):
    for c in range(c1, c2 + 1):
        b = ws.cell(r, c).border
        ws.cell(r, c).border = Border(left=b.left, right=b.right, top=top or b.top, bottom=bottom or b.bottom)


def ecol(ws, r1, r2, c, l=None, rt=None):
    for r in range(r1, r2 + 1):
        b = ws.cell(r, c).border
        ws.cell(r, c).border = Border(left=l or b.left, right=rt or b.right, top=b.top, bottom=b.bottom)


def merge_clean(ws, merges):
    """Merge, then force every cell in the range (anchor + MergedCells) to carry ONLY the
    perimeter. openpyxl otherwise leaves inner edges on the anchor and Excel draws them
    straight through the merged text."""
    for r1, c1, r2, c2 in merges:
        lft = ws.cell(r1, c1).border.left; rgt = ws.cell(r1, c2).border.right
        top = ws.cell(r1, c1).border.top; bot = ws.cell(r2, c1).border.bottom
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                ws.cell(r, c).border = Border(left=lft if c == c1 else NONE, right=rgt if c == c2 else NONE,
                                              top=top if r == r1 else NONE, bottom=bot if r == r2 else NONE)


def block(ws, r, name, lines, M, total_label):
    """One trade / alternate block. Returns (row after block, total row)."""
    ws.row_dimensions[r].height = 7; r += 1
    h1, h2 = r, r + 1
    c = ws.cell(h1, 1, name.upper()); c.font = Font(FONT, 11, bold=True, color=W)
    c.alignment = Alignment("left", "center", wrap_text=True)
    for rr in (h1, h2):
        for cc in range(1, GUT): fill(ws.cell(rr, cc), MED)
    M.append((h1, 1, h2, TOT))
    for col in range(SUB1, NCOL + 1):
        a = ws.cell(h1, col, "Sub #%d" % (col - GUT)); a.font = Font(FONT, 10, bold=True); fill(a, LT)
        a.alignment = Alignment("center", "center")
        b = ws.cell(h2, col); b.number_format = CUR; fill(b, VLT); b.alignment = Alignment("center", "center")
    ws.row_dimensions[h1].height = 19; ws.row_dimensions[h2].height = 19
    r = h2 + 1; first = r; band = 0
    for ln in lines:
        bullets = ln.get("bullets") or []
        a = ws.cell(r, 1, ln.get("csi", "")); a.font = Font(FONT, 9, color=CHAR); a.alignment = Alignment("center", "center")
        s = ws.cell(r, 2, ln["scope"]); s.font = Font(FONT, 10, bold=bool(bullets)); s.alignment = Alignment("left", "center", wrap_text=True)
        q = ws.cell(r, 3, ln["qty"]); q.number_format = '#,##0.00'; q.font = Font(FONT, 10, bold=True); q.alignment = Alignment("right", "center")
        u = ws.cell(r, 4, ln["unit"]); u.font = Font(FONT, 10); u.alignment = Alignment("center", "center")
        ur = ws.cell(r, RATE, ln["rate"]); ur.number_format = CUR; ur.font = Font(FONT, 10); ur.alignment = Alignment("right", "center")
        t = ws.cell(r, TOT, '=IF(AND(ISNUMBER($C%d),ISNUMBER($E%d)),$C%d*$E%d,"")' % (r, r, r, r))
        t.number_format = CUR; t.font = Font(FONT, 10, bold=True); t.alignment = Alignment("right", "center")
        for col in range(SUB1, NCOL + 1):
            p = ws.cell(r, col); p.number_format = CUR; p.alignment = Alignment("right", "center")
        if band % 2:
            for cc in list(range(1, GUT)) + list(range(SUB1, NCOL + 1)): fill(ws.cell(r, cc), VLT)
        band += 1
        ws.row_dimensions[r].height = ch([(ln["scope"], COLW[1])]); r += 1
        for b in bullets:
            bc = ws.cell(r, 2, "- " + b); bc.font = Font(FONT, 9, italic=True, color=CHAR)
            bc.alignment = Alignment("left", "center", wrap_text=True, indent=2)
            ws.row_dimensions[r].height = ch([(b, COLW[1] - 6)], mn=14, lp=11); r += 1
    last = r - 1; tr = r
    c = ws.cell(tr, 1, total_label); c.font = Font(FONT, 10, bold=True); c.alignment = Alignment("left", "center")
    M.append((tr, 1, tr, 4))
    te = ws.cell(tr, TOT, '=IF(SUM(F%d:F%d)=0,"-",SUM(F%d:F%d))' % (first, last, first, last))
    te.number_format = CUR; te.font = Font(FONT, 10, bold=True); te.alignment = Alignment("right", "center")
    for col in range(SUB1, NCOL + 1):
        L = CL(col)
        p = ws.cell(tr, col, '=IF(SUM(%s%d:%s%d)=0,"-",SUM(%s%d:%s%d))' % (L, h2, L, last, L, h2, L, last))
        p.number_format = CUR; p.font = Font(FONT, 10, bold=True); p.alignment = Alignment("right", "center")
    for cc in list(range(1, GUT)) + list(range(SUB1, NCOL + 1)): fill(ws.cell(tr, cc), LT)
    ws.row_dimensions[tr].height = 19
    paint(ws, h1, 1, tr, NCOL, MEDS, inner=HAIR)
    erow(ws, h2, 1, NCOL, bottom=MEDS); erow(ws, tr, 1, NCOL, top=MEDS)
    ecol(ws, h1, tr, GUT, l=MEDS, rt=MEDS); ecol(ws, h1, tr, SUB1, l=MEDS)
    return tr + 1, tr


def alt_ladder(ws, t, r_m1, n_mk, labels, M):
    """General Conditions / Fee / Insurance / TOTAL under an alternate's SUBTOTAL row (t).
    Percentages link to the base-bid ladder, so markups are entered once."""
    r = t + 1
    for i in range(n_mk):
        rr = r + i
        lb = ws.cell(rr, 1, labels[i]); lb.font = Font(FONT, 10); lb.alignment = Alignment("left", "center")
        M.append((rr, 1, rr, 4))
        pc = ws.cell(rr, RATE, "=$E$%d" % (r_m1 + i)); pc.number_format = '0.00%'; pc.font = Font(FONT, 10, bold=True)
        pc.alignment = Alignment("right", "center")
        for col in MONEY:
            L = CL(col)
            base = "SUM(%s%d)" % (L, t) if i == 0 else "SUM(%s%d)+SUM(%s%d:%s%d)" % (L, t, L, r, L, rr - 1)
            c = ws.cell(rr, col, '=IF(OR($E%d="",SUM(%s%d)=0),"",ROUND((%s)*$E%d,2))' % (rr, L, t, base, rr))
            c.number_format = CUR; c.alignment = Alignment("right", "center")
        for cc in list(range(1, GUT)) + list(range(SUB1, NCOL + 1)): fill(ws.cell(rr, cc), VLT)
        ws.row_dimensions[rr].height = 18
    tot = r + n_mk
    c = ws.cell(tot, 1, "TOTAL"); c.font = Font(FONT, 11, bold=True); c.alignment = Alignment("left", "center")
    M.append((tot, 1, tot, 4))
    for col in MONEY:
        L = CL(col); agg = "SUM(%s%d,%s%d:%s%d)" % (L, t, L, r, L, tot - 1)
        c = ws.cell(tot, col, '=IF(%s=0,"-",ROUND(%s,2))' % (agg, agg))
        c.number_format = CUR; c.font = Font(FONT, 11, bold=True); c.alignment = Alignment("right", "center")
    for cc in list(range(1, GUT)) + list(range(SUB1, NCOL + 1)): fill(ws.cell(tot, cc), LT)
    ws.row_dimensions[tot].height = 20
    paint(ws, r, 1, tot, NCOL, MEDS, inner=HAIR)
    erow(ws, tot, 1, NCOL, top=MEDS)
    ecol(ws, r, tot, GUT, l=MEDS, rt=MEDS); ecol(ws, r, tot, SUB1, l=MEDS)
    return tot + 1


def estimate_sheet(wb, d):
    global ALT_START
    ws = wb.active; ws.title = "Estimate"
    for i, w in enumerate(COLW, 1): ws.column_dimensions[CL(i)].width = w
    for col in range(GUT, NCOL + 1): ws.column_dimensions[CL(col)].outline_level = 1
    M = []; LC = NCOL
    ws.cell(1, 1, d["project"]).font = Font(FONT, 18, bold=True); ws.row_dimensions[1].height = 28; M.append((1, 1, 1, LC))
    ws.cell(2, 1, d["address"] + "   |   " + d["subtitle"]).font = Font(FONT, 11, bold=True); M.append((2, 1, 2, LC))
    ws.cell(3, 1, "Date: %s" % d["date"]).font = Font(FONT, 10); M.append((3, 1, 3, 2))
    ws.cell(3, 3, "Prepared By: %s   |   Job %s   |   Client: %s" % (d["prepared_by"], d["job"], d["client"])).font = Font(FONT, 10)
    M.append((3, 3, 3, LC)); ws.row_dimensions[4].height = 5
    x = ws.cell(5, 1)   # legend text written once ALT_START is known
    x.font = Font(FONT, 9, italic=True, color=CHAR)
    for c in range(1, LC + 1): fill(ws.cell(5, c), VLT)
    ws.row_dimensions[5].height = 16; M.append((5, 1, 5, LC))
    x = ws.cell(6, 1, "Sub bids: enter one lump sum on the trade header row or price line by line; both roll up to the trade total.")
    x.font = Font(FONT, 9, bold=True, color=CHAR)
    for c in range(1, LC + 1): fill(ws.cell(6, c), LT)
    ws.row_dimensions[6].height = 18; M.append((6, 1, 6, LC))
    erow(ws, 5, 1, LC, top=MEDS); erow(ws, 6, 1, LC, bottom=MEDS); ws.row_dimensions[7].height = 7
    HDR = 8
    for col, t in enumerate(HEADS, 1):
        c = ws.cell(HDR, col, t); c.font = Font(FONT, 11, bold=True, color=W); fill(c, CHAR)
        c.alignment = Alignment("left" if col == 2 else "center", "center")
    ws.row_dimensions[HDR].height = 22
    paint(ws, HDR, 1, HDR, LC, BLKS, inner=Side("thin", color=CHAR)); ws.freeze_panes = "A%d" % (HDR + 1)

    r = HDR + 1; trade_rows = []
    for tr in d["trades"]:
        r, t = block(ws, r, tr["trade"], tr["lines"], M, "TRADE TOTAL")
        trade_rows.append((tr["trade"], ", ".join(tr["csi"]), t))

    # ---- SUBTOTAL / markup ladder / TOTAL (SUM-wrapped) ----
    ws.row_dimensions[r].height = 7; r += 1
    r_sub = r
    c = ws.cell(r_sub, 1, "SUBTOTAL"); c.font = Font(FONT, 11, bold=True); M.append((r_sub, 1, r_sub, 4))
    for col in MONEY:
        L = CL(col); ref = "SUM(%s)" % ",".join("%s%d" % (L, t) for _, _, t in trade_rows)
        cc = ws.cell(r_sub, col, '=IF(%s=0,"-",%s)' % (ref, ref)); cc.number_format = CUR
        cc.font = Font(FONT, 11, bold=True); cc.alignment = Alignment("right", "center")
    for cc in list(range(1, GUT)) + list(range(SUB1, NCOL + 1)): fill(ws.cell(r_sub, cc), LT)
    ws.row_dimensions[r_sub].height = 20; r += 1
    r_m1 = r; n_mk = len(d["markup_labels"])
    for i in range(n_mk):
        rr = r_m1 + i
        lb = ws.cell(rr, 1, d["markup_labels"][i]); lb.font = Font(FONT, 10); fill(lb, VLT); M.append((rr, 1, rr, 4))
        pc = ws.cell(rr, RATE, d["markup_pcts"][i]); pc.number_format = '0.00%'; pc.font = Font(FONT, 10, bold=True)
        pc.alignment = Alignment("right", "center"); fill(pc, VLT)
        for col in MONEY:
            L = CL(col)
            base = "SUM(%s%d)" % (L, r_sub) if i == 0 else "SUM(%s%d)+SUM(%s%d:%s%d)" % (L, r_sub, L, r_m1, L, rr - 1)
            f = '=IF(OR($E%d="",SUM(%s%d)=0),"",ROUND((%s)*$E%d,2))' % (rr, L, r_sub, base, rr)
            cc = ws.cell(rr, col, f); cc.number_format = CUR; cc.alignment = Alignment("right", "center")
        ws.row_dimensions[rr].height = 18
    r_tot = r_m1 + n_mk
    c = ws.cell(r_tot, 1, "TOTAL -- BASE BID"); c.font = Font(FONT, 12, bold=True); M.append((r_tot, 1, r_tot, 4))
    for col in MONEY:
        L = CL(col); agg = "SUM(%s%d,%s%d:%s%d)" % (L, r_sub, L, r_m1, L, r_tot - 1)
        cc = ws.cell(r_tot, col, '=IF(%s=0,"-",ROUND(%s,2))' % (agg, agg))
        cc.number_format = CUR; cc.font = Font(FONT, 12, bold=True); cc.alignment = Alignment("right", "center")
    for cc in list(range(1, GUT)) + list(range(SUB1, NCOL + 1)): fill(ws.cell(r_tot, cc), LT)
    ws.row_dimensions[r_tot].height = 24
    paint(ws, r_sub, 1, r_tot, NCOL, MEDS, inner=HAIR)
    erow(ws, r_sub, 1, NCOL, bottom=MEDS); erow(ws, r_tot, 1, NCOL, top=MEDS)
    ecol(ws, r_sub, r_tot, GUT, l=MEDS, rt=MEDS); ecol(ws, r_sub, r_tot, SUB1, l=MEDS)
    while r_tot >= ALT_START - 1: ALT_START += 100
    ws.cell(5, 1).value = ("F/I = Furnish and Install  |  F/O = Furnish Only  |  I/O = Install Only  |  Remove = Demolition  |  "
               "Provide = Services  |  LS = Lump Sum  |  Alternates begin at row %d" % ALT_START)

    # ---- ALTERNATES from ALT_START (401 unless the base runs long) ----
    r = ALT_START
    band = ws.cell(r, 1, "ALTERNATES  --  markups link to the base bid percentages. Final Price and proposal comparison are on the SOV.")
    band.font = Font(FONT, 12, bold=True, color=W); band.alignment = Alignment("left", "center")
    for c in range(1, NCOL + 1): fill(ws.cell(r, c), CHAR)
    paint(ws, r, 1, r, NCOL, BLKS, inner=None); M.append((r, 1, r, NCOL)); ws.row_dimensions[r].height = 24
    ws.row_breaks.append(Break(id=ALT_START - 1))
    r += 1; alt_rows = []
    for a in d["alts"]:
        tag = ("PROPOSAL " + a["proposal_no"]) if a["proposal"] is not None else "NOT ON PROPOSAL SCHEDULE"
        r, t = block(ws, r, "%s  |  %s" % (a["trade"], tag), a["lines"], M, "SUBTOTAL")
        ws.cell(t, 1).font = Font(FONT, 11, bold=True)
        r = alt_ladder(ws, t, r_m1, n_mk, d["markup_labels"], M)
        alt_rows.append(t)
    last_row = r - 1

    merge_clean(ws, M)
    ws.print_area = "A1:K%d" % last_row
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4; ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_title_rows = "1:%d" % HDR
    ws.oddHeader.left.text = ("%s -- Estimate" % d["project"]).replace("&", "&&"); ws.oddHeader.right.text = "Page &P of &N"
    ws.oddFooter.center.text = d["prepared_by"]; ws.oddFooter.right.text = d["date"]
    ws.sheet_properties.outlinePr.summaryRight = False; ws.sheet_view.showGridLines = False
    return trade_rows, r_m1, n_mk, alt_rows


def title_block(ws, d, title, ncols, note):
    LC = CL(ncols)
    ws.cell(1, 1, "%s -- %s" % (d["project"], title)).font = Font(FONT, 16, bold=True)
    ws.row_dimensions[1].height = 24; ws.merge_cells("A1:%s1" % LC)
    x = ws.cell(2, 1, "Job %s  |  %s  |  %s  |  %s" % (d["job"], d["address"], d["prepared_by"], d["date"]))
    x.font = Font(FONT, 10, italic=True)
    for c in range(1, ncols + 1): fill(ws.cell(2, c), VLT)
    ws.row_dimensions[2].height = 18; ws.merge_cells("A2:%s2" % LC)
    if note:
        x = ws.cell(3, 1, note); x.font = Font(FONT, 9, italic=True, color=CHAR)
        x.alignment = Alignment("left", "center", wrap_text=True); ws.merge_cells("A3:%s3" % LC)
        for c in range(1, ncols + 1): fill(ws.cell(3, c), VLT)
        ws.row_dimensions[3].height = 26
    erow(ws, 2, 1, ncols, top=MEDS); erow(ws, 3 if note else 2, 1, ncols, bottom=MEDS)


def header_row(ws, r, heads, left_cols):
    for col, t in enumerate(heads, 1):
        c = ws.cell(r, col, t); c.font = Font(FONT, 11, bold=True, color=W); fill(c, CHAR)
        c.alignment = Alignment("left" if col in left_cols else "center", "center", wrap_text=True)
    ws.row_dimensions[r].height = 30
    paint(ws, r, 1, r, len(heads), BLKS, inner=Side("thin", color=CHAR))


def sov_sheet(wb, d, trade_rows, r_m1, n_mk, alt_rows):
    sv = wb.create_sheet("SOV"); NC = 11
    for i, w in enumerate([20, 32, 16, 16, 15, 15, 15, 15, 15, 16, 16], 1): sv.column_dimensions[CL(i)].width = w
    title_block(sv, d, "Schedule of Values", NC,
                "Our Total and the sub columns link to the Estimate sheet. Final Price is entered by hand per trade and runs "
                "through the same Subtotal / General Conditions / Fee / Insurance ladder. Markup percentages are entered once "
                "on the Estimate sheet. The low sub bid in each row is highlighted.")
    shr = 5
    header_row(sv, shr, ["CSI", "Trade", "Our Total", "Final Price", "Sub #1", "Sub #2", "Sub #3", "Sub #4",
                         "Low Bid", "Low Bidder", "Low Bid vs Ours"], left_cols=(1, 2))
    sv.freeze_panes = "C%d" % (shr + 1)
    rr = shr + 1; first = rr
    for name, csi, t in trade_rows:
        a = sv.cell(rr, 1, csi); a.font = Font(FONT, 9, color=CHAR); a.alignment = Alignment("left", "center", wrap_text=True)
        b = sv.cell(rr, 2, name); b.font = Font(FONT, 10, bold=True); b.alignment = Alignment("left", "center", wrap_text=True)
        c = sv.cell(rr, 3, "=Estimate!F%d" % t); c.number_format = CUR; c.font = Font(FONT, 10, bold=True)
        fp = sv.cell(rr, 4); fp.number_format = CUR; fp.font = Font(FONT, 10, bold=True); fill(fp, FINAL)
        for i, col in enumerate(range(SUB1, NCOL + 1)):
            L = CL(col)
            sv.cell(rr, 5 + i, '=IF(Estimate!%s%d="-","",Estimate!%s%d)' % (L, t, L, t)).number_format = CUR
        sv.cell(rr, 9, '=IF(COUNT(E%d:H%d)=0,"",MIN(E%d:H%d))' % (rr, rr, rr, rr)).number_format = CUR
        sv.cell(rr, 9).font = Font(FONT, 10, bold=True)
        sv.cell(rr, 10, '=IF(I%d="","",INDEX($E$%d:$H$%d,MATCH(I%d,E%d:H%d,0)))' % (rr, shr, shr, rr, rr, rr))
        sv.cell(rr, 11, '=IF(OR(I%d="",N(C%d)=0),"",I%d-C%d)' % (rr, rr, rr, rr)).number_format = CUR
        for col in range(3, NC + 1): sv.cell(rr, col).alignment = Alignment("center" if col == 10 else "right", "center")
        sv.row_dimensions[rr].height = ch([(name, 32), (csi, 20)], mn=20, lp=13); rr += 1
    last = rr - 1
    sv.conditional_formatting.add("E%d:H%d" % (first, last), FormulaRule(
        formula=['AND(ISNUMBER(E%d),COUNT($E%d:$H%d)>0,E%d=MIN($E%d:$H%d))' % (first, first, first, first, first, first)],
        fill=PatternFill("solid", fgColor=GRN), font=Font(bold=True)))
    # ladder: C Our Total, D Final Price (always shown), E:H subs (blank until bid)
    s_sub = rr
    sv.cell(s_sub, 1, "SUBTOTAL").font = Font(FONT, 11, bold=True); sv.merge_cells(start_row=s_sub, start_column=1, end_row=s_sub, end_column=2)
    for col in range(3, 9):
        L = CL(col); rng = "%s%d:%s%d" % (L, first, L, last)
        f = "=SUM(%s)" % rng if col in (3, 4) else '=IF(SUM(%s)=0,"",SUM(%s))' % (rng, rng)
        cc = sv.cell(s_sub, col, f); cc.number_format = CUR; cc.font = Font(FONT, 11, bold=True); cc.alignment = Alignment("right", "center")
    for col in range(1, NC + 1): fill(sv.cell(s_sub, col), LT)
    sv.row_dimensions[s_sub].height = 20
    s_m1 = s_sub + 1
    for i in range(n_mk):
        rw = s_m1 + i
        sv.cell(rw, 1, "=Estimate!A%d" % (r_m1 + i)).font = Font(FONT, 10)
        p = sv.cell(rw, 2, "=Estimate!$E$%d" % (r_m1 + i)); p.number_format = '0.00%'; p.font = Font(FONT, 10, bold=True)
        p.alignment = Alignment("right", "center")
        for col in range(3, 9):
            L = CL(col)
            base = "SUM(%s%d)" % (L, s_sub) if i == 0 else "SUM(%s%d)+SUM(%s%d:%s%d)" % (L, s_sub, L, s_m1, L, rw - 1)
            if col in (3, 4):
                f = "=ROUND((%s)*$B%d,2)" % (base, rw)
            else:
                f = '=IF(SUM(%s%d)=0,"",ROUND((%s)*$B%d,2))' % (L, s_sub, base, rw)
            cc = sv.cell(rw, col, f); cc.number_format = CUR; cc.alignment = Alignment("right", "center")
        for col in range(1, NC + 1): fill(sv.cell(rw, col), VLT)
        sv.row_dimensions[rw].height = 18
    s_tot = s_m1 + n_mk
    sv.cell(s_tot, 1, "TOTAL -- BASE BID").font = Font(FONT, 12, bold=True); sv.merge_cells(start_row=s_tot, start_column=1, end_row=s_tot, end_column=2)
    for col in range(3, 9):
        L = CL(col); agg = "SUM(%s%d,%s%d:%s%d)" % (L, s_sub, L, s_m1, L, s_tot - 1)
        f = "=ROUND(%s,2)" % agg if col in (3, 4) else '=IF(%s=0,"",ROUND(%s,2))' % (agg, agg)
        cc = sv.cell(s_tot, col, f); cc.number_format = CUR; cc.font = Font(FONT, 12, bold=True); cc.alignment = Alignment("right", "center")
    for col in range(1, NC + 1): fill(sv.cell(s_tot, col), LT)
    sv.row_dimensions[s_tot].height = 24
    for rw in range(s_sub, s_tot + 1): fill(sv.cell(rw, 4), FINAL)
    paint(sv, first, 1, s_tot, NC, MEDS, inner=HAIR)
    erow(sv, last, 1, NC, bottom=MEDS); erow(sv, s_tot, 1, NC, top=MEDS); ecol(sv, shr, s_tot, 4, l=MEDS, rt=MEDS)

    # ---- ALTERNATES table ----
    ar = s_tot + 3
    sv.row_breaks.append(Break(id=ar - 1))
    t = sv.cell(ar, 1, "ALTERNATES"); t.font = Font(FONT, 14, bold=True)
    x = sv.cell(ar + 1, 1, "Our Direct links to the alternate blocks on the Estimate sheet (row %d on). Enter a Final Price (direct cost) to "
                           "override ours: General Conditions, Fee and Insurance then compound on the Final Price. Variance = Proposal "
                           "Amount minus Burdened Total. Red rows are priced but not on the issued proposal." % ALT_START)
    x.font = Font(FONT, 9, italic=True, color=CHAR); x.alignment = Alignment("left", "center", wrap_text=True)
    sv.merge_cells(start_row=ar + 1, start_column=1, end_row=ar + 1, end_column=NC); sv.row_dimensions[ar + 1].height = 28
    ahr = ar + 2
    header_row(sv, ahr, ["CSI", "Alternate", "Our Direct", "Final Price", "General Conditions", "Fee", "Insurance",
                         "Burdened Total", "Proposal Amount", "Variance", "Proposal Ref"], left_cols=(1, 2))
    pct = ["Estimate!$E$%d" % (r_m1 + i) for i in range(3)]   # requires exactly 3 markups: GC, Fee, Insurance
    rr = ahr + 1; af = rr
    for a, t in zip(d["alts"], alt_rows):
        miss = a["proposal"] is None
        base = "IF(ISNUMBER(D%d),D%d,C%d)" % (rr, rr, rr)
        vals = [", ".join(a["csi"]), a["trade"], "=SUM(Estimate!F%d)" % t, None,
                "=ROUND(%s*%s,2)" % (base, pct[0]),
                "=ROUND((%s+E%d)*%s,2)" % (base, rr, pct[1]),
                "=ROUND((%s+E%d+F%d)*%s,2)" % (base, rr, rr, pct[2]),
                "=ROUND(%s+E%d+F%d+G%d,2)" % (base, rr, rr, rr),
                a["proposal"], '=IF(I%d="","",I%d-H%d)' % (rr, rr, rr), a["proposal_no"]]
        for col, v in enumerate(vals, 1):
            c = sv.cell(rr, col, v)
            c.font = Font(FONT, 9 if col in (1, 11) else 10, bold=(col in (2, 8)), color=(CHAR if col == 1 else None))
            if 3 <= col <= 10: c.number_format = CUR; c.alignment = Alignment("right", "center")
            else: c.alignment = Alignment("left", "center", wrap_text=True)
            fill(c, RED if miss else (VLT if (rr - af) % 2 else W))
        fill(sv.cell(rr, 4), FINAL)
        sv.row_dimensions[rr].height = ch([(a["trade"], 32), (", ".join(a["csi"]), 20)], mn=20, lp=13); rr += 1
    al = rr - 1
    rows = [("TOTAL -- ALTERNATES ON PROPOSAL", '=SUMIFS(H%d:H%d,I%d:I%d,"<>")' % (af, al, af, al), "=SUM(I%d:I%d)" % (af, al), LT),
            ("TOTAL -- PRICED, NOT ON PROPOSAL", '=SUMIFS(H%d:H%d,I%d:I%d,"")' % (af, al, af, al), None, RED)]
    for label, fh, fi, col_ in rows:
        sv.cell(rr, 1, label).font = Font(FONT, 10, bold=True)
        sv.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=2)
        sv.cell(rr, 8, fh); sv.cell(rr, 9, fi)
        for col in range(1, NC + 1):
            c = sv.cell(rr, col); fill(c, col_)
            if col >= 3: c.number_format = CUR; c.font = Font(FONT, 10, bold=True); c.alignment = Alignment("right", "center")
        sv.row_dimensions[rr].height = 20; rr += 1
    paint(sv, af, 1, rr - 1, NC, MEDS, inner=HAIR); erow(sv, al, 1, NC, bottom=MEDS); ecol(sv, ahr, al, 4, l=MEDS, rt=MEDS)
    sv.page_setup.orientation = "landscape"; sv.page_setup.fitToWidth = 1; sv.page_setup.fitToHeight = 0
    sv.sheet_properties.pageSetUpPr.fitToPage = True; sv.print_title_rows = "1:2"; sv.sheet_view.showGridLines = False


def notes_sheet(wb, d):
    n = wb.create_sheet("Notes & Clarifications")
    for col, w in zip("ABC", [16, 34, 104]): n.column_dimensions[col].width = w
    title_block(n, d, "Notes & Clarifications", 3, None)
    hr = 4; header_row(n, hr, ["Type", "Item", "Clarification"], left_cols=(1, 2, 3)); n.freeze_panes = "A5"
    rr = hr + 1
    for nt in d["notes"]:
        for col, v in enumerate([nt["type"], nt["item"], nt["text"]], 1):
            c = n.cell(rr, col, v); c.font = Font(FONT, 10, bold=(col == 1))
            c.alignment = Alignment("left", "top", wrap_text=True); fill(c, RED if nt["type"] == "EXCLUSION" else VLT)
        n.row_dimensions[rr].height = ch([(nt["item"], 34), (nt["text"], 104)], mn=26); rr += 1
    paint(n, hr + 1, 1, rr - 1, 3, MEDS, inner=HAIR)
    n.page_setup.orientation = "landscape"; n.page_setup.fitToWidth = 1; n.page_setup.fitToHeight = 0
    n.sheet_properties.pageSetUpPr.fitToPage = True; n.print_title_rows = "1:%d" % hr; n.sheet_view.showGridLines = False


def review_sheet(wb, d):
    ws = wb.create_sheet("Estimate Review")
    for i, w in enumerate([5, 10, 20, 78, 16, 44], 1): ws.column_dimensions[CL(i)].width = w
    title_block(ws, d, "Estimate Review (Internal)", 6, "Internal only. Items found reconciling the estimate, the cost detail and the issued proposal. Do not send.")
    ws.cell(3, 1).font = Font(FONT, 9, bold=True, color="FFB00000")
    hr = 5; header_row(ws, hr, ["No.", "Priority", "Area", "Finding", "Direct $ at Stake", "Action Before Issue"], left_cols=(3, 4, 6))
    ws.freeze_panes = "A6"; r = hr + 1
    for i, f in enumerate(d["review"], 1):
        for col, v in enumerate([i, f["p"], f["area"], f["finding"], f.get("amt"), f["action"]], 1):
            c = ws.cell(r, col, v); c.font = Font(FONT, 10, bold=(col in (1, 2)))
            c.alignment = Alignment("center" if col <= 2 else "left", "top", wrap_text=True)
            if col == 5: c.number_format = CUR; c.alignment = Alignment("right", "top")
            fill(c, RED if f["p"] == "HIGH" else (VLT if f["p"] == "MED" else W))
        ws.row_dimensions[r].height = ch([(f["finding"], 74), (f["action"], 40)], mn=26); r += 1
    paint(ws, hr + 1, 1, r - 1, 6, MEDS, inner=HAIR)
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.print_title_rows = "1:%d" % hr; ws.sheet_view.showGridLines = False


def basis_sheet(wb, d):
    h = wb.create_sheet("Quantity Basis")
    for col, w in zip("ABCDEF", [10, 22, 54, 12, 8, 70]): h.column_dimensions[col].width = w
    h.cell(1, 1, "Internal -- quantity basis by line").font = Font(FONT, 14, bold=True); h.merge_cells("A1:F1")
    hr = 3; header_row(h, hr, ["CSI", "Trade", "Scope Item", "Qty", "Unit", "Basis"], left_cols=(1, 2, 3, 6)); rr = hr + 1
    for b in d["basis"]:
        for col, v in enumerate([b["csi"], b["trade"], b["scope"], b["qty"], b["unit"], b["basis"]], 1):
            c = h.cell(rr, col, v); c.font = Font(FONT, 10); c.alignment = Alignment("left", "top", wrap_text=True); fill(c, VLT)
        h.row_dimensions[rr].height = ch([(b["scope"], 54), (b["basis"], 70)], mn=24); rr += 1
    paint(h, hr, 1, rr - 1, 6, MEDS, inner=HAIR); h.sheet_view.showGridLines = False; h.sheet_state = "hidden"


def build(d, out):
    wb = Workbook()
    trade_rows, r_m1, n_mk, alt_rows = estimate_sheet(wb, d)
    sov_sheet(wb, d, trade_rows, r_m1, n_mk, alt_rows)
    notes_sheet(wb, d)
    review_sheet(wb, d)
    basis_sheet(wb, d)
    wb.active = 0
    wb.save(out)
    print("wrote", out, "| base total row", r_m1 + n_mk, "| alternates from", ALT_START)


if __name__ == "__main__":
    build(json.load(open(sys.argv[1])), sys.argv[2])
